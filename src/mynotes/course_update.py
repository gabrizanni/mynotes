"""One cumulative Word transaction with conservative three-way preservation."""
from __future__ import annotations

from contextlib import ExitStack
from copy import deepcopy
from difflib import SequenceMatcher
import hashlib
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import uuid

from docx import Document
from lxml import etree

from .utils import atomic_json_write, file_sha256, now_compact_utc
from .workflow import WorkflowState, _generation_lock, parse_lecture_id


def _index_path(workspace, code):
    return workspace.state_dir / 'published-baselines' / (code + '.json')


def baseline_index(workspace, code):
    path = _index_path(workspace, code)
    return json.loads(path.read_text()) if path.is_file() else {}


def _section(document, lecture):
    from .publication import _boundaries
    bounds = _boundaries(document, lecture)
    return list(document._body._element)[bounds[0]:bounds[1]+1] if bounds else []


def _signature(element, document):
    node = deepcopy(element)
    # Word rewrites rsid/paragraph IDs and relationship IDs during an ordinary
    # save. Compare actual relationship content rather than incidental names.
    for child in node.iter():
        for key in list(child.attrib):
            if key.rsplit('}', 1)[-1] in {'rsidR', 'rsidRPr', 'rsidP', 'rsidDel', 'rsidRDefault', 'paraId', 'textId'}:
                del child.attrib[key]
            elif key.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
                rel = document.part.rels.get(child.attrib[key])
                if rel:
                    child.attrib[key] = rel.target_ref if rel.is_external else hashlib.sha256(rel.target_part.blob).hexdigest()
    return etree.tostring(node, method='c14n')


def _edits(before, after):
    return [(a, b, c, d) for tag, a, b, c, d in SequenceMatcher(None, before, after, autojunk=False).get_opcodes() if tag != 'equal']


def _overlap(a, b, c, d):
    if a == b and c == d:
        return a == c
    if a == b:
        return c < a < d
    if c == d:
        return a < c < b
    return max(a, c) < min(b, d)


def _merge_text(base, personal, generated):
    # Token-level diff3 retains independent additions even inside a rewritten
    # paragraph. Overlapping contradictory edits are left for source review.
    import re
    a, b, c = [re.findall(r'\s+|\S+', value) for value in (base, personal, generated)]
    user, writer = _edits(a, b), _edits(a, c)
    if any(_overlap(i, j, x, y) and b[k:l] != c[z:w] for i, j, k, l in user for x, y, z, w in writer):
        return None
    combined = [(i, j, c[k:l]) for i, j, k, l in writer]
    for i, j, k, l in user:
        if not any(i == x and j == y and b[k:l] == c[z:w] for x, y, z, w in writer):
            combined.append((i, j, b[k:l]))
    result = list(a)
    for i, j, replacement in sorted(combined, reverse=True):
        result[i:j] = replacement
    return ''.join(result)


def _paragraph_text(node):
    return ''.join(node.itertext()) if node.tag.endswith('}p') else None


def _plain_text(node):
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    if not node.tag.endswith('}p') or etree.XPath('.//w:drawing|.//w:fldChar|.//w:hyperlink|.//m:oMath', namespaces={**ns, 'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'})(node):
        return None
    return ''.join(etree.XPath('.//w:t/text()', namespaces=ns)(node))


def _set_text(node, text):
    # Keep the generated paragraph/run formatting for unchanged portions.
    elements = etree.XPath('.//w:t', namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})(node)
    if not elements:
        return False
    old = ''.join(element.text or '' for element in elements)
    edits = _edits(list(old), list(text))
    for a, b, c, d in reversed(edits):
        spans = []; position = 0
        for element in elements:
            value = element.text or ''; spans.append((element, position, position+len(value), value)); position += len(value)
        first = next((index for index, (_, start, end, _) in enumerate(spans) if start <= a <= end), len(spans)-1)
        for index, (element, position, end, value) in enumerate(spans):
            if index < first or (position >= b and index != first): continue
            start = min(len(value), max(0, a-position)); stop = min(len(value), max(start, b-position))
            element.text = value[:start] + (text[c:d] if index == first else '') + value[stop:]
            element.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    return True


def merge_section(base, current, generated, base_doc, current_doc, generated_doc):
    """Return merged elements or a concrete conflict; never discard user changes."""
    a = [_signature(node, base_doc) for node in base]
    b = [_signature(node, current_doc) for node in current]
    c = [_signature(node, generated_doc) for node in generated]
    if a == b:
        return generated, []
    if b == c:
        return current, []
    user, writer = _edits(a, b), _edits(a, c)
    operations = [(i, j, list(generated[k:l])) for i, j, k, l in writer]
    preserved = []
    for i, j, k, l in user:
        overlaps = [(x, y, z, w) for x, y, z, w in writer if _overlap(i, j, x, y)]
        if overlaps:
            if len(overlaps) == 1 and j-i == l-k == 1:
                x, y, z, w = overlaps[0]
                if (i, j) == (x, y) and w-z == 1:
                    values = [_plain_text(node) for node in (base[i], current[k], generated[z])]
                    if all(value is not None for value in values):
                        text = _merge_text(*values)
                        if text is not None:
                            node = deepcopy(generated[z])
                            if _set_text(node, text):
                                operations = [op for op in operations if op[:2] != (x, y)]
                                operations.append((i, j, [node])); preserved.append({'from': i, 'to': j}); continue
            text = lambda nodes: '\n'.join(''.join(etree.XPath('.//w:t/text()', namespaces={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})(n)) for n in nodes)
            return None, [{'from': i, 'to': j, 'reason': 'Personal and generated changes overlap; current lesson retained for source-grounded reconciliation.',
                           'baseline_text': text(base[i:j]), 'current_text': text(current[k:l]),
                           'proposed_text': text([node for x,y,z,w in overlaps for node in generated[z:w]]),
                           'source_content_only': True}]
        operations.append((i, j, list(current[k:l]))); preserved.append({'from': i, 'to': j})
    result = list(base)
    # Unchanged base elements are taken from current, whose relationships exist
    # in generated_doc (it was opened from the current Word).
    matcher = SequenceMatcher(None, a, b, autojunk=False)
    for tag, i, j, k, l in matcher.get_opcodes():
        if tag == 'equal': result[i:j] = current[k:l]
    for i, j, replacement in sorted(operations, key=lambda op: (op[0], op[1]), reverse=True):
        result[i:j] = replacement
    return result, preserved


def _baseline(workspace, code, lecture, expected):
    from .publication import _transactions, _target, _course, PublicationError
    record = baseline_index(workspace, code).get(lecture)
    if record:
        path = Path(record['path'])
        if not path.is_file() or file_sha256(path) != record['sha256']:
            raise PublicationError('Generated Word baseline is missing or changed; personal edits cannot be compared safely')
        return Document(path)
    target = _target(_course(workspace.configuration_store().load(), code), workspace)
    for transaction in _transactions(workspace, target):
        path = Path(transaction.get('candidate', ''))
        if transaction.get('target_hash') == expected and path.is_file() and file_sha256(path) == expected:
            return Document(path)
    if target.is_file() and file_sha256(target) == expected:
        return Document(target)
    return None


def build_course_revision(workspace, lectures):
    from . import publication as pub
    lectures = sorted(set(lecture.upper() for lecture in lectures), key=lambda value: parse_lecture_id(value)[1])
    if not lectures or len({parse_lecture_id(value)[0] for value in lectures}) != 1:
        raise pub.PublicationError('A cumulative update requires lessons from one course')
    code = parse_lecture_id(lectures[0])[0]
    course = pub._course(workspace.configuration_store().load(), code)
    target = pub._target(course, workspace)
    state = WorkflowState(workspace.database_path)
    try: expected = state.expected_output_hash(target)
    finally: state.close()
    if expected and not target.is_file():
        raise pub.PublicationError('Saved final handout is missing; locate or restore it before updating')
    baseline_hash = file_sha256(target) if target.is_file() else None
    current = Document(target) if target.is_file() else pub._new_document(course, pub.effective_style(workspace.configuration_store().load(), code))
    generated = Document(target) if target.is_file() else pub._new_document(course, pub.effective_style(workspace.configuration_store().load(), code))
    sources, conflicts, preserved, applied = {}, {}, {}, []
    payloads = {}
    for lecture in lectures:
        path = workspace.runs_dir / lecture / 'lecture_notes.candidate.json'
        payload, pack = pub._load_candidate(workspace, lecture)
        related = set(payload.get('related_lessons', [])) - set(lectures)
        if related:
            raise pub.PublicationError('Related published lessons require the same coordinated revision: ' + ', '.join(sorted(related)))
        profile = pub.effective_style(workspace.configuration_store().load(), code)
        limit = payload.get('request_options', {}).get('image_limit')
        if limit is None: limit = profile.get('image_limit')
        count = sum(v.get('decision') == 'include' for v in payload['notes'].get('visual_decisions', []))
        if (limit is not None and (type(limit) is not int or limit < 0 or count > limit)) or (profile.get('image_policy') == 'none' and count):
            raise pub.PublicationError('Candidate exceeds the effective image policy; review it before cumulative update')
        payloads[lecture] = (payload, pack)
        sources[lecture] = file_sha256(path)
        pub._replace_or_append(generated, payload['notes'], pack, language=course.get('languages', {}).get('notes', 'en'), recap_language=pub.effective_style(workspace.configuration_store().load(), code).get('recap_language', 'it'))
    directory = workspace.runs_dir / '_course_updates' / code / uuid.uuid4().hex
    directory.mkdir(parents=True, exist_ok=True)
    raw = directory / 'generated.docx'; generated.save(raw)
    for lecture in lectures:
        old = _section(current, lecture)
        if not old:
            applied.append(lecture); continue
        base_doc = _baseline(workspace, code, lecture, expected)
        if base_doc is None:
            conflicts[lecture] = [{'reason': 'Existing lesson has no verified generated baseline; its current text was preserved.'}]
            merged = old
        else:
            merged, details = merge_section(_section(base_doc, lecture), old, _section(generated, lecture), base_doc, current, generated)
            if merged is None:
                conflicts[lecture] = details; merged = old
            else:
                preserved[lecture] = details; applied.append(lecture)
        bounds = pub._boundaries(generated, lecture)
        body = generated._body._element
        for node in list(body)[bounds[0]:bounds[1]+1]: body.remove(node)
        for offset, node in enumerate(merged): body.insert(bounds[0]+offset, deepcopy(node))
    if not applied:
        return {'status': 'PERSONAL_CONTENT_CONFLICT', 'lectures': lectures, 'conflicts': conflicts, 'output': str(target), 'published_to_word': False}
    output = directory / 'course.candidate.docx'; generated.save(output)
    structure = {}
    for lecture in applied:
        payload, _ = payloads[lecture]; notes = payload['notes']
        visuals = sum(v.get('decision') == 'include' for v in notes.get('visual_decisions', []))
        formula = bool(notes.get('formula') or notes.get('formulas') or any(b.get('kind') == 'formula' for b in notes.get('blocks', [])))
        tables = [b['rows'] for b in notes.get('blocks', []) if b.get('kind') == 'table'] or ([notes['table']] if notes.get('table') else [])
        structure[lecture] = pub._docx_structure(output, lecture, visuals, formula, tables)
    render = pub._render_and_validate(output, directory/'render')
    if (file_sha256(target) if target.is_file() else None) != baseline_hash or any(file_sha256(workspace.runs_dir/l/'lecture_notes.candidate.json') != digest for l, digest in sources.items()):
        raise pub.PublicationError('Word or notes changed during the cumulative build; no file was overwritten')
    validation = {'lecture_id': lectures[0], 'lecture_ids': applied, 'requested_lectures': lectures, 'mode': 'course_update', 'target': str(target),
        'candidate': str(output), 'candidate_hash': file_sha256(output), 'baseline_hash': baseline_hash,
        'source_candidate_hashes': sources, 'generated_baseline': str(raw), 'generated_baseline_sha256': file_sha256(raw),
        'structure': structure, 'render': render, 'conflicts': conflicts, 'preserved_edits': preserved,
        'content_hashes': {l: pub._content_fingerprint(payloads[l][0]) for l in applied}, 'validated_at': now_compact_utc()}
    atomic_json_write(output.with_suffix('.validation.json'), validation)
    return {'status': 'WORD_CANDIDATE_VALIDATED', **validation}


def publish_course_revision(workspace, candidate_path):
    from . import publication as pub
    validation = json.loads(candidate_path.with_suffix('.validation.json').read_text())
    lectures = validation['lecture_ids']; code = parse_lecture_id(lectures[0])[0]
    course = pub._course(workspace.configuration_store().load(), code); target = pub._target(course, workspace)
    if validation['mode'] != 'course_update' or Path(validation['target']).resolve() != target.resolve():
        raise pub.PublicationError('Cumulative candidate target changed')
    with ExitStack() as locks:
        for lecture in validation['requested_lectures']: locks.enter_context(_generation_lock(workspace, lecture))
        locks.enter_context(pub._lock(target))
        if file_sha256(candidate_path) != validation['candidate_hash'] or file_sha256(Path(validation['generated_baseline'])) != validation['generated_baseline_sha256']:
            raise pub.PublicationError('Validated cumulative candidate or generated baseline changed')
        if (file_sha256(target) if target.is_file() else None) != validation['baseline_hash']:
            raise pub.PublicationError('Cumulative Word changed after validation; rebuild from the current file')
        payloads = {}
        for lecture, digest in validation['source_candidate_hashes'].items():
            if file_sha256(workspace.runs_dir/lecture/'lecture_notes.candidate.json') != digest:
                raise pub.PublicationError('Notes changed after cumulative validation')
            payloads[lecture] = pub._load_candidate(workspace, lecture)
        transaction = {'transaction_id': uuid.uuid4().hex, 'lecture_id': lectures[0], 'lecture_ids': lectures,
            'target': str(target), 'candidate': str(candidate_path), 'candidate_hash': validation['candidate_hash'],
            'baseline_hash': validation['baseline_hash'], 'status': 'PREPARED', 'created_at': now_compact_utc(),
            'content_hashes': validation['content_hashes'], 'content_hash': validation['content_hashes'][lectures[0]],
            'publication_state': pub._publication_state(workspace, target)}
        originals_dir = candidate_path.parent / 'original-candidates'
        originals_dir.mkdir(exist_ok=True)
        transaction['original_candidates'] = {}
        for lecture in lectures:
            original = workspace.runs_dir/lecture/'lecture_notes.candidate.json'
            saved = originals_dir/(lecture+'.json'); saved.write_bytes(original.read_bytes())
            transaction['original_candidates'][lecture] = {'path': str(saved), 'sha256': file_sha256(saved)}
        journal = pub._journal(workspace, transaction)
        transaction['backup'] = pub._archive(workspace, target, code, reason='before coordinated update ' + ', '.join(lectures))
        transaction['status'] = 'BACKED_UP'; pub._journal(workspace, transaction)
        target.parent.mkdir(parents=True, exist_ok=True)
        fd, name = tempfile.mkstemp(suffix='.docx', dir=target.parent); os.close(fd)
        temporary = Path(name)
        try:
            temporary.write_bytes(candidate_path.read_bytes())
            if file_sha256(temporary) != validation['candidate_hash'] or (file_sha256(target) if target.is_file() else None) != validation['baseline_hash']:
                raise pub.PublicationError('Cumulative Word changed during publication')
            os.replace(temporary, target)
        finally: temporary.unlink(missing_ok=True)
        transaction['status'] = 'PUBLISHED'; transaction['target_hash'] = file_sha256(target); pub._journal(workspace, transaction)
        originals = {lecture: (workspace.runs_dir/lecture/'lecture_notes.candidate.json').read_bytes() for lecture in lectures}
        try:
            index = baseline_index(workspace, code)
            with sqlite3.connect(workspace.database_path) as connection:
                for lecture in lectures:
                    payload, pack = payloads[lecture]
                    payload.update(published_to_word=True, published_path=str(target), published_hash=transaction['target_hash'])
                    path = workspace.runs_dir/lecture/'lecture_notes.candidate.json'; atomic_json_write(path, payload)
                    requested = payload.get('requested_model', {}); actual = payload.get('actual_model', {})
                    row = connection.execute('SELECT input_hash FROM candidate_runs WHERE lecture_id=?', (lecture,)).fetchone()
                    values = (lecture, row[0] if row else payload.get('source_pack_hash', ''), 'PUBLISHED', actual.get('provider', requested.get('provider', 'unknown')), actual.get('model', requested.get('model', 'unknown')), actual.get('effort', requested.get('effort', 'unknown')), payload.get('quality', 'standard'), str(target), transaction['target_hash'], json.dumps(pack.get('availability', {}), sort_keys=True), now_compact_utc())
                    connection.execute('INSERT OR REPLACE INTO lesson_runs VALUES(?,?,?,?,?,?,?,?,?,?,?)', values)
                    connection.execute("UPDATE candidate_runs SET status='PUBLISHED',candidate_hash=?,updated_at=? WHERE lecture_id=?", (file_sha256(path), now_compact_utc(), lecture))
                    index[lecture] = {'path': validation['generated_baseline'], 'sha256': validation['generated_baseline_sha256']}
            atomic_json_write(_index_path(workspace, code), index)
            state = WorkflowState(workspace.database_path)
            try:
                from .workflow import _write_status_markdown
                _write_status_markdown(workspace, state.rows())
            finally: state.close()
            transaction['status'] = 'COMPLETED'; transaction['completed_at'] = now_compact_utc(); pub._journal(workspace, transaction)
        except Exception:
            backup = transaction.get('backup')
            if backup:
                saved = Path(backup['path'])
                if file_sha256(saved) != backup['sha256']:
                    transaction['status'] = 'RECOVERY_REQUIRED'; pub._journal(workspace, transaction); raise
                fd, name = tempfile.mkstemp(suffix='.docx', dir=target.parent); os.close(fd)
                temporary = Path(name)
                try:
                    temporary.write_bytes(saved.read_bytes()); os.replace(temporary, target)
                finally: temporary.unlink(missing_ok=True)
            else:
                target.unlink(missing_ok=True)
            for lecture, content in originals.items():
                (workspace.runs_dir/lecture/'lecture_notes.candidate.json').write_bytes(content)
            with sqlite3.connect(workspace.database_path) as connection:
                for lecture in originals:
                    path = workspace.runs_dir/lecture/'lecture_notes.candidate.json'
                    connection.execute('UPDATE candidate_runs SET candidate_hash=? WHERE lecture_id=?', (file_sha256(path), lecture))
            pub._restore_publication_state(workspace, target, transaction['publication_state'])
            transaction['status'] = 'ROLLED_BACK'; pub._journal(workspace, transaction)
            raise
    return {'status': 'PARTIALLY_PUBLISHED' if validation['conflicts'] else 'PUBLISHED', 'lecture_id': lectures[0], 'lectures': lectures,
        'output': str(target), 'backup': transaction['backup'], 'transaction': str(journal),
        'conflicts': validation['conflicts'], 'preserved_edits': validation['preserved_edits'], 'validation': validation}


def recover_course_revision(workspace, transaction):
    """Rollback an interrupted multi-lesson commit, preserving its staged output."""
    from . import publication as pub
    target = Path(transaction['target'])
    with pub._restore_locks(workspace, target, transaction['lecture_ids']):
        current = file_sha256(target) if target.is_file() else None
        if current not in {transaction.get('baseline_hash'), transaction.get('candidate_hash')}:
            raise pub.PublicationError('Word changed after an interrupted update; current file preserved')
        originals = {}
        for lecture, record in transaction['original_candidates'].items():
            saved = Path(record['path']); actual = workspace.runs_dir/lecture/'lecture_notes.candidate.json'
            if not saved.resolve().is_relative_to(workspace.root.resolve()) or file_sha256(saved) != record['sha256']:
                raise pub.PublicationError('Interrupted update original candidate is missing or changed')
            original = json.loads(saved.read_text()); present = json.loads(actual.read_text())
            if pub._content_fingerprint(original) != pub._content_fingerprint(present):
                raise pub.PublicationError('Notes changed after an interrupted update; current files preserved')
            originals[actual] = saved.read_bytes()
        backup = transaction.get('backup')
        if backup:
            saved = Path(backup['path'])
            if file_sha256(saved) != backup['sha256']:
                raise pub.PublicationError('Interrupted update backup is missing or changed')
            fd, name = tempfile.mkstemp(suffix='.docx', dir=target.parent); os.close(fd)
            temporary = Path(name)
            try:
                temporary.write_bytes(saved.read_bytes()); os.replace(temporary, target)
            finally: temporary.unlink(missing_ok=True)
        elif transaction.get('baseline_hash') is None:
            target.unlink(missing_ok=True)
        for path, content in originals.items(): path.write_bytes(content)
        with sqlite3.connect(workspace.database_path) as connection:
            for path in originals:
                connection.execute('UPDATE candidate_runs SET candidate_hash=? WHERE lecture_id=?', (file_sha256(path), path.parent.name))
        pub._restore_publication_state(workspace, target, transaction['publication_state'])
        transaction['status'] = 'ROLLED_BACK'; transaction['recovered_at'] = now_compact_utc(); pub._journal(workspace, transaction)
    return transaction


def update_course(workspace, lectures):
    from . import publication as pub
    if isinstance(lectures, str): lectures = [lectures]
    lectures = [value.upper() for value in lectures]
    code = parse_lecture_id(lectures[0])[0]
    course = pub._course(workspace.configuration_store().load(), code); target = pub._target(course, workspace)
    transactions = pub._transactions(workspace, target)
    state = WorkflowState(workspace.database_path)
    try:
        expected = state.expected_output_hash(target)
        published = {r[0] for r in state.connection.execute("SELECT lecture_id FROM lesson_runs WHERE status='PUBLISHED'")}
    finally: state.close()
    if target.is_file() and expected == file_sha256(target) and all(l in published for l in lectures):
        unchanged = True
        for lecture in lectures:
            payload, _ = pub._load_candidate(workspace, lecture)
            previous = next((t.get('content_hashes', {t.get('lecture_id'): t.get('content_hash')}).get(lecture)
                             for t in transactions if t.get('status') == 'COMPLETED' and lecture in t.get('content_hashes', {t.get('lecture_id'): t.get('content_hash')})), None)
            unchanged &= bool(payload.get('published_to_word')) and previous == pub._content_fingerprint(payload)
        if unchanged:
            return {'status': 'SKIPPED', 'lecture_id': lectures[0], 'lectures': lectures, 'output': str(target), 'reason': 'already_published'}
    built = build_course_revision(workspace, lectures)
    if built['status'] == 'PERSONAL_CONTENT_CONFLICT': return built
    if course['publication']['mode'] == 'review':
        return {**built, 'status': 'AWAITING_APPROVAL', 'next_action': 'Show the cumulative candidate and publish only after approval.'}
    return publish_course_revision(workspace, Path(built['candidate']))
