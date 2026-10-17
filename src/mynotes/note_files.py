"""Stable course notebook and handout paths, independent of generated drafts."""
from __future__ import annotations

import os
import tempfile
from pathlib import Path

from docx import Document

from .configuration import ConfigurationError, clean_path

NOTE_FORMATS = {'.docx', '.txt', '.md'}


def handout_path(course, workspace):
    explicit = course.get('files', {}).get('handout')
    if explicit:
        return Path(explicit)
    folder = course['folders'].get('handout')
    return (Path(folder) if folder else workspace.runs_dir / 'handouts') / f"{course['code']}_COURSE_NOTES.docx"


def file_path(value, *, handout=False):
    path = Path(clean_path(str(value))).resolve()
    allowed = {'.docx'} if handout else NOTE_FORMATS
    if path.suffix.casefold() not in allowed or path.is_dir():
        raise ConfigurationError('Choose a .docx handout or a .docx, .txt or .md personal notebook.')
    return path


def create_notebook(path, course):
    """Create only a new Word; an existing user notebook is never rewritten."""
    if path.exists():
        return False
    if path.suffix.casefold() != '.docx':
        raise ConfigurationError('A new personal notebook must be a .docx file; select an existing text file instead.')
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(course['name'], level=1)
    document.add_heading('Lezione 1', level=2)
    document.add_paragraph('')
    fd, name = tempfile.mkstemp(suffix='.docx', dir=path.parent)
    os.close(fd)
    temporary = Path(name)
    try:
        document.save(temporary)
        try:
            os.link(temporary, path)  # Exclusive creation: never replace a concurrent user's file.
        except FileExistsError:
            return False
    finally:
        temporary.unlink(missing_ok=True)
    return True


def files_guidance(course, language):
    notes, handout = (course['files'][key] for key in ('personal_notes', 'handout'))
    automatic = course['publication']['mode'] == 'automatic'
    if language == 'it':
        action = 'Aggiornerò' if automatic else 'Dopo la tua approvazione aggiornerò'
        return (f'Scrivi i tuoi appunti in {notes}, separandoli con titoli come «Lezione 1», «Lezione 2». '
                'Puoi aggiungere «NOTA AI: chiarisci questo passaggio» per richiamare l’attenzione durante l’elaborazione. '
                f'{action} {handout} salvando un backup; puoi cambiare entrambi i nomi o chiedermi di annullare l’ultimo aggiornamento.')
    action = 'I will update' if automatic else 'After your approval, I will update'
    return (f'Write your notes in {notes}, using headings such as “Lesson 1”, “Lesson 2”. '
            'Add “NOTA AI: clarify this step” to highlight a point during preparation. '
            f'{action} {handout}, saving a backup; you can change either filename or ask me to undo the last update.')
