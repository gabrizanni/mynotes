"""Transactional Word candidate building, validation, publication and restore."""

from __future__ import annotations

import json
import hashlib
from io import BytesIO
from importlib.resources import files
import os
import re
import shutil
import time
import sqlite3
import tempfile
import uuid
import zipfile
from contextlib import contextmanager, ExitStack
from copy import deepcopy
from pathlib import Path
from typing import Any, Iterator

import fcntl
import pypdfium2 as pdfium
import pypdfium2.raw as pdfium_raw
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from .configuration import effective_style
from .office import OfficeError, convert_to_pdf
from .utils import atomic_json_write, file_sha256, now_compact_utc
from .workspace import Workspace


class PublicationError(RuntimeError):
    """Publication stopped before damaging a cumulative handout."""


def word_renderer_hash() -> str:
    template = Path(str(files('mynotes').joinpath('resources/base-template.docx')))
    return hashlib.sha256((file_sha256(Path(__file__)) + file_sha256(template)).encode()).hexdigest()


_GREEK = {
    "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε",
    "theta": "θ", "lambda": "λ", "mu": "μ", "pi": "π", "rho": "ρ",
    "sigma": "σ", "tau": "τ", "phi": "φ", "omega": "ω", "Delta": "Δ",
    "Sigma": "Σ", "Pi": "Π", "Phi": "Φ", "Omega": "Ω",
}


def _m(tag: str, text: str | None = None) -> OxmlElement:
    element = OxmlElement(f"m:{tag}")
    if text is not None:
        element.text = text
    return element


def _math_run(text: str, *, normal: bool = False) -> OxmlElement:
    run = _m("r")
    if normal:
        properties = _m("rPr")
        normal_text = _m("nor")
        normal_text.set(qn("m:val"), "1")
        properties.append(normal_text)
        run.append(properties)
    run.append(_m("t", text))
    return run


class _FormulaParser:
    """Small strict LaTeX parser for the equation forms used in lecture notes."""

    def __init__(self, source: str) -> None:
        # Providers sometimes double-escape LaTeX inside a JSON text field.
        # A single backslash is the actual TeX command introducer.
        value = source.strip().replace("\\\\", "\\")
        for left, right in (("$$", "$$"), ("\\[", "\\]"), ("$", "$")):
            if value.startswith(left) and value.endswith(right):
                value = value[len(left):-len(right)].strip()
                break
        self.source = value
        self.position = 0

    def parse(self) -> list[OxmlElement]:
        result = self._sequence(None)
        if self.position != len(self.source):
            raise PublicationError(f"Unsupported formula near: {self.source[self.position:]}")
        if not result:
            raise PublicationError("Formula is empty")
        return result

    def _sequence(self, stop: str | None) -> list[OxmlElement]:
        result: list[OxmlElement] = []
        plain = ""
        while self.position < len(self.source):
            char = self.source[self.position]
            if stop and char == stop:
                break
            if char == "\\":
                if plain:
                    result.append(_math_run(plain)); plain = ""
                result.extend(self._command())
                continue
            if char in "^_":
                if plain:
                    result.append(_math_run(plain)); plain = ""
                if not result:
                    raise PublicationError("A script has no base")
                base = result.pop()
                self.position += 1
                scripts: dict[str, list[OxmlElement]] = {char: self._argument()}
                if self.position < len(self.source) and self.source[self.position] in "^_" and self.source[self.position] != char:
                    other = self.source[self.position]; self.position += 1
                    scripts[other] = self._argument()
                tag = "sSubSup" if len(scripts) == 2 else ("sSup" if "^" in scripts else "sSub")
                node = _m(tag); node.append(_m("e")); node[-1].append(base)
                if "_" in scripts:
                    sub = _m("sub"); [sub.append(item) for item in scripts["_"]]; node.append(sub)
                if "^" in scripts:
                    sup = _m("sup"); [sup.append(item) for item in scripts["^"]]; node.append(sup)
                result.append(node)
                continue
            if char == "{":
                if plain:
                    result.append(_math_run(plain)); plain = ""
                result.extend(self._argument())
                continue
            if char == "}":
                break
            plain += char
            self.position += 1
        if plain:
            result.append(_math_run(plain))
        return result

    def _argument(self) -> list[OxmlElement]:
        while self.position < len(self.source) and self.source[self.position].isspace():
            self.position += 1
        if self.position >= len(self.source):
            raise PublicationError("Formula argument is missing")
        if self.source[self.position] == "{":
            self.position += 1
            value = self._sequence("}")
            if self.position >= len(self.source) or self.source[self.position] != "}":
                raise PublicationError("Unclosed formula group")
            self.position += 1
            return value
        if self.source[self.position] == "\\":
            return self._command()
        char = self.source[self.position]; self.position += 1
        return [_math_run(char)]

    def _literal_argument(self) -> str:
        """Keep words inside LaTeX text commands literal in Word equations."""
        while self.position < len(self.source) and self.source[self.position].isspace():
            self.position += 1
        if self.position >= len(self.source) or self.source[self.position] != "{":
            raise PublicationError("Formula text requires a braced argument")
        self.position += 1
        start = self.position
        depth = 1
        while self.position < len(self.source) and depth:
            if self.source[self.position] == "{":
                depth += 1
            elif self.source[self.position] == "}":
                depth -= 1
            self.position += 1
        if depth:
            raise PublicationError("Unclosed formula text")
        return self.source[start:self.position - 1]

    def _command(self) -> list[OxmlElement]:
        self.position += 1
        match = re.match(r"[A-Za-z]+", self.source[self.position:])
        if not match:
            if self.position >= len(self.source):
                raise PublicationError("Dangling formula escape")
            char = self.source[self.position]; self.position += 1
            return [_math_run(char)]
        name = match.group(0); self.position += len(name)
        if name in _GREEK:
            return [_math_run(_GREEK[name])]
        symbols = {
            "cdot": "·", "times": "×", "le": "≤", "leq": "≤", "ge": "≥", "geq": "≥",
            "neq": "≠", "ne": "≠", "infty": "∞", "pm": "±", "to": "→", "approx": "≈", "sim": "∼", "propto": "∝",
            "implies": "⇒", "quad": "  ", "dots": "…", "int": "∫",
            "max": "max", "min": "min",
            "ln": "ln", "log": "log", "exp": "exp",
            "ell": "ℓ",
            "in": "∈",
        }
        if name in symbols:
            return [_math_run(symbols[name])]
        if name in {"left", "right"}:
            # Delimiter size is layout only; retain the following delimiter as
            # editable math, except TeX's explicit invisible delimiter.
            if self.source[self.position:self.position + 1] == ".":
                self.position += 1
            return []
        if name == "hat":
            accent = _m("acc"); properties = _m("accPr")
            character = _m("chr"); character.set(qn("m:val"), "̂")
            properties.append(character); accent.append(properties)
            expression = _m("e")
            for item in self._argument(): expression.append(item)
            accent.append(expression)
            return [accent]
        if name == "mathcal":
            items = self._argument()
            for item in items:
                for run in item.iter(qn("m:r")):
                    properties = run.find(qn("m:rPr"))
                    if properties is None:
                        properties = _m("rPr"); run.insert(0, properties)
                    script = _m("scr"); script.set(qn("m:val"), "script")
                    properties.append(script)
            return items
        if name in {"frac", "binom"}:
            numerator, denominator = self._argument(), self._argument()
            fraction = _m("f"); num, den = _m("num"), _m("den")
            if name == "binom":
                properties = _m("fPr"); kind = _m("type"); kind.set(qn("m:val"), "noBar")
                properties.append(kind); fraction.append(properties)
            [num.append(item) for item in numerator]; [den.append(item) for item in denominator]
            fraction.append(num); fraction.append(den)
            if name == "binom":
                delimiter = _m("d"); properties = _m("dPr")
                for tag, value in (("begChr", "("), ("endChr", ")")):
                    character = _m(tag); character.set(qn("m:val"), value); properties.append(character)
                delimiter.append(properties); expression = _m("e"); expression.append(fraction); delimiter.append(expression)
                return [delimiter]
            return [fraction]
        if name == "sqrt":
            rad = _m("rad"); props = _m("radPr"); hide = _m("degHide"); hide.set(qn("m:val"), "1")
            props.append(hide); rad.append(props); rad.append(_m("deg")); expression = _m("e")
            [expression.append(item) for item in self._argument()]; rad.append(expression)
            return [rad]
        if name in {"sum", "prod"}:
            # Word displays empty placeholders for an nary element without a
            # full expression. A symbol with native scripts stays editable and
            # renders correctly even when the model omits summation limits.
            return [_math_run("∑" if name == "sum" else "∏")]
        if name == "text":
            return [_math_run(self._literal_argument(), normal=True)]
        if name == "mathbb":
            items = self._argument()
            for item in items:
                for run in item.iter(qn("m:r")):
                    properties = run.find(qn("m:rPr"))
                    if properties is None:
                        properties = _m("rPr")
                        run.insert(0, properties)
                    script = _m("scr"); script.set(qn("m:val"), "double-struck")
                    style = _m("sty"); style.set(qn("m:val"), "p")
                    properties.append(script); properties.append(style)
                    # E/P script styling is ignored by some renderers; Unicode
                    # math letters keep the same notation and remain editable.
                    exceptions = {"C": "ℂ", "H": "ℍ", "N": "ℕ", "P": "ℙ", "Q": "ℚ", "R": "ℝ", "Z": "ℤ"}
                    text = run.find(qn("m:t"))
                    if text is not None and text.text:
                        text.text = "".join(
                            exceptions.get(char, chr(0x1D538 + ord(char) - ord("A"))) if "A" <= char <= "Z"
                            else chr(0x1D552 + ord(char) - ord("a")) if "a" <= char <= "z"
                            else chr(0x1D7D8 + ord(char) - ord("0")) if "0" <= char <= "9"
                            else char for char in text.text
                        )
            return items
        if name in {"mathrm", "mathbf"}:
            return self._argument()
        raise PublicationError(f"Unsupported formula command: \\{name}")


def append_equation(document: Document, formula: str) -> None:
    paragraph = document.add_paragraph(style="MyNotes Body" if "MyNotes Body" in document.styles else None)
    math_para = _m("oMathPara"); math = _m("oMath")
    aligned = re.fullmatch(r"\\begin\{(aligned|gathered)\}(.*?)\\end\{\1\}", formula.strip(), flags=re.DOTALL)
    if aligned:
        # Preserve multiline provider equations as editable Word equations.
        # TeX alignment markers control placement, not mathematical content.
        array = _m("eqArr")
        for row in aligned.group(2).split("\\\\"):
            expression = _m("e")
            for item in _FormulaParser(row.replace("&", "")).parse():
                expression.append(item)
            array.append(expression)
        math.append(array)
    else:
        for item in _FormulaParser(formula).parse():
            math.append(item)
    math_para.append(math); paragraph._p.append(math_para)


def _marker(document: Document, value: str, *, page_break_before: bool = False) -> None:
    paragraph = document.add_paragraph(style="MyNotes Body" if "MyNotes Body" in document.styles else None); run = paragraph.add_run(value)
    run.font.hidden = True; run.font.size = Pt(1)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(1)
    paragraph.paragraph_format.page_break_before = page_break_before


def _boundaries(document: Document, lecture_id: str) -> tuple[int, int] | None:
    children = list(document._body._element)
    starts = [index for index, item in enumerate(children) if f"MYNOTES START {lecture_id}" in "".join(item.itertext())]
    ends = [index for index, item in enumerate(children) if f"MYNOTES END {lecture_id}" in "".join(item.itertext())]
    if not starts and not ends:
        return None
    if len(starts) != 1 or len(ends) != 1 or starts[0] >= ends[0]:
        raise PublicationError("Existing lecture boundaries are missing, duplicated or ambiguous")
    return starts[0], ends[0]


def _append_visual(document: Document, visual: dict[str, Any], decision: dict[str, Any]) -> None:
    image = Path(str(visual.get("image_path", ""))).expanduser()
    if not image.is_file():
        raise PublicationError(f"Selected visual is missing: {image}")
    paragraph = document.add_paragraph(style="MyNotes Body"); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image), width=Inches(5.4))
    caption = document.add_paragraph(style="MyNotes Caption")
    caption.paragraph_format.keep_with_next = False
    page = f", page {visual['page']}" if visual.get("page") is not None else ""
    source_name = Path(str(visual.get("source_file") or image.name)).name
    caption.add_run(f"Figure — {decision.get('reason', 'Relevant course visual')} ({source_name}{page})")


def _append_rows(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    if (not isinstance(rows, list) or any(not isinstance(row, (list, tuple)) or not row
                                        or any(not isinstance(cell, str) for cell in row) for row in rows)):
        raise PublicationError("Word table needs non-empty rows of text")
    width = len(rows[0])
    if any(len(row) != width for row in rows):
        raise PublicationError("Word table rows must have the same number of cells; no missing cells were inferred")
    if any(_formatting_residue(cell, inline_supported=True) for row in rows for cell in row):
        raise PublicationError("Word table cell contains unresolved Markdown or escaped newlines")
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # A cumulative candidate can have several sections. Fit the table within
    # every section before its lesson block is moved into its final position.
    available = min(int(section.page_width - section.left_margin - section.right_margin)
                    for section in document.sections)
    if available <= 0:
        raise PublicationError("Word section has no usable width for a table")
    table.autofit = False
    column_width = available // width
    for column in table.columns:
        column.width = column_width
    table_width = table._tbl.tblPr.find(qn("w:tblW"))
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(round(available / 635)))
    repeat_header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
    for row_index, row in enumerate(rows):
        table.rows[row_index]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.width = column_width
            _append_inline(cell.paragraphs[0], str(value))
            for paragraph in cell.paragraphs:
                paragraph.style = document.styles["MyNotes Body"]
                for run in paragraph.runs:
                    if run.style.name == "Termine tecnico":
                        continue
                    if width >= 6:
                        run.font.size = Pt(8)
                    elif width >= 4:
                        run.font.size = Pt(9)
                    if row_index == 0 and run.style.name != "Termine tecnico":
                        run.bold = True


def _key_terms(blocks: list[dict[str, Any]]) -> tuple[str, ...]:
    """Use declared terms as a per-lecture vocabulary, never a hard-coded course list."""
    declared: set[str] = set()
    for block in blocks:
        for value in (block.get("text", ""), block.get("answer", ""), *block.get("items", [])):
            declared.update(match.group(1).strip() for match in re.finditer(r"\[\[term:([^\]]+)\]\]", str(value)))
    return tuple(sorted(declared, key=len, reverse=True))


def _inline_math_tokens(line: str) -> list[str]:
    """Find TeX delimiters without consuming currency signs or escaped dollars.

    Numeric currency can precede a real inline equation. Rejected openings are
    skipped one at a time so the following equation's opening stays available.
    """
    tokens: list[str] = []
    cursor = 0
    while cursor < len(line):
        opening = line.find("$", cursor)
        if opening < 0:
            break
        if opening and line[opening - 1] == "\\":
            cursor = opening + 1
            continue
        delimiter = "$$" if line.startswith("$$", opening) else "$"
        content = opening + len(delimiter)
        closing = content
        while True:
            closing = line.find(delimiter, closing)
            if closing < 0 or closing == 0 or line[closing - 1] != "\\":
                break
            closing += len(delimiter)
        if closing < 0:
            break
        body = line[content:closing].strip()
        numeric = bool(re.match(r"\d", body))
        next_character = line[closing + len(delimiter):closing + len(delimiter) + 1]
        mathematics = bool(re.search(r"\\(?:[A-Za-z]+|%)|[+*/=<>^_-]", body))
        simple = bool(re.fullmatch(r"[\w.,]+", body))
        incomplete = bool(re.search(r"[+*/=<>_-]\s*$", body))
        currency = numeric and (next_character.isalnum() or not (mathematics or simple))
        if delimiter == "$" and (not body or currency or incomplete or "\n" in body):
            cursor = opening + 1
            continue
        tokens.append(line[opening:closing + len(delimiter)])
        cursor = closing + len(delimiter)
    return tokens


def _append_inline(paragraph: Any, line: str, key_terms: tuple[str, ...] = ()) -> None:
    """Accent declared technical terms on recurring narrative uses."""
    term_pattern = (re.compile(r"(?<!\w)(?:" + "|".join(re.escape(term) for term in key_terms) + r")(?!\w)", re.IGNORECASE)
                    if key_terms else None)

    def append_plain(value: str) -> None:
        value = value.replace(r"\\$", "$" ).replace(r"\$", "$").replace(r'\"', '"')
        if not value:
            return
        # Recover only an unambiguous parenthesized bounded sum/product.
        # Code spans bypass append_plain; ordinary prose and variable names
        # must never acquire math semantics merely because they contain '_'.
        bare = re.search(r"\((?P<operator>sum|prod)(?P<body>_\{[^{}\n]+\}\^\{[^{}\n]+\}\s*[A-Za-z]\^\{[^{}\n]+\})\)", value)
        if bare:
            append_plain(value[:bare.start()] + "(")
            math = _m("oMath")
            for item in _FormulaParser("\\" + bare.group("operator") + bare.group("body")).parse():
                math.append(item)
            paragraph._p.append(math)
            append_plain(")" + value[bare.end():])
            return
        if term_pattern is None:
            paragraph.add_run(value)
            return
        cursor = 0
        for found in term_pattern.finditer(value):
            if found.start() > cursor:
                paragraph.add_run(value[cursor:found.start()])
            paragraph.add_run(found.group()).style = "Termine tecnico"
            cursor = found.end()
        if cursor < len(value):
            paragraph.add_run(value[cursor:])

    math_tokens = _inline_math_tokens(line)
    math_pattern = "|".join(re.escape(value) for value in math_tokens) or "(?!)"
    token = re.compile(
        "(?P<math>" + math_pattern + ")"
        r"|\[\[term:(?P<term>.+?)\]\]"
        r"|\*\*(?P<bold>.+?)\*\*"
        r"|(?<![\w*])\*(?!\*)(?P<italic>.+?)(?<!\*)\*(?![\w*$])"
        r"|`(?P<code>[^`]+)`"
    )
    cursor = 0
    for match in token.finditer(line):
        if match.start() > cursor:
            append_plain(line[cursor:match.start()])
        if match.group("math") is not None:
            math = _m("oMath")
            for item in _FormulaParser(match.group("math")).parse():
                math.append(item)
            paragraph._p.append(math)
            cursor = match.end()
            continue
        if match.group("bold") is not None and any(marker in match.group("bold") for marker in ("$", "[[term:", "`")):
            first = len(paragraph.runs)
            _append_inline(paragraph, match.group("bold"), key_terms)
            for nested in paragraph.runs[first:]:
                if nested.style.name != "Termine tecnico":
                    nested.bold = True
            cursor = match.end()
            continue
        if match.group("italic") is not None and any(marker in match.group("italic") for marker in ("$", "[[term:", "`")):
            first = len(paragraph.runs)
            _append_inline(paragraph, match.group("italic"), key_terms)
            for nested in paragraph.runs[first:]:
                if nested.style.name != "Termine tecnico":
                    nested.italic = True
            cursor = match.end()
            continue
        run = paragraph.add_run(next(match.group(name) for name in ("term", "bold", "italic", "code") if match.group(name) is not None))
        is_term = match.group("term") is not None or (match.group("bold") is not None and term_pattern is not None and term_pattern.fullmatch(run.text))
        if is_term:
            run.style = "Termine tecnico"
        elif match.group("bold") is not None:
            run.bold = True
        elif match.group("italic") is not None:
            run.italic = True
        cursor = match.end()
    if cursor < len(line):
        append_plain(line[cursor:])


def _formatting_residue(value: str, *, inline_supported: bool = False) -> bool:
    """Recognize formatting errors without mistaking TeX or paths for escapes."""
    text = value
    for token in _inline_math_tokens(text):
        text = text.replace(token, "")
    text = re.sub(r"(?<!`)`(?!`)[^`]+`(?!`)", "", text)
    # Windows/UNC paths are literal data, not serialized paragraph breaks.
    text = re.sub(r"(?:(?<!\w)[A-Za-z]:\\|\\\\|(?<!\w)\.\.?\\)[^\s]+", "", text)
    text = re.sub(r"\\n(?:u|eq|ot|abla)\b", "", text)
    if inline_supported:
        text = re.sub(r"\[\[term:([^\]]+)\]\]", r"\1", text)
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    return bool("\\n" in text or "**" in text or "[[term:" in text
                or re.search(r"\|\s*:?-{3,}:?\s*\|", text)
                or re.search(r"(?m)^\s*\|[^|]+\|", text)
                or "```" in text)


def _legacy_table_text(value: str) -> str:
    r"""Decode only serialized line boundaries adjacent to a pipe-table row.

    Do not globally decode backslashes: \\nu, \\neq and filesystem paths can
    be valid course content. Ambiguous malformed tables are rejected later.
    """
    protected = _inline_math_tokens(value)
    protected.extend(match.group() for match in re.finditer(
        r"(?<!`)`(?!`)[^`]+`(?!`)|(?:(?<!\w)[A-Za-z]:\\|\\\\|(?<!\w)\.\.?\\)[^\s]+", value))
    text = value
    for index, token in enumerate(protected):
        text = text.replace(token, f"\x00MYNOTES_LITERAL_{index}\x00")
    text = re.sub(r"(?:\\n)+(?=\s*\|)", lambda match: "\n" * (len(match.group()) // 2), text)
    text = re.sub(r"(?<=\|)(?:\\n)+", lambda match: "\n" * (len(match.group()) // 2), text)
    # Retain the common legacy sentence-break case without interpreting a TeX
    # command starting with n as a newline. Other ambiguous escapes fail QA.
    text = re.sub(r"(?<=[.!?;:])(?:\\n)+(?!u\b|eq\b|ot\b|abla\b)",
                  lambda match: "\n" * (len(match.group()) // 2), text)
    for index, token in enumerate(protected):
        text = text.replace(f"\x00MYNOTES_LITERAL_{index}\x00", token)
    return text


def _markdown_cells(line: str) -> list[str]:
    cells = re.split(r"(?<!\\)\|", line.strip()[1:-1])
    return [cell.strip().replace(r"\|", "|") for cell in cells]


def _append_text(document: Document, value: str, key_terms: tuple[str, ...] = ()) -> None:
    """Convert unambiguous legacy Markdown tables to native Word elements."""
    text = _legacy_table_text(str(value))
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if (line.startswith("$$") and line.endswith("$$") and len(line) > 4) or (line.startswith("\\[") and line.endswith("\\]") and len(line) > 4):
            append_equation(document, line)
            index += 1
            continue
        if line.startswith("|") and line.endswith("|"):
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|") and lines[index].strip().endswith("|"):
                rows.append(_markdown_cells(lines[index]))
                index += 1
            if (len(rows) < 2 or not all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1])
                    or any(len(row) != len(rows[0]) for row in rows)):
                raise PublicationError("Malformed Markdown table: expected a rectangular header, separator and rows")
            _append_rows(document, [rows[0], *rows[2:]])
            continue
        if _formatting_residue(line, inline_supported=True):
            raise PublicationError("Paragraph contains unresolved Markdown table or escaped newlines")
        if line.startswith(("- ", "* ")):
            paragraph = document.add_paragraph(style="MyNotes List Bullet")
            paragraph._p.get_or_add_pPr().append(deepcopy(document.styles["List Bullet"]._element.pPr.numPr))
            line = line[2:]
        else:
            paragraph = document.add_paragraph(style="MyNotes Body")
        # The model sometimes returns Markdown emphasis inside JSON strings.
        if line.startswith("*") and line.endswith("*") and not line.startswith("**") and len(line) > 2:
            first_run = len(paragraph.runs)
            _append_inline(paragraph, line[1:-1], key_terms)
            for run in paragraph.runs[first_run:]:
                if run.style.name != "Termine tecnico":
                    run.italic = True
        else:
            _append_inline(paragraph, line, key_terms)
        index += 1


def _heading(document: Document, text: str, level: int) -> None:
    if level not in {2, 3, 4}:
        raise PublicationError(f"Unsupported heading level: {level}")
    document.add_paragraph(text, style=f"Heading {level - 1}")


def _new_numbered_list(document: Document) -> int:
    """Give each lecture overview its own Word numbering sequence."""
    numbering = document.part.numbering_part.element
    style_num = document.styles["List Number"]._element.pPr.numPr.numId.val
    existing = [item for item in numbering if item.tag == qn("w:num")]
    source = next((item for item in existing if int(item.get(qn("w:numId"))) == style_num), None)
    if source is None:
        raise PublicationError("Word numbering definition is missing")
    source_abstract_id = source.find(qn("w:abstractNumId")).get(qn("w:val"))
    abstract_definition = deepcopy(next(item for item in numbering if item.tag == qn('w:abstractNum')
                                       and item.get(qn('w:abstractNumId')) == source_abstract_id))
    abstract_id = str(max((int(item.get(qn('w:abstractNumId'))) for item in numbering if item.tag == qn('w:abstractNum')),default=-1)+1)
    abstract_definition.set(qn('w:abstractNumId'),abstract_id)
    for tag in ('nsid','tmpl'):
        element=abstract_definition.find(qn('w:'+tag))
        if element is not None:element.set(qn('w:val'),uuid.uuid4().hex[:8].upper())
    for linked_style in list(abstract_definition.iter(qn('w:pStyle'))):
        linked_style.getparent().remove(linked_style)
    # Word requires every abstract definition before all concrete instances.
    numbering.insert(next((i for i,item in enumerate(numbering) if item.tag==qn('w:num')),len(numbering)),abstract_definition)
    number_id = max((int(item.get(qn("w:numId"))) for item in existing), default=0) + 1
    number = OxmlElement("w:num"); number.set(qn("w:numId"), str(number_id))
    abstract = OxmlElement("w:abstractNumId"); abstract.set(qn("w:val"), abstract_id)
    override = OxmlElement("w:lvlOverride"); override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride"); start.set(qn("w:val"), "1")
    override.append(start); number.append(abstract); number.append(override); numbering.append(number)
    return number_id


def _set_numbered_item(paragraph: Any, number_id: int) -> None:
    properties = paragraph._p.get_or_add_pPr()
    numbering = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl"); level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId"); number.set(qn("w:val"), str(number_id))
    numbering.append(level); numbering.append(number); properties.append(numbering)


def _lesson_heading(lecture_id: str, title: str) -> str:
    match = re.search(r"-L0*(\d+)$", lecture_id, flags=re.IGNORECASE)
    if not match:
        return title.strip()
    lesson_number = int(match.group(1))
    clean_title = re.sub(r"^\s*lecture\s+\d+\s*[:\-–—]\s*", "", title, flags=re.IGNORECASE).strip()
    return f"Lecture {lesson_number}: {clean_title}" if clean_title else f"Lecture {lesson_number}"


def append_note_blocks(document: Document, notes: dict[str, Any], pack: dict[str, Any], *, language: str = "en", recap_language: str = "it") -> None:
    """Shared native Word narrative renderer for lectures and standalone study guides."""
    italian = language.lower().startswith("it")
    visuals = {str(item.get("evidence_id")): item for item in pack.get("visuals", [])}
    decisions = {str(item.get("evidence_id")): item for item in notes.get("visual_decisions", [])}
    blocks = notes.get("blocks") or []
    if blocks:
        key_terms = _key_terms(blocks)
        qa_heading_added = False
        for block in blocks:
            kind = block["kind"]
            if kind == "overview":
                _heading(document, "Argomenti della lezione" if italian else "Lecture overview", 3)
                number_id = _new_numbered_list(document)
                for item in block["items"]:
                    paragraph = document.add_paragraph(style="MyNotes List Number")
                    _set_numbered_item(paragraph, number_id)
                    _append_inline(paragraph, str(item))
                    for run in paragraph.runs:
                        if run.style.name != "Termine tecnico":
                            run.bold = True
            elif kind == "heading":
                _heading(document, block["text"], int(block["level"]))
            elif kind in {"paragraph", "recap"}:
                if kind == "recap":
                    recap_heading = "Riepilogo in italiano" if recap_language == "it" else ("Riepilogo" if italian else "Recap")
                    _heading(document, recap_heading, 3)
                first_paragraph = len(document.paragraphs)
                _append_text(document, block["text"], key_terms)
                if kind == "recap":
                    recap_paragraphs = document.paragraphs[first_paragraph:]
                    for paragraph in recap_paragraphs[:-1]:
                        paragraph.paragraph_format.keep_with_next = True
                    if recap_paragraphs:
                        recap_paragraphs[-1].paragraph_format.keep_together = True
            elif kind == "bullets":
                for item in block["items"]:
                    _append_text(document, f"- {item}", key_terms)
            elif kind == "qa":
                if not qa_heading_added:
                    _heading(document, "Risposte alle mie domande" if italian else "Answers to my questions", 3)
                    qa_heading_added = True
                question = document.add_paragraph(style="MyNotes Question")
                question.paragraph_format.keep_with_next = True
                question.add_run(str(block["question"]))
                _append_text(document, str(block["answer"]), key_terms)
            elif kind == "table":
                if block.get("label"):
                    _append_text(document, block["label"])
                    document.paragraphs[-1].paragraph_format.keep_with_next = True
                _append_rows(document, block["rows"])
            elif kind == "formula":
                if block.get("label"):
                    _append_text(document, block["label"])
                    document.paragraphs[-1].paragraph_format.keep_with_next = True
                append_equation(document, str(block["expression"]))
            elif kind == "visual":
                identifier = str(block["visual_id"])
                if identifier not in visuals or decisions.get(identifier, {}).get("decision") != "include":
                    raise PublicationError("Ordered visual is unavailable or not approved")
                _append_visual(document, visuals[identifier], decisions[identifier])


def _append_block(document: Document, notes: dict[str, Any], pack: dict[str, Any], *, page_break_before: bool = False, language: str = "en", recap_language: str = "it") -> None:
    lecture_id = notes["lecture_id"]
    _marker(document, f"MYNOTES START {lecture_id}", page_break_before=page_break_before)
    _heading(document, _lesson_heading(lecture_id, str(notes["title"])), 2)
    italian = language.lower().startswith("it")
    visuals = {str(item.get("evidence_id")): item for item in pack.get("visuals", [])}
    decisions = {str(item.get("evidence_id")): item for item in notes.get("visual_decisions", [])}
    if notes.get("blocks"):
        append_note_blocks(document, notes, pack, language=language, recap_language=recap_language)
    else:
        if notes.get("summary"):
            _append_text(document, notes["summary"])
        if notes.get("bullets"):
            _heading(document, "Key takeaways", 3)
            for bullet in notes["bullets"]:
                _append_text(document, f"- {bullet}")
        for section in notes.get("sections", []):
            _heading(document, section["heading"], int(section["level"]))
            if section.get("body"):
                _append_text(document, section["body"])
            for bullet in section.get("bullets", []):
                _append_text(document, f"- {bullet}")
        _append_rows(document, notes.get("table", []))
        formulas = list(notes.get("formulas", []))
        if notes.get("formula") and all(item.get("expression") != notes["formula"] for item in formulas):
            formulas.insert(0, {"label": "Formula", "expression": notes["formula"]})
        if formulas:
            _heading(document, "Formulas", 3)
            for item in formulas:
                if item.get("label"):
                    _append_text(document, str(item["label"]))
                append_equation(document, str(item["expression"]))
        for identifier, decision in decisions.items():
            if decision.get("decision") == "include":
                visual = visuals.get(identifier)
                if not visual:
                    raise PublicationError("Selected visual does not exist in the source pack")
                _append_visual(document, visual, decision)
    _marker(document, f"MYNOTES END {lecture_id}")


def _profile_color(value: str, fallback: str) -> RGBColor:
    try:
        return RGBColor.from_string(value.lstrip("#"))
    except (TypeError, ValueError):
        return RGBColor.from_string(fallback)


def _ensure_template_styles(document: Document) -> None:
    """Import absent base styles/numbering into a candidate, never its original.

    Existing cumulative documents may come from older/localized Word exports
    without List Number/List Bullet. Remap imported numbering IDs so existing
    paragraphs retain their own definitions.
    """
    template = Document(str(files("mynotes").joinpath("resources/base-template.docx")))
    style_ids = {s.style_id for s in document.styles}
    mapping = {}; missing = []
    for style in template.styles:
        if style.name in document.styles:
            mapping[style.style_id] = document.styles[style.name].style_id
        else:
            identifier = style.style_id
            while identifier in style_ids: identifier = "MyNotesBase" + identifier
            style_ids.add(identifier);mapping[style.style_id] = identifier
            missing.append(style)
    numbering = document.part.numbering_part.element
    source_numbering = template.part.numbering_part.element
    num_map = {}; abstract_map = {}
    def number(old_id):
        if old_id in num_map:return num_map[old_id]
        source = next(n for n in source_numbering if n.tag==qn('w:num') and n.get(qn('w:numId'))==old_id)
        abstract_id = source.find(qn('w:abstractNumId')).get(qn('w:val'))
        if abstract_id not in abstract_map:
            abstract = deepcopy(next(a for a in source_numbering if a.tag==qn('w:abstractNum') and a.get(qn('w:abstractNumId'))==abstract_id))
            new_abstract = str(max((int(a.get(qn('w:abstractNumId'))) for a in numbering if a.tag==qn('w:abstractNum')),default=-1)+1)
            abstract_map[abstract_id]=new_abstract;abstract.set(qn('w:abstractNumId'),new_abstract)
            for ref in abstract.iter(qn('w:pStyle')):
                ref.set(qn('w:val'),mapping.get(ref.get(qn('w:val')),ref.get(qn('w:val'))))
            numbering.insert(next((i for i,item in enumerate(numbering) if item.tag==qn('w:num')),len(numbering)),abstract)
        num = deepcopy(source);new_num=str(max((int(n.get(qn('w:numId'))) for n in numbering if n.tag==qn('w:num')),default=0)+1)
        num.set(qn('w:numId'),new_num);num.find(qn('w:abstractNumId')).set(qn('w:val'),abstract_map[abstract_id]);numbering.append(num);num_map[old_id]=new_num
        return new_num
    for style in missing:
        element=deepcopy(style.element);element.set(qn('w:styleId'),mapping[style.style_id]);element.attrib.pop(qn('w:default'),None)
        for tag in ('basedOn','next','link'):
            ref=element.find(qn('w:'+tag))
            if ref is not None:ref.set(qn('w:val'),mapping.get(ref.get(qn('w:val')),ref.get(qn('w:val'))))
        for ref in element.iter(qn('w:numId')):ref.set(qn('w:val'),number(ref.get(qn('w:val'))))
        document.styles.element.append(element)
    for name in ('List Number','List Bullet'):
        style=document.styles[name];ppr=style.element.get_or_add_pPr();numpr=ppr.find(qn('w:numPr'))
        identifier=numpr.find(qn('w:numId')) if numpr is not None else None
        if identifier is None or not any(n.tag==qn('w:num') and n.get(qn('w:numId'))==identifier.get(qn('w:val')) for n in numbering):
            source=deepcopy(template.styles[name].element.pPr.numPr)
            source.numId.val=int(number(str(source.numId.val)))
            if numpr is not None:ppr.remove(numpr)
            ppr.append(source)


def _configure_document(document: Document, profile: dict[str, Any], *, set_page_layout: bool = False) -> None:
    """Apply the selected study design and native Word heading hierarchy."""
    _ensure_template_styles(document)
    def named_style(name: str, base: str) -> Any:
        if name in document.styles:
            return document.styles[name]
        result = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        result.base_style = document.styles[base]
        return result

    def named_character_style(name: str) -> Any:
        if name in document.styles:
            return document.styles[name]
        return document.styles.add_style(name, WD_STYLE_TYPE.CHARACTER)

    def set_font_family(style: Any, family: str) -> None:
        style.font.name = family
        fonts = style._element.get_or_add_rPr().rFonts
        if fonts is not None:
            for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "csTheme", "cstheme"):
                fonts.attrib.pop(qn(f"w:{attribute}"), None)
            fonts.set(qn("w:eastAsia"), family)
            fonts.set(qn("w:cs"), family)

    body_font = str(profile.get("body_font", "Aptos"))
    heading_font = str(profile.get("heading_font", "Aptos Display"))
    body = named_style("MyNotes Body", "Normal")
    set_font_family(body, body_font)
    body.font.size = Pt(float(profile.get("body_pt", 11)))
    body.paragraph_format.space_after = Pt(float(profile.get("body_after_pt", 3)))
    body.paragraph_format.line_spacing = float(profile.get("body_line_spacing", 1.15))
    term = named_character_style("Termine tecnico")
    set_font_family(term, body_font)
    term.font.bold = True
    term.quick_style = True
    term.hidden = False
    term.unhide_when_used = True
    term.priority = 20
    term.font.color.rgb = _profile_color(str(profile.get("accent_color", "ED7D31")), "ED7D31")
    heading_color = _profile_color(str(profile.get("heading_color", "0F4761")), "0F4761")
    for level, point_size in ((2, "lesson_heading_pt"), (3, "section_heading_pt"), (4, "subsection_heading_pt")):
        style = document.styles[f"Heading {level - 1}"]
        set_font_family(style, heading_font)
        style.font.size = Pt(float(profile.get(point_size, {2: 16, 3: 14, 4: 12}[level])))
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = heading_color
        style.paragraph_format.space_before = Pt(float(profile.get("heading_before_pt", 8)))
        style.paragraph_format.space_after = Pt(float(profile.get("heading_after_pt", 4)))
        style.paragraph_format.keep_with_next = True
        properties = style._element.get_or_add_pPr()
        outline = properties.find(qn("w:outlineLvl"))
        if outline is None:
            outline = OxmlElement("w:outlineLvl")
            properties.append(outline)
        outline.set(qn("w:val"), str(level - 2))
        style.quick_style = True
        style.hidden = False
        style.priority = level - 1
    for name in ("List Number", "List Bullet"):
        style = named_style(f"MyNotes {name}", name)
        set_font_family(style, body_font)
        style.font.size = Pt(float(profile.get("body_pt", 11)))
        style.paragraph_format.space_after = Pt(3)
        if name == "List Bullet":
            # Unicode bullets avoid a private-use Symbol glyph disappearing in PDF.
            numbering = document.part.numbering_part.element
            num_id = document.styles[name]._element.pPr.numPr.numId.val
            number = next(item for item in numbering if item.tag == qn("w:num") and item.get(qn("w:numId")) == str(num_id))
            abstract_id = number.find(qn("w:abstractNumId")).get(qn("w:val"))
            abstract = next(item for item in numbering if item.tag == qn("w:abstractNum") and item.get(qn("w:abstractNumId")) == abstract_id)
            level = abstract.find(qn("w:lvl"))
            level.find(qn("w:lvlText")).set(qn("w:val"), "•")
            fonts = level.find(qn("w:rPr")).find(qn("w:rFonts"))
            fonts.set(qn("w:ascii"), body_font)
            fonts.set(qn("w:hAnsi"), body_font)
    caption = named_style("MyNotes Caption", "Caption")
    set_font_family(caption, body_font)
    caption.font.size = Pt(9)
    caption.font.bold = True
    caption.font.italic = True
    caption.font.color.rgb = _profile_color(str(profile.get("caption_color", "595959")), "595959")
    question = named_style("MyNotes Question", "Intense Quote")
    question.base_style = document.styles["Intense Quote"]
    set_font_family(question, body_font)
    question.font.size = Pt(float(profile.get("body_pt", 11)))
    question.font.color.rgb = _profile_color(str(profile.get("question_color", "4F81BD")), "4F81BD")
    question.font.bold = True
    question.font.italic = True
    question.paragraph_format.left_indent = Pt(46.8)
    question.paragraph_format.right_indent = Pt(46.8)
    question.paragraph_format.space_before = Pt(10)
    question.paragraph_format.space_after = Pt(14)
    question.paragraph_format.keep_with_next = True
    for section in document.sections:
        # New outputs use the product's fixed margins even when their baseline
        # comes from an older cumulative Word. This never edits that baseline.
        section.top_margin = section.bottom_margin = Cm(1.5)
        section.left_margin = section.right_margin = Cm(1.5)
        if set_page_layout:
            section.page_width = Cm(21)
            section.page_height = Cm(29.7)
        if profile.get("page_numbers", True):
            _ensure_page_number(section)
    if profile.get("page_numbers", True):
        settings = document.settings._element
        update = settings.find(qn("w:updateFields"))
        if update is None:
            update = OxmlElement("w:updateFields")
            settings.append(update)
        update.set(qn("w:val"), "true")


def _ensure_page_number(section: Any) -> None:
    footer = section.footer
    if "PAGE" in footer._element.xml:
        return
    paragraph = footer.paragraphs[0] if footer.paragraphs and not footer.paragraphs[0].text else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for kind in ("begin", "separate", "end"):
        run = OxmlElement("w:r")
        field = OxmlElement("w:fldChar")
        field.set(qn("w:fldCharType"), kind)
        run.append(field)
        paragraph._p.append(run)
        if kind == "begin":
            instruction = OxmlElement("w:r")
            text = OxmlElement("w:instrText")
            text.set(qn("xml:space"), "preserve")
            text.text = " PAGE "
            instruction.append(text)
            paragraph._p.append(instruction)
        elif kind == "separate":
            fallback = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "1"
            fallback.append(text)
            paragraph._p.append(fallback)


def _new_document(course: dict[str, Any], profile: dict[str, Any]) -> Document:
    template = files("mynotes").joinpath("resources/base-template.docx")
    document = Document(BytesIO(template.read_bytes()))
    _configure_document(document, profile, set_page_layout=True)
    return document


def preview_style_document(profile: dict[str, Any], output_path: Path) -> Path:
    """Build a self-contained Word preview using synthetic teaching content."""
    from PIL import Image, ImageDraw, ImageFont

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = _new_document({"name": "Anteprima"}, profile)
    writing = profile.get("writing_style", "clear_study")
    explanations = {
        "concise": "Inventory helps meet demand, but carrying it has a cost.",
        "clear_study": "[[term:Inventory]] helps meet demand when replenishment is not immediate. The chosen quantity balances **service level** against holding cost.",
        "detailed": "[[term:Inventory]] helps meet demand when replenishment is not immediate. A larger quantity may reduce missed sales, but raises tied-up capital and holding cost. The choice therefore compares **service level**, demand and lead time together.",
    }
    bullet_counts = {"prose_first": 1, "balanced": 2, "list_first": 4}
    bullets = [
        "Estimate expected demand.",
        "Assess replenishment lead time.",
        "Compare service and cost.",
        "Revisit the choice when conditions change.",
    ][:bullet_counts.get(str(profile.get("list_density")), 2)]
    blocks: list[dict[str, Any]] = [
        {"kind": "overview", "items": ["Inventory and demand", "Service level", "Holding cost"]},
        {"kind": "heading", "text": "Inventory and service level", "level": 3},
        {"kind": "paragraph", "text": explanations.get(str(writing), explanations["clear_study"])},
        {"kind": "heading", "text": "How to assess the choice", "level": 4},
        {"kind": "bullets", "items": bullets},
        {"kind": "table", "rows": [["Period", "Demand", "Stock"], ["1", "100", "120"], ["2", "110", "120"]]},
        {"kind": "formula", "label": "Available stock after demand", "expression": "I_{end}=I_{start}-D"},
    ]
    visuals: list[dict[str, Any]] = []
    decisions: list[dict[str, Any]] = []
    with tempfile.TemporaryDirectory(prefix="mynotes-style-preview-") as temporary:
        if profile.get("image_policy", "essential") != "none":
            image_path = Path(temporary) / "diagramma-sintetico.png"
            image = Image.new("RGB", (900, 240), "white")
            draw = ImageDraw.Draw(image)
            fill, outline = "#E7F0F5", "#0F4761"
            font = ImageFont.load_default(size=28)
            for left, label in ((30, "Demand"), (340, "Inventory"), (650, "Service")):
                draw.rounded_rectangle((left, 55, left + 220, 185), radius=18, fill=fill, outline=outline, width=5)
                text_width = draw.textbbox((0, 0), label, font=font)[2]
                draw.text((left + (220 - text_width) / 2, 106), label, fill=outline, font=font)
            draw.line((250, 120, 335, 120), fill="#ED7D31", width=7)
            draw.line((560, 120, 645, 120), fill="#ED7D31", width=7)
            image.save(image_path)
            visuals.append({"evidence_id": "PREVIEW", "image_path": str(image_path), "source_file": "diagramma-sintetico.png"})
            decisions.append({"evidence_id": "PREVIEW", "decision": "include", "reason": "Synthetic preview diagram"})
            blocks.append({"kind": "visual", "visual_id": "PREVIEW"})
        if profile.get("image_policy") == "explanatory":
            blocks.append({"kind": "paragraph", "text": "Explanatory images accompany a step when they add useful information."})
        sample_questions = [
            ("Perché un'impresa mantiene una scorta quando il rifornimento non è immediato?", "La scorta copre la domanda mentre il rifornimento è in corso."),
            ("Come cambia il compromesso tra livello di servizio e costo quando aumenta la scorta?", "Una scorta maggiore può migliorare il servizio, ma aumenta il costo di mantenimento."),
        ]
        for question, answer in sample_questions[:min(int(profile.get("question_count", 5)), 2)]:
            blocks.append({"kind": "qa", "question": question, "answer": answer})
        if profile.get("recap", True):
            recap = "Inventory supports service, but carrying it has a cost." if profile.get("recap_language") in {"en", "same_as_notes"} else "Le scorte sostengono il servizio, ma comportano un costo."
            blocks.append({"kind": "recap", "text": recap})
        notes = {"lecture_id": "PREVIEW-L01", "title": "Layout example", "blocks": blocks, "visual_decisions": decisions}
        _append_block(document, notes, {"visuals": visuals}, recap_language=str(profile.get("recap_language", "it")))
        # This synthetic preview is never published: hidden lesson boundaries
        # are unnecessary and can create a blank trailing page when it fills.
        for paragraph in list(document.paragraphs):
            if paragraph.text in {"MYNOTES START PREVIEW-L01", "MYNOTES END PREVIEW-L01"}:
                paragraph._p.getparent().remove(paragraph._p)
        document.save(output_path)
    return output_path


def _replace_or_append(document: Document, notes: dict[str, Any], pack: dict[str, Any], *, language: str = "en", recap_language: str = "it") -> None:
    lecture_id = notes["lecture_id"]
    boundary = _boundaries(document, lecture_id)
    body = document._body._element
    before = list(body)
    preceding = before[:boundary[0]] if boundary is not None else before
    page_break_before = any(item.tag != qn("w:sectPr") and "".join(item.itertext()).strip() for item in preceding)
    _append_block(document, notes, pack, page_break_before=page_break_before, language=language, recap_language=recap_language)
    added = [item for item in list(body)[len(before):] if item.tag != qn("w:sectPr")]
    for item in added:
        body.remove(item)
    if boundary is None:
        insertion = len(body) - (1 if len(body) and body[-1].tag == qn("w:sectPr") else 0)
    else:
        start, end = boundary; insertion = start
        for item in before[start:end + 1]:
            body.remove(item)
    for item in added:
        body.insert(insertion, item); insertion += 1


def _course(configuration: dict[str, Any], code: str) -> dict[str, Any]:
    for course in configuration["courses"]:
        if course["code"] == code:
            return course
    raise PublicationError(f"Course is not configured: {code}")


def _target(course: dict[str, Any], workspace: Workspace) -> Path:
    from .note_files import handout_path
    return handout_path(course, workspace)


def _load_candidate(workspace: Workspace, lecture_id: str) -> tuple[dict[str, Any], dict[str, Any]]:
    path = workspace.runs_dir / lecture_id / "lecture_notes.candidate.json"
    try:
        candidate = json.loads(path.read_text(encoding="utf-8"))
        pack = json.loads(Path(candidate["source_pack"]).read_text(encoding="utf-8"))
    except (OSError, KeyError, json.JSONDecodeError) as error:
        raise PublicationError(f"Cannot load candidate/source pack for {lecture_id}: {error}") from error
    if candidate.get("published_to_word") is True:
        pass
    if candidate.get("source_pack_hash") != file_sha256(Path(candidate["source_pack"])):
        raise PublicationError("Source pack changed after candidate generation")
    if candidate.get("repair_lineage"):
        try:
            from .batch_recovery import verify_repair_lineage
            verify_repair_lineage(candidate,workspace.root)
        except (OSError, ValueError, KeyError, TypeError, RuntimeError) as error:
            raise PublicationError("Targeted correction provenance cannot be verified: " + str(error)) from error
    # Earlier interactive repairs accepted ERROR envelopes as candidates.
    # Reject those historical transports before rendering or publication.
    for name in candidate.get("shared_attempt_paths", []):
        attempt = Path(name)
        if not attempt.is_absolute():
            continue  # Legacy/injected adapters have no persisted transport.
        if not attempt.resolve().is_relative_to(workspace.root.resolve()):
            raise PublicationError("Candidate transport is outside its private workspace")
        try:
            record = json.loads(attempt.read_text())
            if record.get("transport") == "agy_local_recovery_v1":
                from .local_recovery import verify_local_receipt
                verify_local_receipt(attempt, workspace.root)
                continue
            if record.get("envelope", {}).get("transport") == "agy_parts_v1":
                from .agy_parts import verify_receipt
                verify_receipt(record, workspace.root)
                envelope = record["envelope"]
            else:
                envelope = json.loads(record.get("stdout", "{}")) if record.get("requested", {}).get("provider") == "agy" else {}
        except (OSError, ValueError, TypeError, KeyError, RuntimeError) as error:
            raise PublicationError("Candidate transport cannot be verified") from error
        if record.get("error") or record.get("returncode") != 0 or (isinstance(envelope, dict) and str(envelope.get("status", "SUCCESS")).upper() != "SUCCESS"):
            raise PublicationError("Candidate came from a failed provider transport; retained for review, not Word publication")
    return candidate, pack


def _soffice() -> str | None:
    bundled = sorted((Path.home() / ".cache" / "codex-runtimes").glob("*/dependencies/bin/override/soffice"))
    from .onboarding import managed_tool
    candidates = [os.environ.get("MYNOTES_SOFFICE"), managed_tool("soffice"), shutil.which("soffice"), "/Applications/LibreOffice.app/Contents/MacOS/soffice", *(str(item) for item in bundled)]
    return next((str(Path(item)) for item in candidates if item and Path(item).is_file()), None)


def _render_and_validate(docx: Path, output_dir: Path, *, timeout_seconds: float = 180) -> dict[str, Any]:
    soffice = _soffice()
    if not soffice:
        raise PublicationError("LibreOffice/soffice is required to validate Word rendering before publication")
    try:
        pdf = convert_to_pdf(soffice, docx, output_dir, timeout=timeout_seconds)
    except OfficeError as error:
        raise PublicationError(str(error)) from error
    rendered = pdfium.PdfDocument(pdf)
    pages = []
    try:
        for index, page in enumerate(rendered):
            text = page.get_textpage().get_text_range().strip()
            images = sum(1 for obj in page.get_objects() if obj.type == pdfium_raw.FPDF_PAGEOBJ_IMAGE)
            if not text and not images:
                raise PublicationError(f"Rendered Word contains an empty page at {index + 1}")
            pages.append({"page": index + 1, "characters": len(text), "images": images})
    finally:
        rendered.close()
    if not pages:
        raise PublicationError("Rendered Word has no pages")
    return {"pdf": str(pdf), "pages": pages, "renderer": soffice}


def _docx_structure(path: Path, lecture_id: str, expected_visuals: int, expected_formula: bool,
                    expected_tables: list[list[list[str]]] | None = None) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml")
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    start = document_xml.count(f"MYNOTES START {lecture_id}".encode())
    end = document_xml.count(f"MYNOTES END {lecture_id}".encode())
    if (start, end) != (1, 1):
        raise PublicationError("Candidate Word does not contain exactly one lecture boundary pair")
    equations = document_xml.count(b"<m:oMath")
    if expected_formula and not equations:
        raise PublicationError("Formula was not written as a native Word equation")
    if len(media) < expected_visuals:
        raise PublicationError("One or more selected visuals were not embedded in Word")
    document = Document(path)
    in_lesson = False
    tables: list[list[list[str]]] = []
    for element in document._body._element:
        value = "".join(node.text or "" for node in element.iter(qn("w:t")))
        if f"MYNOTES START {lecture_id}" in value:
            in_lesson = True
            continue
        if f"MYNOTES END {lecture_id}" in value:
            break
        if not in_lesson:
            continue
        # Include cell paragraphs, unlike Document.paragraphs, which excludes
        # them. Scope the guard to this generated lesson, not historical notes.
        for paragraph in element.iter(qn("w:p")):
            text = "".join(node.text or "" for node in paragraph.iter(qn("w:t")))
            if _formatting_residue(text):
                raise PublicationError("Candidate Word contains literal Markdown or escaped newlines")
        if element.tag == qn("w:tbl"):
            tables.append([["".join(node.text or "" for node in cell.iter(qn("w:t")))
                            for cell in row.findall(qn("w:tc"))] for row in element.findall(qn("w:tr"))])
    # Inline emphasis and native equations can change plain-text extraction;
    # compare the native table shapes here and cell content in renderer tests.
    unmatched = list(tables)
    for expected in expected_tables or []:
        shape = [len(row) for row in expected]
        match = next((index for index, actual in enumerate(unmatched) if [len(row) for row in actual] == shape), None)
        if match is None:
            raise PublicationError("Candidate Word is missing an expected native table")
        unmatched.pop(match)
    for section in document.sections:
        if any(abs(margin.cm - 1.5) >= 0.01 for margin in
               (section.top_margin, section.bottom_margin, section.left_margin, section.right_margin)):
            raise PublicationError("Candidate Word must use 1.5 cm margins in every section")
    return {"markers": {"start": start, "end": end}, "equations": equations, "embedded_media": len(media),
            "lesson_tables": len(tables), "margin_cm": 1.5, "sections": len(document.sections)}


def build_word_candidate(workspace: Workspace, lecture_id: str, *, timeout_seconds: float | None = None, mode: str = "update") -> dict[str, Any]:
    if mode not in {"update","standalone"}:raise PublicationError("Unknown Word build mode")
    build_started = time.monotonic()
    lecture_id = lecture_id.upper(); code = lecture_id.split("-L", 1)[0]
    configuration = workspace.configuration_store().load(); course = _course(configuration, code)
    profile = effective_style(configuration, code)
    source_candidate = workspace.runs_dir / lecture_id / "lecture_notes.candidate.json"
    source_candidate_hash = file_sha256(source_candidate)
    candidate, pack = _load_candidate(workspace, lecture_id); notes = candidate["notes"]
    image_limit = candidate.get("request_options", {}).get("image_limit")
    if image_limit is None:
        image_limit = profile.get("image_limit")
    if image_limit is not None:
        if type(image_limit) is not int or image_limit < 0:
            raise PublicationError("Candidate image_limit must be a non-negative integer")
        if sum(item.get("decision") == "include" for item in notes.get("visual_decisions", [])) > image_limit:
            raise PublicationError("Candidate exceeds the effective image_limit; review the content before building")
    if profile.get("image_policy") == "none" and any(item.get("decision") == "include" for item in notes.get("visual_decisions", [])):
        raise PublicationError("Image policy is none, but this candidate includes visuals; regenerate the notes with the current course style")
    target = _target(course, workspace)
    if mode == "update" and not target.exists() and workspace.database_path.is_file():
        from .workflow import WorkflowState
        check = WorkflowState(workspace.database_path)
        try:
            if check.connection.execute("SELECT 1 FROM lesson_runs WHERE output_path=? AND status='PUBLISHED'", (str(target),)).fetchone():
                raise PublicationError("Saved final handout is missing; locate it or restore a backup before updating")
        finally:
            check.close()
    run_dir = workspace.runs_dir / lecture_id / "word"
    run_dir.mkdir(parents=True, exist_ok=True)
    output = run_dir / f"{lecture_id}-{uuid.uuid4().hex[:10]}.candidate.docx"
    baseline_hash = file_sha256(target) if mode == "update" and target.is_file() else None
    if mode == "update" and target.is_file():
        document = Document(target)
        _configure_document(document, profile)
    else:
        document = _new_document(course, profile)
    _replace_or_append(document, notes, pack, language=str(course.get("languages", {}).get("notes", "en")), recap_language=str(profile.get("recap_language", "it")))
    document.save(output)
    expected_visuals = sum(1 for item in notes.get("visual_decisions", []) if item.get("decision") == "include")
    expected_formula = bool(notes.get("formula") or notes.get("formulas") or any(block.get("kind") == "formula" for block in notes.get("blocks", [])))
    expected_tables = ([block["rows"] for block in notes["blocks"] if block.get("kind") == "table"]
                       if notes.get("blocks") else ([notes["table"]] if notes.get("table") else []))
    structure = _docx_structure(output, lecture_id, expected_visuals, expected_formula, expected_tables)
    remaining = timeout_seconds - (time.monotonic()-build_started) if timeout_seconds is not None else None
    if remaining is not None and remaining <= 0:raise PublicationError('Word build exhausted its active budget; candidate retained')
    render = (_render_and_validate(output, run_dir / f"{output.stem}-render", timeout_seconds=remaining)
              if timeout_seconds is not None else _render_and_validate(output, run_dir / f"{output.stem}-render"))
    if file_sha256(source_candidate) != source_candidate_hash:
        raise PublicationError("Notes candidate changed during Word build; rebuild from the current notes")
    validation = {"lecture_id": lecture_id, "mode": mode, "source_candidate_hash": source_candidate_hash, "renderer_hash": word_renderer_hash(), "candidate": str(output), "candidate_hash": file_sha256(output), "target": str(target), "baseline_hash": baseline_hash, "structure": structure, "render": render, "validated_at": now_compact_utc()}
    atomic_json_write(output.with_suffix(".validation.json"), validation)
    from .workflow import WorkflowState
    state = WorkflowState(workspace.database_path)
    try:
        state.connection.execute("UPDATE candidate_runs SET status='WORD_CANDIDATE_VALIDATED', updated_at=? WHERE lecture_id=?", (now_compact_utc(), lecture_id))
        state.connection.commit()
    finally:
        state.close()
    return {"status": "WORD_CANDIDATE_VALIDATED", **validation}


@contextmanager
def _lock(target: Path) -> Iterator[None]:
    lock = target.with_suffix(target.suffix + ".mynotes.lock"); lock.parent.mkdir(parents=True, exist_ok=True)
    with lock.open("a+", encoding="utf-8") as handle:
        try:
            fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError as error:
            raise PublicationError(f"Another publication is active for {target}") from error
        try:
            yield
        finally:
            fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def _journal(workspace: Workspace, payload: dict[str, Any]) -> Path:
    directory = workspace.state_dir / "publication-transactions"; directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{payload['transaction_id']}.json"; atomic_json_write(path, payload); return path


def _content_fingerprint(payload):
    return hashlib.sha256(json.dumps({k: v for k, v in payload.items()
        if k not in {"published_to_word", "published_path", "published_hash"}}, sort_keys=True).encode()).hexdigest()


def _publication_state(workspace, target):
    from .workflow import WorkflowState
    bootstrap = WorkflowState(workspace.database_path)
    bootstrap.close()
    with sqlite3.connect(workspace.database_path) as connection:
        connection.row_factory = sqlite3.Row
        rows = [dict(row) for row in connection.execute("SELECT * FROM lesson_runs")
                if Path(row["output_path"]).resolve() == target.resolve()]
    candidates = {}
    for row in rows:
        path = workspace.runs_dir / row['lecture_id'] / 'lecture_notes.candidate.json'
        if path.is_file():
            payload = json.loads(path.read_text())
            candidates[row['lecture_id']] = {"content_hash": _content_fingerprint(payload),
                "publication": {k: payload[k] for k in ('published_to_word', 'published_path', 'published_hash') if k in payload}}
    return {'lessons': rows, 'candidates': candidates}


def _restore_publication_state(workspace, target, snapshot):
    """Restore Word membership without discarding newer, reusable content candidates."""
    from .workflow import WorkflowState, _write_status_markdown
    bootstrap = WorkflowState(workspace.database_path)
    bootstrap.close()
    snapshot = snapshot or {'lessons': [], 'candidates': {}}
    digest = file_sha256(target) if target.is_file() else ''
    with sqlite3.connect(workspace.database_path) as connection:
        affected = {row[0] for row in connection.execute('SELECT lecture_id,output_path FROM lesson_runs')
                    if Path(row[1]).resolve() == target.resolve()}
        affected.update(row['lecture_id'] for row in snapshot['lessons'])
        connection.execute("UPDATE lesson_runs SET status='ROLLED_BACK',updated_at=? WHERE output_path=?", (now_compact_utc(), str(target)))
        for row in snapshot['lessons']:
            row = dict(row)
            if row['status'] == 'PUBLISHED':
                row['output_hash'] = digest
            row['updated_at'] = now_compact_utc()
            columns = ['lecture_id','input_hash','status','provider','model','effort','quality','output_path','output_hash','sources_json','updated_at']
            connection.execute('INSERT OR REPLACE INTO lesson_runs VALUES(?,?,?,?,?,?,?,?,?,?,?)', [row[k] for k in columns])
        for lecture in affected:
            path = workspace.runs_dir / lecture / 'lecture_notes.candidate.json'
            if not path.is_file():
                connection.execute("UPDATE candidate_runs SET status='NEEDS_REVIEW' WHERE lecture_id=?", (lecture,))
                continue
            payload = json.loads(path.read_text())
            recorded = connection.execute('SELECT candidate_hash FROM candidate_runs WHERE lecture_id=?', (lecture,)).fetchone()
            if recorded and recorded[0] != file_sha256(path):
                # Preserve external edits, but never leave them certified as published.
                connection.execute("UPDATE candidate_runs SET status='NEEDS_REVIEW' WHERE lecture_id=?", (lecture,))
                continue
            saved = snapshot.get('candidates', {}).get(lecture)
            matches = saved and saved['content_hash'] == _content_fingerprint(payload)
            for key in ('published_to_word', 'published_path', 'published_hash'):
                payload.pop(key, None)
            payload['published_to_word'] = False
            if matches:
                payload.update(saved['publication'])
                if payload.get('published_to_word'):
                    payload['published_hash'] = digest
            atomic_json_write(path, payload)
            connection.execute('UPDATE candidate_runs SET status=?,candidate_hash=?,updated_at=? WHERE lecture_id=?',
                               ('PUBLISHED' if payload.get('published_to_word') else 'CANDIDATE_READY', file_sha256(path), now_compact_utc(), lecture))
    state = WorkflowState(workspace.database_path)
    try:
        _write_status_markdown(workspace, state.rows())
    finally:
        state.close()


def _archive(workspace: Workspace, target: Path, course_code: str, *, reason: str) -> dict[str, Any] | None:
    if not target.is_file():
        return None
    directory = workspace.archives_dir / "word" / course_code; directory.mkdir(parents=True, exist_ok=True)
    original_hash = file_sha256(target)
    backup_id = f"{now_compact_utc()}-{original_hash[:12]}-{uuid.uuid4().hex[:8]}"
    path = directory / f"{backup_id}.docx"; shutil.copy2(target, path)
    if file_sha256(path) != original_hash or file_sha256(target) != original_hash:
        raise PublicationError("Word changed while saving its safety backup; no update was made")
    metadata = {"backup_id": backup_id, "course": course_code, "path": str(path), "sha256": file_sha256(path), "source": str(target), "reason": reason, "created_at": now_compact_utc(), "publication_state": _publication_state(workspace, target)}
    atomic_json_write(directory / f"{backup_id}.json", metadata); return metadata


def publish_word_candidate(workspace: Workspace, lecture_id: str, candidate_path: Path | None = None, *, accept_current_word: bool = False) -> dict[str, Any]:
    lecture_id = lecture_id.upper(); code = lecture_id.split("-L", 1)[0]
    from .workflow import WorkflowState
    state_bootstrap = WorkflowState(workspace.database_path); state_bootstrap.close()
    course = _course(workspace.configuration_store().load(), code); target = _target(course, workspace)
    if candidate_path is None:
        matches = sorted((workspace.runs_dir / lecture_id / "word").glob("*.candidate.docx"), key=lambda path: path.stat().st_mtime)
        if not matches:
            raise PublicationError("No validated Word candidate exists; run mynotes build first")
        candidate_path = matches[-1]
    validation_path = candidate_path.with_suffix(".validation.json")
    try:
        validation = json.loads(validation_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise PublicationError("Candidate has no valid render/structure report") from error
    if validation.get("lecture_id") != lecture_id:
        raise PublicationError("Validated Word candidate belongs to a different lecture")
    if not isinstance(validation.get("target"), str) or Path(validation["target"]).resolve() != target.resolve():
        raise PublicationError("Validated Word candidate belongs to a different course target; rebuild for this course")
    candidate_hash = file_sha256(candidate_path)
    if validation.get("candidate_hash") != candidate_hash:
        raise PublicationError("Candidate changed after validation")
    if validation.get("mode","update") != "update":
        raise PublicationError("Standalone or cumulative candidates require their own explicit publication workflow")
    from .workflow import _generation_lock
    with _generation_lock(workspace, lecture_id), _lock(target):
        current = file_sha256(target) if target.is_file() else None
        with sqlite3.connect(workspace.database_path) as connection:
            published = connection.execute("SELECT status,output_path,output_hash FROM lesson_runs WHERE lecture_id=?", (lecture_id,)).fetchone()
        if current == candidate_hash and published and published[0] == "PUBLISHED" and Path(published[1]).resolve() == target.resolve() and published[2] == current:
            return {"status": "SKIPPED", "lecture_id": lecture_id, "output": str(target), "reason": "validated_candidate_already_published"}
        source_candidate = workspace.runs_dir / lecture_id / "lecture_notes.candidate.json"
        source_candidate_hash = validation.get("source_candidate_hash")
        if not source_candidate_hash:
            raise PublicationError("Word candidate has no notes fingerprint; rebuild before publishing")
        if file_sha256(source_candidate) != source_candidate_hash:
            raise PublicationError("Notes candidate changed after Word build; rebuild before publishing")
        payload, source_pack = _load_candidate(workspace, lecture_id)
        sources = source_pack.get("availability", {})
        transaction = {"transaction_id": uuid.uuid4().hex, "lecture_id": lecture_id, "target": str(target), "candidate": str(candidate_path), "candidate_hash": candidate_hash, "baseline_hash": validation.get("baseline_hash"), "status": "PREPARED", "created_at": now_compact_utc()}
        journal = _journal(workspace, transaction)
        if current != transaction["baseline_hash"] and not accept_current_word:
            transaction["status"] = "NEEDS_REVIEW"; transaction["reason"] = "target_changed"; _journal(workspace, transaction)
            raise PublicationError("Cumulative Word changed after candidate creation; rebuild on the current file or explicitly accept it")
        transaction["content_hash"] = _content_fingerprint(payload)
        transaction["publication_state"] = _publication_state(workspace, target)
        transaction["backup"] = _archive(workspace, target, code, reason=f"before {lecture_id} publication")
        transaction["status"] = "BACKED_UP"; _journal(workspace, transaction)
        descriptor, temporary_name = tempfile.mkstemp(suffix=".docx", dir=target.parent); os.close(descriptor)
        temporary = Path(temporary_name)
        try:
            shutil.copy2(candidate_path, temporary)
            if file_sha256(temporary) != candidate_hash:
                raise PublicationError("Word candidate changed during publication")
            if file_sha256(source_candidate) != source_candidate_hash:
                raise PublicationError("Notes candidate changed during publication; rebuild before publishing")
            if (file_sha256(target) if target.is_file() else None) != current:
                raise PublicationError("Cumulative Word changed during publication")
            os.replace(temporary, target)
        finally:
            temporary.unlink(missing_ok=True)
        transaction["status"] = "PUBLISHED"; transaction["target_hash"] = file_sha256(target); _journal(workspace, transaction)
        if transaction["target_hash"] != transaction["candidate_hash"]:
            raise PublicationError("Published Word hash does not match validated candidate")
        if file_sha256(source_candidate) != source_candidate_hash:
            raise PublicationError("Notes candidate changed during publication; inspect the published Word before recovery")
        payload["published_to_word"] = True
        payload["published_path"] = str(target); payload["published_hash"] = transaction["target_hash"]
        atomic_json_write(source_candidate, payload)
        requested = payload.get("requested_model", {}); actual = payload.get("actual_model", {})
        published_source_hash = file_sha256(source_candidate)
        with sqlite3.connect(workspace.database_path) as connection:
            row = connection.execute("SELECT input_hash FROM candidate_runs WHERE lecture_id=?", (lecture_id,)).fetchone()
            input_hash = row[0] if row else payload.get("source_pack_hash", "")
            connection.execute(
                """INSERT INTO lesson_runs VALUES(?,?,?,?,?,?,?,?,?,?,?)
                   ON CONFLICT(lecture_id) DO UPDATE SET input_hash=excluded.input_hash,status=excluded.status,
                   provider=excluded.provider,model=excluded.model,effort=excluded.effort,quality=excluded.quality,
                   output_path=excluded.output_path,output_hash=excluded.output_hash,sources_json=excluded.sources_json,
                   updated_at=excluded.updated_at""",
                (lecture_id, input_hash, "PUBLISHED", actual.get("provider", requested.get("provider", "unknown")), actual.get("model", requested.get("model", "unknown")), actual.get("effort", requested.get("effort", "unknown")), payload.get("quality", requested.get("quality", "standard")), str(target), transaction["target_hash"], json.dumps(sources, sort_keys=True), now_compact_utc()),
            )
            connection.execute("UPDATE candidate_runs SET status='PUBLISHED', candidate_hash=?, updated_at=? WHERE lecture_id=?", (published_source_hash, now_compact_utc(), lecture_id))
        from .workflow import WorkflowState, _write_status_markdown
        state = WorkflowState(workspace.database_path)
        try:
            _write_status_markdown(workspace, state.rows())
        finally:
            state.close()
        transaction["status"] = "COMPLETED"; transaction["completed_at"] = now_compact_utc(); _journal(workspace, transaction)
    return {"status": "PUBLISHED", "lecture_id": lecture_id, "output": str(target), "backup": transaction.get("backup"), "transaction": str(journal), "undo_hint": "Puoi chiedermi di annullare l’ultimo aggiornamento."}


def update_course_notes(workspace: Workspace, lecture_id: str) -> dict[str, Any]:
    """Build from the current cumulative and publish unless preview was requested."""
    from .workflow import WorkflowState
    lecture_id = lecture_id.upper()
    course = _course(workspace.configuration_store().load(), lecture_id.split('-L', 1)[0])
    target = _target(course, workspace)
    # Repeating a completed request must not create another backup or Word version.
    source = workspace.runs_dir / lecture_id / 'lecture_notes.candidate.json'
    if source.is_file() and target.is_file():
        payload, _ = _load_candidate(workspace, lecture_id)
        state = WorkflowState(workspace.database_path)
        try:
            row = state.connection.execute("SELECT status FROM lesson_runs WHERE lecture_id=?", (lecture_id,)).fetchone()
            expected = state.expected_output_hash(target)
        finally:
            state.close()
        if expected and expected != file_sha256(target):
            raise PublicationError('Final Word was edited externally; review it before an automatic update. No file was overwritten.')
        if row and row[0] == 'PUBLISHED' and payload.get('published_to_word') and expected == file_sha256(target):
            for record in _transactions(workspace, target):
                if record.get('status') == 'COMPLETED' and record.get('lecture_id') == lecture_id:
                    if record.get('content_hash') == _content_fingerprint(payload):
                        return {'status': 'SKIPPED', 'lecture_id': lecture_id, 'output': str(target), 'reason': 'already_published'}
                    break
    built = build_word_candidate(workspace, lecture_id)
    if course['publication']['mode'] == 'review':
        return {**built, 'status': 'AWAITING_APPROVAL', 'next_action': 'Show this candidate and publish only after approval.'}
    result = publish_word_candidate(workspace, lecture_id, Path(built['candidate']))
    return {**result, 'validation': built}


def _transactions(workspace, target):
    root = workspace.state_dir / 'publication-transactions'
    records = []
    for path in root.glob('*.json'):
        record = json.loads(path.read_text())
        if Path(record['target']).resolve() == target.resolve():
            records.append(record)
    return sorted(records, key=lambda r: r.get('created_at', ''), reverse=True)


@contextmanager
def _restore_locks(workspace, target, extra=()):
    """Freeze affected candidate metadata as well as the cumulative Word."""
    from .workflow import WorkflowState, _generation_lock
    bootstrap = WorkflowState(workspace.database_path)
    bootstrap.close()
    with sqlite3.connect(workspace.database_path) as connection:
        lectures = {row[0] for row in connection.execute('SELECT lecture_id,output_path FROM lesson_runs')
                    if Path(row[1]).resolve() == target.resolve()}
    lectures.update(extra)
    with ExitStack() as stack:
        for lecture in sorted(lectures):
            stack.enter_context(_generation_lock(workspace, lecture))
        stack.enter_context(_lock(target))
        yield


def undo_last_update(workspace: Workspace, course_code: str) -> dict[str, Any]:
    """Undo one completed publication, including a document's first creation."""
    course = _course(workspace.configuration_store().load(), course_code.upper())
    target = _target(course, workspace)
    records = [r for r in _transactions(workspace, target) if r.get('status') == 'COMPLETED']
    if not records:
        raise PublicationError('No completed update is available to undo for this course')
    digest = file_sha256(target) if target.is_file() else None
    record = next((r for r in records if r.get('target_hash') == digest), None)
    if record is None:
        raise PublicationError('Final Word changed since the last update; review it before undoing. No file was overwritten.')
    with _restore_locks(workspace, target, [record['lecture_id']]):
        current = file_sha256(target) if target.is_file() else None
        if current != record.get('target_hash'):
            raise PublicationError('Final Word changed since the last update; review it before undoing. No file was overwritten.')
        backup = record.get('backup')
        source = Path(backup['path']) if backup else None
        if source is not None and (not source.is_file() or file_sha256(source) != backup['sha256']):
            raise PublicationError('Backup is missing or corrupted; the final Word was preserved')
        safety = _archive(workspace, target, course['code'], reason=f"before undo {record['transaction_id']}")
        if safety is None:
            raise PublicationError('Could not save the current Word before undo')
        record['undo'] = {'status': 'PREPARED', 'backup': safety,
                          'restored_hash': backup['sha256'] if backup else None}
        _journal(workspace, record)
        if source is not None:
            fd, name = tempfile.mkstemp(suffix='.docx', dir=target.parent)
            os.close(fd)
            temporary = Path(name)
            try:
                shutil.copy2(source, temporary)
                if file_sha256(temporary) != backup['sha256'] or file_sha256(target) != current:
                    raise PublicationError('Word or backup changed during undo')
                os.replace(temporary, target)
            finally:
                temporary.unlink(missing_ok=True)
        else:
            if file_sha256(target) != current:
                raise PublicationError('Word changed during undo')
            target.unlink()  # The removed first creation remains in the verified safety backup.
        _restore_publication_state(workspace, target, record.get('publication_state'))
        record['status'] = 'UNDONE'
        record['undo']['status'] = 'COMPLETED'
        record['undo']['completed_at'] = now_compact_utc()
        _journal(workspace, record)
    return {'status': 'UNDONE', 'course': course['code'], 'lecture_id': record['lecture_id'],
            'output': str(target), 'output_exists': target.is_file(), 'pre_restore_backup': safety}


def publication_history(workspace: Workspace, course_code: str | None = None) -> list[dict[str, Any]]:
    root = workspace.archives_dir / "word"
    files = root.glob("*/*.json") if not course_code else (root / course_code.upper()).glob("*.json")
    result = []
    for path in sorted(files):
        try:
            result.append(json.loads(path.read_text(encoding="utf-8")))
        except (OSError, json.JSONDecodeError):
            continue
    return result


def restore_backup(workspace: Workspace, backup_id: str) -> dict[str, Any]:
    matches = [item for item in publication_history(workspace) if item.get("backup_id") == backup_id]
    if len(matches) != 1:
        raise PublicationError("Backup id was not found or is ambiguous")
    metadata = matches[0]; source = Path(metadata["path"]); target = Path(metadata["source"])
    if not source.is_file() or file_sha256(source) != metadata["sha256"]:
        raise PublicationError("Backup is missing or corrupted")
    with _restore_locks(workspace, target, [row["lecture_id"] for row in metadata.get("publication_state", {}).get("lessons", [])]):
        current_hash = file_sha256(target) if target.is_file() else None
        safety = _archive(workspace, target, metadata["course"], reason=f"before restore {backup_id}")
        descriptor, name = tempfile.mkstemp(suffix=".docx", dir=target.parent); os.close(descriptor); temporary = Path(name)
        try:
            shutil.copy2(source, temporary)
            if file_sha256(temporary) != metadata['sha256'] or (file_sha256(target) if target.is_file() else None) != current_hash:
                raise PublicationError("Word or backup changed during restore")
            os.replace(temporary, target)
            _restore_publication_state(workspace, target, metadata.get('publication_state'))
        finally:
            temporary.unlink(missing_ok=True)
    return {"status": "RESTORED", "backup_id": backup_id, "output": str(target), "pre_restore_backup": safety}


def recover_publications(workspace: Workspace) -> list[dict[str, Any]]:
    directory = workspace.state_dir / "publication-transactions"
    results = []
    for path in sorted(directory.glob("*.json")) if directory.is_dir() else []:
        payload = json.loads(path.read_text(encoding="utf-8"))
        if payload.get('undo', {}).get('status') == 'PREPARED':
            target = Path(payload['target'])
            with _restore_locks(workspace, target, [payload['lecture_id']]):
                current = file_sha256(target) if target.is_file() else None
                if current == payload['undo']['restored_hash']:
                    _restore_publication_state(workspace, target, payload.get('publication_state'))
                    payload['status'] = 'UNDONE'
                    payload['undo']['status'] = 'COMPLETED'
                elif current == payload.get('target_hash'):
                    payload['undo']['status'] = 'ABORTED'
                else:
                    payload['status'] = 'NEEDS_REVIEW'
                    payload['reason'] = 'ambiguous_undo_state'
                atomic_json_write(path, payload)
                results.append(payload)
            continue
        if payload.get("status") in {"COMPLETED", "ABORTED", "NEEDS_REVIEW", "UNDONE"}:
            continue
        target = Path(payload["target"]); current = file_sha256(target) if target.is_file() else None
        if current == payload.get("candidate_hash"):
            payload["status"] = "COMPLETED"; payload["recovered_at"] = now_compact_utc()
        elif current == payload.get("baseline_hash") or (current is None and payload.get("baseline_hash") is None):
            payload["status"] = "ABORTED"; payload["recovered_at"] = now_compact_utc()
        else:
            payload["status"] = "NEEDS_REVIEW"; payload["reason"] = "ambiguous_recovery_state"
        atomic_json_write(path, payload); results.append(payload)
    return results
