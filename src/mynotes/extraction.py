"""Local extraction of lecture material and user-authored note blocks."""

from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from docx import Document

from .utils import file_sha256


class ExtractionError(RuntimeError):
    """Raised when a declared source cannot be extracted truthfully."""


LECTURE_HEADING_RE = re.compile(r"^\s*(?:lecture|lesson|lezione|l)\s*0*(\d{1,3})(?:\b|\s*[-:])", re.IGNORECASE)
INSTRUCTION_RE = re.compile(r"^\s*(?:NOTE|DOMANDA|ISTRUZIONI)\s+AI\s*:\s*(.*)$", re.IGNORECASE)
INSTRUCTION_END_RE = re.compile(r"^\s*FINE\s+ISTRUZIONI\s+AI\s*$", re.IGNORECASE)
URL_RE = re.compile(r"https://[^\s<>\]\[)\"']+", re.IGNORECASE)


def extract_pdf(path: Path, visual_dir: Path | None = None) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    import pypdfium2 as pdfium
    import pypdfium2.raw as pdfium_raw

    document = pdfium.PdfDocument(path)
    chunks: list[dict[str, Any]] = []
    visuals: list[dict[str, Any]] = []
    try:
        for number, page in enumerate(document, start=1):
            textpage = page.get_textpage()
            raw_text = textpage.get_text_range()
            text = re.sub(r"\s+", " ", raw_text).strip()
            if text:
                chunks.append({"content": text, "page": number})
            image_count = 0
            drawing_count = 0
            for obj in page.get_objects():
                if obj.type == pdfium_raw.FPDF_PAGEOBJ_IMAGE:
                    image_count += 1
                elif obj.type == pdfium_raw.FPDF_PAGEOBJ_PATH:
                    drawing_count += 1
            table_count = 0
            score = min(1.0, (image_count * 0.35) + (table_count * 0.70) + min(drawing_count, 8) * 0.08)
            decision = "candidate" if table_count or drawing_count >= 3 else "review" if image_count else "exclude"
            image_path: str | None = None
            if visual_dir is not None and decision in {"candidate", "review"}:
                target = visual_dir / f"page-{number:03d}.png"
                target.parent.mkdir(parents=True, exist_ok=True)
                img = page.render(scale=2).to_pil()
                img.save(target)
                image_path = str(target)
            visuals.append(
                {
                    "page": number, "source_file": str(path), "image_path": image_path,
                    "visual_score": round(score, 3), "visual_decision": decision,
                    "signals": {"images": image_count, "drawings": drawing_count, "tables": table_count},
                    "reason": (
                        "Contains a structural diagram or table" if decision == "candidate"
                        else "Raster image requires semantic review before inclusion" if decision == "review"
                        else "Predominantly text; prose extraction is preferable"
                    ),
                }
            )
    finally:
        document.close()
    return chunks, visuals


def _pptx_visual_scores(path: Path) -> dict[int, dict[str, int | float]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    scores: dict[int, dict[str, int | float]] = {}
    for number, slide in enumerate(Presentation(path).slides, start=1):
        pictures = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
        charts = sum(bool(getattr(shape, "has_chart", False)) for shape in slide.shapes)
        tables = sum(bool(getattr(shape, "has_table", False)) for shape in slide.shapes)
        groups = sum(shape.shape_type == MSO_SHAPE_TYPE.GROUP for shape in slide.shapes)
        score = min(1.0, pictures * 0.45 + charts * 0.8 + tables * 0.8 + groups * 0.25)
        scores[number] = {"pictures": pictures, "charts": charts, "tables": tables, "groups": groups, "score": score}
    return scores


def _render_pptx_as_pdf(path: Path, output_dir: Path) -> Path | None:
    executable = os.environ.get("MYNOTES_LIBREOFFICE") or shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        return None
    output_dir.mkdir(parents=True, exist_ok=True)
    completed = subprocess.run(
        [executable, "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(path)],
        capture_output=True, text=True, timeout=300,
    )
    candidate = output_dir / f"{path.stem}.pdf"
    return candidate if completed.returncode == 0 and candidate.is_file() else None


def extract_ppt(path: Path, visual_dir: Path | None = None) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Extract legacy PowerPoint through a local LibreOffice PDF conversion."""
    with tempfile.TemporaryDirectory(prefix="mynotes-ppt-") as temporary:
        pdf = _render_pptx_as_pdf(path, Path(temporary))
        if pdf is None:
            raise ExtractionError(
                "Legacy .ppt extraction requires LibreOffice; configure MYNOTES_LIBREOFFICE or install soffice"
            )
        chunks, visuals = extract_pdf(pdf, visual_dir)
    for visual in visuals:
        visual["source_file"] = str(path)
    return chunks, visuals


def extract_xlsx(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Read formulas and cached values with sheet and cell provenance."""
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise ExtractionError("Excel extraction requires openpyxl") from error
    formulas = load_workbook(path, read_only=True, data_only=False)
    values = load_workbook(path, read_only=True, data_only=True)
    chunks: list[dict[str, Any]] = []
    try:
        for sheet in formulas.worksheets:
            cached = values[sheet.title]
            for row_formula, row_cached in zip(sheet.iter_rows(), cached.iter_rows()):
                parts: list[str] = []
                for cell, value_cell in zip(row_formula, row_cached):
                    if cell.value is None:
                        continue
                    if cell.data_type == "f":
                        cached_value = value_cell.value if value_cell.value is not None else "not available"
                        parts.append(f"{cell.coordinate}: formula {cell.value}; cached value {cached_value}")
                    else:
                        parts.append(f"{cell.coordinate}: {cell.value}")
                if parts:
                    chunks.append({"content": "; ".join(parts), "sheet": sheet.title, "row": row_formula[0].row})
    finally:
        formulas.close()
        values.close()
    return chunks, []


def extract_pptx(path: Path, visual_dir: Path | None = None) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    from pptx import Presentation

    chunks: list[dict[str, Any]] = []
    for number, slide in enumerate(Presentation(path).slides, start=1):
        texts: list[str] = []
        table_number = 0
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                table_number += 1
                rows = [" | ".join(" ".join(cell.text.split()) for cell in row.cells) for row in shape.table.rows]
                texts.append(f"Table {table_number} (source rows):\n" + "\n".join(rows))
            elif hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        if texts:
            chunks.append({"content": "\n".join(texts), "page": number})
    signals = _pptx_visual_scores(path)
    rendered: dict[int, str] = {}
    if visual_dir is not None and any(float(item["score"]) >= 0.25 for item in signals.values()):
        with tempfile.TemporaryDirectory(prefix="mynotes-pptx-") as temporary:
            pdf = _render_pptx_as_pdf(path, Path(temporary))
            if pdf:
                _, pages = extract_pdf(pdf, visual_dir)
                rendered = {int(item["page"]): str(item["image_path"]) for item in pages if item.get("image_path")}
    visuals = []
    for number, item in signals.items():
        decision = "candidate" if int(item["charts"]) or int(item["tables"]) or int(item["groups"]) else "review" if int(item["pictures"]) else "exclude"
        visuals.append(
            {
                "page": number, "source_file": str(path), "image_path": rendered.get(number),
                "visual_score": round(float(item["score"]), 3), "visual_decision": decision,
                "signals": {key: value for key, value in item.items() if key != "score"},
                "reason": (
                    "Contains a chart, table, or grouped diagram" if decision == "candidate"
                    else "Raster image requires semantic review before inclusion" if decision == "review"
                    else "Text-only slide; use extracted prose"
                ),
                **({"render_error": "LibreOffice rendering unavailable"} if decision in {"candidate", "review"} and number not in rendered else {}),
            }
        )
    return chunks, visuals


def _docx_items(path: Path) -> list[dict[str, Any]]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(path)
    result: list[dict[str, Any]] = []
    paragraph_number = 0
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph_number += 1
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if text:
                result.append({"content": text, "paragraph": paragraph_number, "style": paragraph.style.name})
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            text = "\n".join(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows).strip()
            if text:
                result.append({"content": text, "paragraph": paragraph_number, "style": "Table"})
    return result


def extract_docx(path: Path, lesson_number: int | None = None) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    items = _docx_items(path)
    if lesson_number is None:
        return items, []
    starts = [index for index, item in enumerate(items) if (match := LECTURE_HEADING_RE.match(item["content"])) and int(match.group(1)) == lesson_number]
    if not starts:
        return [], []
    if len(starts) > 1:
        raise ExtractionError(f"More than one heading identifies lesson {lesson_number} in {path}")
    start = starts[0]
    end = next(
        (index for index in range(start + 1, len(items)) if LECTURE_HEADING_RE.match(items[index]["content"])),
        len(items),
    )
    return items[start:end], []


def extract_material(path: Path, *, visual_dir: Path | None = None, lesson_number: int | None = None) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not path.is_file():
        raise ExtractionError(f"Source file does not exist: {path}")
    suffix = path.suffix.casefold()
    try:
        if suffix == ".pdf":
            return extract_pdf(path, visual_dir)
        if suffix == ".pptx":
            return extract_pptx(path, visual_dir)
        if suffix == ".ppt":
            return extract_ppt(path, visual_dir)
        if suffix == ".xlsx":
            return extract_xlsx(path)
        if suffix == ".docx":
            return extract_docx(path, lesson_number)
        if suffix in {".txt", ".md"}:
            content = path.read_text(encoding="utf-8").strip()
            return ([{"content": content, "paragraph": 1}] if content else []), []
        if suffix == ".ipynb":
            payload = json.loads(path.read_text(encoding="utf-8"))
            chunks = [
                {"content": "".join(cell.get("source", [])).strip(), "cell": number}
                for number, cell in enumerate(payload.get("cells", []), start=1)
                if "".join(cell.get("source", [])).strip()
            ]
            return chunks, []
    except (OSError, ValueError, KeyError) as error:
        raise ExtractionError(f"Could not extract {path}: {error}") from error
    raise ExtractionError(f"Unsupported source format: {path.suffix}")


def split_directives(chunks: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Separate explicit user directives without interpreting source prose as commands."""
    evidence: list[dict[str, Any]] = []
    directives: list[dict[str, Any]] = []
    multiline: list[str] | None = None
    for chunk in chunks:
        lines = str(chunk.get("content", "")).splitlines()
        kept: list[str] = []
        for line in lines:
            if multiline is not None:
                if INSTRUCTION_END_RE.match(line):
                    directives.append({"kind": "instruction", "text": "\n".join(multiline).strip(), "locator": chunk.get("paragraph")})
                    multiline = None
                else:
                    multiline.append(line)
                continue
            match = INSTRUCTION_RE.match(line)
            if match:
                label = line.split(":", 1)[0].strip().casefold()
                kind = "question" if "domanda" in label else "note" if label.startswith("note") else "instruction"
                value = match.group(1).strip()
                if not value and kind == "instruction":
                    multiline = []
                elif value:
                    directives.append({"kind": kind, "text": value, "locator": chunk.get("paragraph")})
                continue
            kept.append(line)
        content = "\n".join(kept).strip()
        if content:
            evidence.append({**chunk, "content": content})
    if multiline is not None:
        directives.append({"kind": "instruction", "text": "\n".join(multiline).strip(), "warning": "missing explicit end marker"})
    return evidence, directives


def source_descriptor(path: Path) -> dict[str, Any]:
    return {"path": str(path), "sha256": file_sha256(path), "size": path.stat().st_size, "modified_ns": path.stat().st_mtime_ns}


def linked_urls(chunks: list[dict[str, Any]]) -> list[str]:
    found: list[str] = []
    for chunk in chunks:
        for url in URL_RE.findall(str(chunk.get("content", ""))):
            cleaned = url.rstrip(".,;:")
            if cleaned not in found:
                found.append(cleaned)
    return found
