"""Minimal first-use defaults and honest, read-only component diagnostics."""
from __future__ import annotations

import copy
import hashlib
import json
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from .utils import atomic_json_write, file_sha256
from .workspace import Workspace

INSTALL_URLS = {
    "codex": "https://developers.openai.com/codex/cli/",
    "agy": "https://antigravity.google/docs/cli/",
    "claude": "https://code.claude.com/docs/en/setup",
}


def runtime_home() -> Path:
    return Path(os.environ.get("MYNOTES_HOME", str(Path.home() / "Library/Application Support/MyNotes"))).expanduser()


def _managed_registry() -> tuple[dict[str, Any], str | None]:
    """Validate diagnostic container types without repairing user files."""
    path = runtime_home() / "tools.json"
    try:
        data = json.loads(path.read_text())
    except FileNotFoundError:
        return {}, None
    except (OSError, ValueError):
        return {}, "Managed component registry cannot be read as JSON."
    if not isinstance(data, dict):
        return {}, "Managed component registry must be a JSON object."
    if "schema_version" in data and data["schema_version"] != 1:
        return {}, "Managed component registry uses an unsupported schema version."
    for key in ("argmax", "soffice", "model_cache"):
        if key in data and (not isinstance(data[key], str) or not data[key] or "\x00" in data[key]):
            return {}, f"Managed component registry has an invalid {key} path."
    if "execution_proof" in data and not isinstance(data["execution_proof"], dict):
        return {}, "Managed component execution proof must be a JSON object."
    return data, None


def managed_tools() -> dict[str, Any]:
    return _managed_registry()[0]


def managed_tool(name: str) -> str | None:
    value = managed_tools().get(name)
    return value if isinstance(value, str) and Path(value).is_file() and os.access(value, os.X_OK) else None


def documents_root() -> Path:
    return Path(os.environ.get("MYNOTES_USER_HOME", str(Path.home()))).expanduser() / "Documents/MyNotes"


def _prepare_defaults(settings: dict, provider: str | None = None) -> dict:
    completed = settings["onboarding"]["completed_groups"]
    docs = documents_root()
    if "paths" not in completed and not settings["paths"]["recording_dirs"]:
        folder = docs / "Recordings"
        folder.mkdir(parents=True, exist_ok=True)
        settings["paths"]["recording_dirs"] = [str(folder)]
    if provider:
        from .configuration import SUPPORTED_PROVIDERS
        if provider not in SUPPORTED_PROVIDERS:
            raise ValueError(f"Unsupported provider: {provider}")
        settings["providers"]["default"] = provider
        settings["providers"]["configured"].setdefault(provider, {})
    # Preparing a computer is not an answer about the learner. Repair the old
    # automatic completion only on an empty profile with no existing courses.
    personal = settings["personal"]
    known_profile = personal.get("background_confirmed") or any(
        personal.get(key) for key in ("background", "institution", "study_program", "study_period")
    ) or bool(settings["courses"])
    if known_profile and "personal" not in completed:
        completed.append("personal")
    elif not known_profile and "personal" in completed:
        completed.remove("personal")
    for group in ("paths", "style"):
        if group not in completed:
            completed.append(group)
    if settings["providers"]["default"] and "providers" not in completed:
        completed.append("providers")
    remaining = [g for g in ("personal", "courses", "providers") if g not in completed]
    settings["onboarding"].update(status="in_progress" if remaining else "complete", next_group=remaining[0] if remaining else None)
    return settings


def prepare_conversation(workspace: Workspace, provider: str | None = None) -> dict:
    """Set machine defaults atomically, leaving first-course answers pending."""
    workspace.initialize()
    return workspace.configuration_store().update(lambda settings: _prepare_defaults(settings, provider))


def prepare_setup_preferences(workspace: Workspace, provider: str) -> dict:
    """Change setup defaults and return a scoped, concurrent-safe undo receipt."""
    workspace.initialize()
    before = {}
    def edit(settings):
        before.update({key: copy.deepcopy(settings[key]) for key in ("providers", "paths", "onboarding")})
        return _prepare_defaults(settings, provider)
    written = workspace.configuration_store().update(edit)
    content = (json.dumps(written, ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode("utf-8")
    return {"before": before, "written_sha256": hashlib.sha256(content).hexdigest()}


def restore_setup_preferences(workspace: Workspace, receipt: dict) -> bool:
    """Undo setup defaults only if no later user edit changed the configuration.

    Course data is never backed up or replaced; the receipt contains only the
    provider/path/onboarding fields touched by setup.
    """
    restored = []
    def edit(settings):
        if file_sha256(workspace.configuration_path) == receipt["written_sha256"]:
            for key in ("providers", "paths", "onboarding"):
                settings[key] = copy.deepcopy(receipt["before"][key])
            restored.append(True)
        return settings
    workspace.configuration_store().update(edit)
    return bool(restored)


def course_setup_questions(language: str) -> str:
    if language == "it":
        return (
            "1. Che corso è e cosa ti serve per studiarlo? Raccontami nome e argomento, come sono le lezioni (teoria, esercizi o casi), come sarà l’esame e cosa trovi più difficile.\n"
            "2. Come vorresti gli appunti? Più discorsivi o schematici, sintetici o approfonditi; nessuna immagine, poche essenziali o più figure spiegate? Indica anche la lingua delle lezioni e degli appunti, ed eventuali preferenze su esempi, formule, riepiloghi e domande di ripasso.\n"
            "Puoi rispondere a entrambe in un unico messaggio, come ti viene naturale."
        )
    return (
        "1. What course is this, and what do you need to study it? Tell me its name and subject, whether lessons focus on theory, exercises or cases, the exam format and what you find difficult.\n"
        "2. How would you like your notes? More prose or outlines, concise or detailed; no images, a few essential ones or more explained figures? Include the lecture and notes languages, plus any preferences for examples, formulas, recaps and review questions.\n"
        "You can answer both in one message, in your own words."
    )


def add_first_course(workspace: Workspace, name: str, notes_language: str = "en", materials: str | None = None, recordings: str | None = None,
                     context_answer: str | None = None, course_settings: dict | None = None) -> dict:
    from .configuration import ConfigurationError, clean_path, merge_course, normalize_course, normalize_style_profile
    from .path_links import PathLinkError, resolve_course_directory
    patch = copy.deepcopy(course_settings) if course_settings is not None else {}
    allowed = {"kind", "languages", "assessment", "preparation", "difficult_topics", "style", "preferences"}
    if not isinstance(patch, dict) or set(patch) - allowed:
        raise ConfigurationError("Course settings may include kind, languages, assessment, preparation, difficult_topics, style and preferences")
    for field in ("assessment", "preparation"):
        if field in patch and not isinstance(patch[field], str):
            raise ConfigurationError(f"Course {field} must be text")
    if "style" in patch:
        patch["style"] = normalize_style_profile(patch["style"], partial=True)
    if context_answer is not None:
        if not isinstance(context_answer, str) or not context_answer.strip():
            raise ConfigurationError("Course context answer must be non-empty text")
        patch["context"] = context_answer
    normalize_course({"code": "COURSE", "name": name, **patch})
    def apply_answer(course):
        updated = merge_course(course, patch)
        if isinstance(patch.get("languages"), dict) and "lesson" in patch["languages"]:
            updated["transcription"]["language"] = updated["languages"]["lesson"]
        return updated
    try:
        materials = str(resolve_course_directory(Path(clean_path(materials)))) if materials else None
        recordings = str(resolve_course_directory(Path(clean_path(recordings)))) if recordings else None
    except PathLinkError as error:
        from .configuration import ConfigurationError
        raise ConfigurationError(str(error)) from error
    workspace.initialize()
    selected_code = []
    def edit(settings):
        settings = _prepare_defaults(settings)
        existing = next((c for c in settings["courses"] if c["name"].casefold() == name.strip().casefold()), None)
        if existing:
            if patch:
                existing.update(apply_answer(existing))
            selected_code.append(existing["code"])
            return settings
        base = re.sub(r"[^A-Z0-9]", "", name.upper())[:12] or "COURSE"
        used = {c["code"] for c in settings["courses"]}
        used.update(alias.upper() for course in settings["courses"] for alias in course["aliases"])
        code, suffix = base, 2
        while code in used:
            code, suffix = f"{base}{suffix}", suffix + 1
        root = documents_root() / code
        material_dir = Path(clean_path(materials)) if materials else root / "Materials"
        note_dir = root / "Personal Notes"
        material_dir.mkdir(parents=True, exist_ok=True)
        note_dir.mkdir(parents=True, exist_ok=True)
        (root / "Handouts").mkdir(parents=True, exist_ok=True)
        course = normalize_course({"code": code, "name": name, "languages": {"lesson": "en", "notes": notes_language},
                                   "folders": {"materials": str(material_dir), "personal_notes": str(note_dir), "handout": str(root / "Handouts")}})
        course = apply_answer(course)
        tools = managed_tools()
        course["transcription"]["tools"].update({k: tools[k] for k in ("argmax", "model_cache") if isinstance(tools.get(k), str)})
        settings["courses"].append(course)
        if recordings:
            folder = clean_path(recordings)
            Path(folder).mkdir(parents=True, exist_ok=True)
            if folder not in settings["paths"]["recording_dirs"]:
                settings["paths"]["recording_dirs"].append(folder)
        completed = settings["onboarding"]["completed_groups"]
        if "courses" not in completed:
            completed.append("courses")
        remaining = [g for g in ("personal", "courses", "providers") if g not in completed]
        settings["onboarding"].update(status="in_progress" if remaining else "complete", next_group=remaining[0] if remaining else None)
        selected_code.append(code)
        return settings
    updated = workspace.configuration_store().update(edit)
    return next(course for course in updated["courses"] if course["code"] == selected_code[0])


def provider_check(provider: str | None) -> dict:
    if not provider:
        return {"status": "missing", "action": "Open MyNotes setup and choose an AI agent."}
    from .providers import provider_executable, provider_environment
    executable = provider_executable(provider)
    if not executable:
        return {"status": "missing", "provider": provider, "action": f"Install {provider}, sign in, then recheck.", "install_url": INSTALL_URLS.get(provider)}
    result = {"status": "unverified", "provider": provider, "executable": executable,
              "action": f"Open {provider} and sign in; then recheck. No paid generation was performed."}
    try:
        version = subprocess.run([executable, "--version"], capture_output=True, text=True, timeout=10, env=provider_environment())
        if version.returncode:
            return {**result, "status": "incompatible", "action": f"Update {provider} using its official installer.", "install_url": INSTALL_URLS.get(provider)}
        result["version"] = (version.stdout or version.stderr).strip()[:200]
        # Help probes verify the interfaces actually used by chat/generation.
        # A successful --version alone establishes neither API compatibility nor auth.
        probes = {
            "codex": [(["--help"], ("--add-dir", "--cd")),
                      (["exec", "--help"], ("--output-schema", "--output-last-message", "--json", "--ephemeral", "--sandbox", "--skip-git-repo-check"))],
            "claude": [(["--help"], ("--print", "--output-format", "--permission-mode", "--no-session-persistence", "--json-schema", "--append-system-prompt-file", "--add-dir"))],
            "agy": [(["--help"], ("--agent", "--mode", "--add-dir", "--prompt-interactive", "--print", "--output-format", "--json-schema"))],
        }
        missing = []
        result["compatibility_verified"] = False
        for arguments, required in probes.get(provider, []):
            help_result = subprocess.run([executable, *arguments], capture_output=True, text=True, timeout=10, env=provider_environment())
            help_text = help_result.stdout + help_result.stderr
            if help_result.returncode:
                missing.extend(required)
            else:
                for flag in required:
                    present = re.search(re.escape(flag) + r"(?=[\s=,\[<]|$)", help_text)
                    # Claude documents this supported hidden flag in the help
                    # text as a combined spelling, rather than its own row.
                    if provider == "claude" and flag == "--append-system-prompt-file":
                        present = present or re.search(r"--append-system-prompt\[-file\](?=[\s=,\[<]|$)", help_text)
                    if not present:
                        missing.append(flag)
        if missing:
            return {**result, "status": "incompatible", "missing_capabilities": sorted(set(missing)),
                    "action": f"Update {provider} using its official installer, then recheck.", "install_url": INSTALL_URLS.get(provider)}
        result["compatibility_verified"] = True
        command = {"codex": ["login", "status"], "claude": ["auth", "status"]}.get(provider)
        if command:
            auth = subprocess.run([executable, *command], capture_output=True, text=True, timeout=15, env=provider_environment())
            if provider == "claude":
                payload = json.loads(auth.stdout)
                if not isinstance(payload, dict) or type(payload.get("loggedIn")) is not bool:
                    return {**result, "authentication_error": "Agent returned an invalid authentication status.", "action": f"Open {provider}, check sign-in, then recheck. Authentication is unverified."}
                ready = auth.returncode == 0 and payload.get("loggedIn") is True
            else:
                ready = auth.returncode == 0 and "logged in" in (auth.stdout + auth.stderr).lower()
            result.update(status="ready" if ready else "login_required", action="Open the chat." if ready else f"Sign in using {provider}, then recheck.")
    except (OSError, subprocess.TimeoutExpired, ValueError):
        pass
    return result


def _check(ready: bool, action: str, **extra: Any) -> dict:
    return {"status": "ready" if ready else "missing", "action": action, **extra}


def doctor(workspace: Workspace, provider: str | None = None) -> dict:
    """Report readiness without changing settings or sending provider prompts."""
    tools, registry_error = _managed_registry()
    settings = {}
    config_error = None
    if workspace.is_initialized():
        try:
            settings = workspace.configuration_store().load()
        except Exception as error:
            config_error = str(error)
    argmax = managed_tool("argmax") or shutil.which("argmax-cli")
    model_cache = Path(tools.get("model_cache", workspace.root / "local-models"))
    model_ready = any(model_cache.glob("whisperkit/**/AudioEncoder.mlmodelc"))
    soffice = managed_tool("soffice") or shutil.which("soffice")
    proof = tools.get("execution_proof", {})
    asr_verified = bool(argmax and proof.get("transcription") is True and proof.get("model_cache") == str(model_cache) and proof.get("argmax_sha256") == file_sha256(Path(argmax)))
    word_verified = bool(soffice and proof.get("word") is True and proof.get("soffice_sha256") == file_sha256(Path(soffice)))
    checks = {
        "installation": _check(workspace.is_initialized() and config_error is None, "Open MyNotes setup.", workspace=str(workspace.root), configuration_error=config_error),
        "agent": provider_check(provider or settings.get("providers", {}).get("default")),
        "transcription": _check(bool(argmax and model_ready and asr_verified), "Open MyNotes setup to download local transcription components.", executable=argmax, model_present=model_ready, execution_verified=asr_verified),
        "word": _check(bool(soffice and word_verified), "Open MyNotes setup to install the Word renderer.", executable=soffice, execution_verified=word_verified),
        "capture": {"status": "optional", "action": "Configure Capture, record a microphone test, listen to it and confirm audibility."},
    }
    if registry_error:
        for component in ("transcription", "word"):
            checks[component].update(status="missing", execution_verified=False, registry_error=registry_error, action="Open MyNotes setup to repair its component registry. The original registry was preserved.")
    evidence = workspace.state_dir / "capture/audio-readiness.json"
    try:
        record = json.loads(evidence.read_text())
        if not isinstance(record, dict) or not isinstance(record.get("path"), str):
            raise ValueError("Invalid Capture microphone test evidence; run and confirm a new test.")
        audio = Path(record["path"])
        from .capture import app_path
        binary = app_path() / "Contents/MacOS/MyNotesCapture"
        if record.get("audible_confirmed") is True and record.get("validated") is True and audio.is_file() and record.get("sha256") == file_sha256(audio) and record.get("capture_binary_sha256") == file_sha256(binary):
            checks["capture"] = {"status": "ready", "action": "Capture microphone test verified and audibility confirmed.", "path": str(audio)}
    except (OSError, ValueError, KeyError, TypeError) as error:
        if evidence.exists():
            checks["capture"]["evidence_error"] = str(error)
    ready = all(checks[k]["status"] == "ready" for k in ("installation", "agent", "transcription", "word"))
    return {"ready": ready, "checks": checks, "capture_optional": True, "managed_registry_error": registry_error}


def confirm_capture_audio(workspace: Workspace, audio: Path, *, audible: bool) -> dict:
    """Store explicit listening acknowledgement only after native audio validation."""
    if not audible:
        raise ValueError("Listen to the microphone test before confirming audibility.")
    audio = audio.expanduser().resolve()
    sessions = workspace.state_dir / "capture/sessions"
    verified_session = False
    for metadata in sessions.glob("*.json"):
        try:
            record = json.loads(metadata.read_text())
            if record.get("isTest") is True and record.get("status") == "COMPLETED" and record.get("outputPath") and Path(record["outputPath"]).resolve() == audio:
                verified_session = True
                break
        except (OSError, ValueError):
            continue
    if not verified_session:
        raise ValueError("Select the completed microphone test from Capture status; ordinary recordings cannot confirm readiness.")
    checked = subprocess.run(["/usr/bin/afinfo", str(audio)], capture_output=True, text=True, timeout=15)
    if checked.returncode or not audio.is_file() or audio.stat().st_size == 0:
        raise ValueError("Microphone test is missing or is not valid audio.")
    from .capture import app_path
    binary = app_path() / "Contents/MacOS/MyNotesCapture"
    evidence = {"path": str(audio), "validated": True, "audible_confirmed": True, "sha256": file_sha256(audio), "capture_binary_sha256": file_sha256(binary)}
    atomic_json_write(workspace.state_dir / "capture/audio-readiness.json", evidence)
    return evidence


def setup_cli(workspace: Workspace, provider: str | None = None) -> int:
    settings = prepare_conversation(workspace, provider)
    selected = provider or settings["providers"]["default"]
    native_env = ["--env", f"MYNOTES_PROVIDER={selected}"] if selected else []
    try:
        subprocess.run(["open", "-a", os.environ.get("MYNOTES_SETUP_APP", "/Applications/MyNotes.app"), "--env", f"MYNOTES_WORKSPACE={workspace.root}", "--env", f"MYNOTES_HOME={runtime_home()}", *native_env], check=True)
    except subprocess.CalledProcessError as error:
        raise ValueError("Install MyNotes.app in Applications, then open it to complete setup.") from error
    return 0



def verify_local_components(workspace: Workspace) -> dict:
    """Run real local ASR and Word conversion on generated public test inputs."""
    from docx import Document
    from .asr import ArgmaxEngine, normalize_audio
    from .publication import _render_and_validate
    tools = managed_tools()
    argmax, soffice = managed_tool("argmax"), managed_tool("soffice")
    if not argmax or not soffice or not tools.get("model_cache"):
        raise ValueError("Setup components are missing; retry component installation.")
    signature = {"argmax_sha256": file_sha256(Path(argmax)), "soffice_sha256": file_sha256(Path(soffice)), "model_cache": tools["model_cache"]}
    proof = tools.get("execution_proof", {})
    if all(proof.get(k) == v for k, v in signature.items()) and proof.get("transcription") is True and proof.get("word") is True:
        return proof
    check_root = runtime_home() / "checks"
    check_root.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=check_root, prefix="local-smoke-") as temporary:
        root = Path(temporary)
        subprocess.run(["/usr/bin/say", "-o", str(root / "spoken.aiff"), "My Notes records lectures and prepares clear study notes."], check=True, capture_output=True, timeout=30)
        audio = normalize_audio(root / "spoken.aiff", root / "spoken.wav")
        cache = Path(tools["model_cache"])
        engine = ArgmaxEngine(argmax, cache / "whisperkit", cache / "tokenizers")
        result = engine.transcribe(audio, root / "transcript", model="large-v3-v20240930_626MB", language="en", vocabulary=[], prompt="")
        if not result.segments or not any(segment.get("text", "").strip() for segment in result.segments):
            raise ValueError("Local transcription test produced no text; retry setup.")
        document = Document()
        document.add_heading("MyNotes setup test", 0)
        document.add_paragraph("Your local Word renderer is working.")
        document.save(root / "check.docx")
        rendered = _render_and_validate(root / "check.docx", root / "render")
        proof = {**signature, "transcription": True, "word": True, "asr_engine_version": result.version, "word_pages": len(rendered["pages"]), "test_inputs": "system speech and generated Word; no personal sources"}
    tools["execution_proof"] = proof
    atomic_json_write(runtime_home() / "tools.json", tools)
    return proof


# Alias retained for GUI/CLI integrations.
doctor_report = doctor
