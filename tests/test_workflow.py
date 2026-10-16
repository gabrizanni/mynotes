from __future__ import annotations

import sqlite3
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from mynotes.cli import main
from mynotes.configuration import course_template
from mynotes.providers import ProviderQuotaError, ProviderRequest, ProviderResponse
from mynotes.utils import file_sha256
from mynotes.workflow import WorkflowError, _publication_lock, normalize_quality, process, resolve_model, status
from mynotes.workspace import Workspace


class FakeProvider:
    """Test-only provider. Production MyNotes does not ship a fallback generator."""

    def generate(self, request: ProviderRequest) -> ProviderResponse:
        transcript = request.sources["transcript"].strip()
        personal = request.sources["personal_notes"].strip()
        formula = None
        if "FORMULA:" in transcript:
            transcript, formula = transcript.split("FORMULA:", 1)
            formula = formula.splitlines()[0].strip()
        combined = " ".join(item for item in (transcript.strip(), personal) if item)
        notes = {
            "title": "Provider-generated test notes",
            "summary": combined or request.sources["materials"],
            "bullets": ["This bullet was returned by the test provider."],
            "formula": formula,
            "table": [["Source", "Available"], ["Transcript", "yes"]]
            if request.course["kind"] == "quantitative"
            else [],
        }
        return ProviderResponse(
            notes=notes,
            provider=request.requested.provider,
            model=request.requested.model,
            effort=request.requested.effort,
            usage={"total_tokens": 123},
        )


class QuotaProvider:
    def generate(self, request: ProviderRequest) -> ProviderResponse:
        raise ProviderQuotaError("test allowance exhausted")


FAKE_PROVIDER = FakeProvider()


def configured_workspace(tmp_path) -> tuple[Workspace, dict]:
    workspace = Workspace.from_value(tmp_path / "private")
    workspace.initialize()
    configuration = workspace.configuration_store().load()
    quantitative = course_template()
    quantitative.update({"code": "STAT1", "name": "Statistics", "kind": "quantitative"})
    quantitative["folders"]["handout"] = str(tmp_path / "handouts")
    discursive = course_template()
    discursive.update({"code": "HIST2", "name": "History", "kind": "theoretical"})
    discursive["folders"]["handout"] = str(tmp_path / "handouts")
    configuration["courses"] = [quantitative, discursive]
    configuration["providers"]["default"] = "codex"
    configuration["providers"]["configured"] = {
        "codex": {"model": "test-model", "effort": "medium"}
    }
    workspace.configuration_store().save(configuration)
    return workspace, configuration


def test_two_new_courses_publish_and_identical_second_run_skips(tmp_path) -> None:
    workspace, _ = configured_workspace(tmp_path)
    quantitative = tmp_path / "stat.txt"
    quantitative.write_text(
        "A line has slope m. FORMULA: y = m x + b\nThe table compares observed values. The graph shows the trend.",
        encoding="utf-8",
    )
    discursive = tmp_path / "history.txt"
    discursive.write_text("The source describes the institutional change. The evidence is incomplete.", encoding="utf-8")

    first = process(workspace, "STAT1-L01", quantitative, provider_adapter=FAKE_PROVIDER)
    second = process(workspace, "HIST2-L01", discursive, provider_adapter=FAKE_PROVIDER)
    repeated = process(workspace, "STAT1-L01", quantitative, provider_adapter=FAKE_PROVIDER)
    assert first["status"] == second["status"] == "PUBLISHED"
    assert repeated["status"] == "SKIPPED"


def test_redo_replaces_only_target_lesson_and_preserves_other_content(tmp_path) -> None:
    workspace, _ = configured_workspace(tmp_path)
    first = tmp_path / "first.txt"
    second = tmp_path / "second.txt"
    first.write_text("Initial explanation. FORMULA: a = b", encoding="utf-8")
    second.write_text("Second lesson remains in place.", encoding="utf-8")
    output = process(workspace, "STAT1-L01", first, provider_adapter=FAKE_PROVIDER)["output"]
    process(workspace, "STAT1-L02", second, provider_adapter=FAKE_PROVIDER)
    document = Document(output)
    document.add_paragraph("Manual note outside lesson boundaries.")
    document.save(output)
    first.write_text("Revised explanation with clearer steps. FORMULA: c = d", encoding="utf-8")

    result = process(
        workspace,
        "STAT1-L01",
        first,
        redo=True,
        accept_current_word=True,
        provider_adapter=FAKE_PROVIDER,
    )
    rendered = Document(result["output"])
    visible = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert "Revised explanation" in visible
    assert "Second lesson remains" in visible
    assert "Manual note outside lesson boundaries" in visible
    assert visible.index("MYNOTES START STAT1-L01") < visible.index("MYNOTES START STAT1-L02")


def test_test_provider_output_uses_native_headings_lists_formula_and_table(tmp_path) -> None:
    workspace, _ = configured_workspace(tmp_path)
    transcript = tmp_path / "source.txt"
    transcript.write_text(
        "An input produces an output. FORMULA: y = m x + b\nA graph connects the two quantities.",
        encoding="utf-8",
    )
    output = process(workspace, "STAT1-L01", transcript, provider_adapter=FAKE_PROVIDER)["output"]
    document = Document(output)
    assert any(paragraph.style.name == "Heading 2" for paragraph in document.paragraphs)
    assert any(paragraph._p.find(".//" + qn("w:numPr")) is not None for paragraph in document.paragraphs)
    assert document.styles["MyNotes Body"].font.name == "Aptos"
    assert document.styles["MyNotes Body"].font.size.pt == 11
    assert document.styles["MyNotes Body"].paragraph_format.line_spacing == 1.15
    assert all(abs(section.left_margin.cm - 1.5) < .01 for section in document.sections)
    assert "m:oMath" in document.part.element.xml
    assert not document.inline_shapes
    assert document.tables


def test_quality_aliases_and_actual_resolved_values_are_stable(tmp_path) -> None:
    workspace, configuration = configured_workspace(tmp_path)
    configuration["providers"]["configured"] = {"agy": {"model": "gemini-3.8-flash-high", "effort": "high"}}
    configuration["providers"]["default"] = "agy"
    workspace.configuration_store().save(configuration)
    resolved = resolve_model(workspace.configuration_store().load(), None, "scoped-review")
    assert normalize_quality("fast") == "base"
    assert resolved.quality == "approfondito"
    assert resolved.model == "gemini-3.8-flash-high"


def test_external_word_change_is_blocked_until_explicitly_accepted(tmp_path) -> None:
    workspace, _ = configured_workspace(tmp_path)
    transcript = tmp_path / "source.txt"
    transcript.write_text("Initial explanation.", encoding="utf-8")
    output = process(workspace, "STAT1-L01", transcript, provider_adapter=FAKE_PROVIDER)["output"]
    document = Document(output)
    document.add_paragraph("External edit that must not be lost.")
    document.save(output)
    external_hash = file_sha256(Path(output))
    transcript.write_text("Revised explanation.", encoding="utf-8")

    with pytest.raises(WorkflowError, match="changed after"):
        process(workspace, "STAT1-L01", transcript, redo=True, provider_adapter=FAKE_PROVIDER)
    assert file_sha256(Path(output)) == external_hash

    result = process(
        workspace,
        "STAT1-L01",
        transcript,
        redo=True,
        accept_current_word=True,
        provider_adapter=FAKE_PROVIDER,
    )
    visible = "\n".join(paragraph.text for paragraph in Document(result["output"]).paragraphs)
    assert "External edit that must not be lost" in visible
    assert "Revised explanation" in visible


def test_publication_lock_rejects_a_second_writer(tmp_path) -> None:
    workspace, configuration = configured_workspace(tmp_path)
    transcript = tmp_path / "source.txt"
    transcript.write_text("One controlled input.", encoding="utf-8")
    target = tmp_path / "handouts" / f"{configuration['courses'][0]['code']}_COURSE_NOTES.docx"
    with _publication_lock(target):
        with pytest.raises(WorkflowError, match="already active"):
            process(workspace, "STAT1-L01", transcript, provider_adapter=FAKE_PROVIDER)
    assert not target.exists()


def test_provider_quota_pause_is_saved_and_recoverable(tmp_path) -> None:
    workspace, _ = configured_workspace(tmp_path)
    transcript = tmp_path / "source.txt"
    transcript.write_text("A resumable lesson.", encoding="utf-8")
    with pytest.raises(WorkflowError, match="PAUSED_QUOTA"):
        process(workspace, "STAT1-L01", transcript, provider_adapter=QuotaProvider())
    assert status(workspace)[0]["status"] == "PAUSED_QUOTA"
    assert process(workspace, "STAT1-L01", transcript, provider_adapter=FAKE_PROVIDER)["status"] == "PUBLISHED"


def test_unavailable_cloud_source_fails_without_a_published_state(tmp_path) -> None:
    workspace, _ = configured_workspace(tmp_path)
    missing = tmp_path / "cloud placeholder" / "missing.txt"
    with pytest.raises(WorkflowError, match="Cannot read text source"):
        process(workspace, "STAT1-L01", missing, provider_adapter=FAKE_PROVIDER)
    assert status(workspace) == []


def test_no_institution_and_incomplete_optional_sources_are_explicit(tmp_path) -> None:
    workspace, _ = configured_workspace(tmp_path)
    assert workspace.configuration_store().load()["personal"]["institution"] is None
    transcript = tmp_path / "source.txt"
    web = tmp_path / "online source snapshot.txt"
    transcript.write_text("A discursive explanation with controlled evidence.", encoding="utf-8")
    web.write_text("Public reference snapshot supplied by the user.", encoding="utf-8")
    result = process(workspace, "HIST2-L01", transcript, web_source=web, provider_adapter=FAKE_PROVIDER)
    assert result["sources"] == {
        "transcript": True,
        "materials": False,
        "personal_notes": False,
        "web": True,
    }


def test_unavailable_provider_never_publishes_or_claims_invocation(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr("mynotes.providers.shutil.which", lambda _, **kwargs: None)
    workspace, _ = configured_workspace(tmp_path)
    transcript = tmp_path / "source.txt"
    transcript.write_text("Content that must not be turned into notes by a fallback.", encoding="utf-8")

    with pytest.raises(WorkflowError, match="Codex CLI is not installed"):
        process(workspace, "STAT1-L01", transcript)

    assert status(workspace)[0]["status"] == "PROVIDER_UNAVAILABLE"
    assert not list((tmp_path / "handouts").glob("*.docx"))
    assert "PROVIDER_UNAVAILABLE" in workspace.status_path.read_text(encoding="utf-8")
    with sqlite3.connect(workspace.database_path) as connection:
        attempt = connection.execute(
            "SELECT requested_provider,actual_provider,status FROM provider_attempts"
        ).fetchone()
    assert attempt == ("codex", None, "UNAVAILABLE")


def test_normal_cli_has_no_simulated_provider_fallback(tmp_path, capsys, monkeypatch) -> None:
    monkeypatch.setattr("mynotes.providers.shutil.which", lambda _, **kwargs: None)
    workspace, _ = configured_workspace(tmp_path)
    transcript = tmp_path / "source.txt"
    transcript.write_text("This must not be generated by a hidden fallback.", encoding="utf-8")
    exit_code = main(
        [
            "process",
            "STAT1-L01",
            "--workspace",
            str(workspace.root),
            "--transcript",
            str(transcript),
        ]
    )
    assert exit_code == 2
    assert "Codex CLI is not installed" in capsys.readouterr().out
    assert not list((tmp_path / "handouts").glob("*.docx"))


def test_provider_attempt_records_requested_and_actual_values_and_usage(tmp_path) -> None:
    workspace, _ = configured_workspace(tmp_path)
    transcript = tmp_path / "source.txt"
    transcript.write_text("Provider-backed content.", encoding="utf-8")
    result = process(workspace, "STAT1-L01", transcript, provider_adapter=FAKE_PROVIDER)
    assert result["requested_model"] == {
        "provider": "codex",
        "model": "test-model",
        "effort": "medium",
        "quality": "standard",
    }
    assert result["actual_model"] == {
        "provider": "codex",
        "model": "test-model",
        "effort": "medium",
    }
    assert result["usage"] == {"total_tokens": 123}
    assert "PUBLISHED" in workspace.status_path.read_text(encoding="utf-8")
