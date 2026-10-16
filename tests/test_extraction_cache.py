from __future__ import annotations

import json
import threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest

from mynotes import extraction, extraction_cache
from mynotes.sourcepack import build_source_pack
from mynotes.configuration import course_template
from mynotes.workspace import Workspace


@pytest.fixture
def fake_extract(monkeypatch):
    calls = []
    def extract(source, *, visual_dir=None, lesson_number=None):
        calls.append((str(source), lesson_number, visual_dir is not None))
        chunks = [{"content": source.read_text(), "page": lesson_number or 1}]
        visuals = []
        if visual_dir is not None:
            visual_dir.mkdir(parents=True, exist_ok=True)
            for page in (1, 2):
                image = visual_dir / f"page-{page:03d}.png"
                image.write_bytes(b"identical image fixture")
                visuals.append({"page": page, "source_file": str(source), "image_path": str(image), "visual_decision": "candidate"})
        return chunks, visuals
    monkeypatch.setattr(extraction, "extract_material", extract)
    return calls


def test_repeated_content_is_extracted_once_and_original_unchanged(tmp_path, fake_extract):
    source = tmp_path / "lecture.md"
    source.write_text("Original evidence")
    cache = tmp_path / "cache"
    first = extraction_cache.extract_cached(source, cache)
    second = extraction_cache.extract_cached(source, cache)
    assert first == second
    assert len(fake_extract) == 1
    assert source.read_text() == "Original evidence"
    source.write_text("Changed evidence")
    assert extraction_cache.extract_cached(source, cache)[0][0]["content"] == "Changed evidence"
    assert len(fake_extract) == 2


@pytest.mark.parametrize("damage", ["bad_json", "payload_corruption", "missing_visual", "corrupt_visual"])
def test_damaged_cache_rebuilds_extraction_safely(tmp_path, fake_extract, damage):
    source = tmp_path / "lecture.md"
    source.write_text("Evidence")
    cache, visuals = tmp_path / "cache", tmp_path / "visuals"
    expected = extraction_cache.extract_cached(source, cache, visual_dir=visuals)
    result = next(cache.glob("*/result.json"))
    envelope = json.loads(result.read_text())
    if damage == "bad_json":
        result.write_text("{incomplete")
    elif damage == "payload_corruption":
        envelope["data"]["chunks"][0]["content"] = "Corrupted cache evidence"
        result.write_text(json.dumps(envelope))
    else:
        asset = result.parent / next(iter(envelope["data"]["assets"]))
        if damage == "missing_visual": asset.unlink()
        else: asset.write_bytes(b"corrupted image")
    assert extraction_cache.extract_cached(source, cache, visual_dir=visuals) == expected
    assert len(fake_extract) == 2
    assert source.read_text() == "Evidence"


def test_valid_cached_visual_restores_missing_or_corrupt_run_asset_without_extraction(tmp_path, fake_extract):
    source = tmp_path / "lecture.md"
    source.write_text("Evidence")
    cache = tmp_path / "cache"
    first = extraction_cache.extract_cached(source, cache, visual_dir=tmp_path / "run-a")
    second = extraction_cache.extract_cached(source, cache, visual_dir=tmp_path / "run-b")
    assert len(fake_extract) == 1
    assert [Path(v["image_path"]).name for v in second[1]] == ["page-001.png", "page-002.png"]
    Path(second[1][0]["image_path"]).unlink()
    Path(second[1][1]["image_path"]).write_bytes(b"corrupt")
    restored = extraction_cache.extract_cached(source, cache, visual_dir=tmp_path / "run-b")
    assert len(fake_extract) == 1
    assert all(Path(v["image_path"]).read_bytes() == b"identical image fixture" for v in restored[1])
    assert all(v["source_file"] == str(source) for v in restored[1])


def test_options_and_extractor_identity_changes_invalidate_cache(tmp_path, fake_extract, monkeypatch):
    source = tmp_path / "lecture.docx"
    source.write_text("Evidence")
    cache = tmp_path / "cache"
    extraction_cache.extract_cached(source, cache, lesson_number=1)
    extraction_cache.extract_cached(source, cache, lesson_number=2)
    extraction_cache.extract_cached(source, cache, lesson_number=2, visual_dir=tmp_path / "visuals")
    assert len(fake_extract) == 3
    original = extraction_cache.extractor_identity
    monkeypatch.setattr(extraction_cache, "extractor_identity", lambda *args,**kwargs: {**original(*args,**kwargs), "new_dependency_version": "next"})
    extraction_cache.extract_cached(source, cache, lesson_number=2, visual_dir=tmp_path / "visuals")
    assert len(fake_extract) == 4


def test_renderer_binary_change_invalidates_cache(tmp_path, fake_extract, monkeypatch):
    source = tmp_path / "lecture.pptx"
    source.write_text("Evidence")
    renderer = tmp_path / "soffice"
    renderer.write_bytes(b"binary version one")
    monkeypatch.setenv("MYNOTES_LIBREOFFICE", str(renderer))
    cache = tmp_path / "cache"
    extraction_cache.extract_cached(source, cache, visual_dir=tmp_path / "visuals")
    extraction_cache.extract_cached(source, cache, visual_dir=tmp_path / "visuals")
    assert len(fake_extract) == 1
    renderer.write_bytes(b"binary version two")
    extraction_cache.extract_cached(source, cache, visual_dir=tmp_path / "visuals")
    assert len(fake_extract) == 2


def test_concurrent_writers_extract_once_and_rebind_same_content_provenance(tmp_path, fake_extract):
    source = tmp_path / "lecture.md"
    source.write_text("Same evidence")
    other = tmp_path / "other.md"
    other.write_text(source.read_text())
    gate = threading.Barrier(2)
    def run(path):
        gate.wait(timeout=5)
        return extraction_cache.extract_cached(path, tmp_path / "cache", visual_dir=tmp_path / path.stem)
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(run, (source, other)))
    assert len(fake_extract) == 1
    for path, result in zip((source, other), results):
        assert all(v["source_file"] == str(path) for v in result[1])
        assert all(Path(v["image_path"]).parent == tmp_path / path.stem for v in result[1])


def test_source_pack_semantic_bytes_stay_identical_across_cache_hit(tmp_path, monkeypatch):
    workspace = Workspace(tmp_path / "workspace")
    workspace.initialize()
    config = workspace.configuration_store().load()
    course = course_template()
    course.update(code="OPS", name="Operations")
    config["courses"] = [course]
    workspace.configuration_store().save(config)
    notes = tmp_path / "notes.md"
    notes.write_text("Inventory trade-off.\nNOTE AI: Explain clearly")
    actual = extraction.extract_material
    calls = []
    def counted(*args, **kwargs):
        calls.append(args[0])
        return actual(*args, **kwargs)
    monkeypatch.setattr(extraction, "extract_material", counted)
    first = build_source_pack(workspace, "OPS-L01", personal_notes=notes)
    snapshot = (workspace.runs_dir / "OPS-L01/source_pack.json").read_bytes()
    second = build_source_pack(workspace, "OPS-L01", personal_notes=notes)
    assert len(calls) == 1
    assert (workspace.runs_dir / "OPS-L01/source_pack.json").read_bytes() == snapshot
    assert first["evidence"] == second["evidence"]
    assert first["directives"] == second["directives"]
    assert notes.read_text() == "Inventory trade-off.\nNOTE AI: Explain clearly"


def _workspace(tmp_path):
    workspace = Workspace(tmp_path / "workspace")
    workspace.initialize()
    config = workspace.configuration_store().load()
    course = course_template()
    course.update(code="OPS", name="Operations")
    config["courses"] = [course]
    workspace.configuration_store().save(config)
    return workspace


def test_identical_visual_bytes_keep_distinct_names_and_sourcepack_snapshot(tmp_path, fake_extract):
    workspace = _workspace(tmp_path)
    material = tmp_path / "slides.md"
    material.write_text("Evidence")
    first = build_source_pack(workspace, "OPS-L01", materials=[material])
    names = [Path(v["image_path"]).name for v in first["visuals"]]
    assert names == ["page-001.png", "page-002.png"]
    assert len({v["image_path"] for v in first["visuals"]}) == 2
    snapshot = (workspace.runs_dir / "OPS-L01/source_pack.json").read_bytes()
    second = build_source_pack(workspace, "OPS-L01", materials=[material])
    assert second["visuals"] == first["visuals"]
    assert (workspace.runs_dir / "OPS-L01/source_pack.json").read_bytes() == snapshot
    assert len(fake_extract) == 1


def test_cache_implementation_identity_change_invalidates_entry(tmp_path, fake_extract, monkeypatch):
    implementation = tmp_path / "cache-implementation.py"
    implementation.write_text("implementation revision one")
    monkeypatch.setattr(extraction_cache, "__file__", str(implementation))
    source = tmp_path / "source.md"
    source.write_text("Evidence")
    extraction_cache.extract_cached(source, tmp_path / "cache")
    extraction_cache.extract_cached(source, tmp_path / "cache")
    assert len(fake_extract) == 1
    implementation.write_text("implementation revision two")
    extraction_cache.extract_cached(source, tmp_path / "cache")
    assert len(fake_extract) == 2


def test_source_changed_while_waiting_for_key_lock_is_rejected(tmp_path, fake_extract, monkeypatch):
    import fcntl
    source = tmp_path / "source.md"
    source.write_text("Old evidence")
    cache = tmp_path / "cache"
    extraction_cache.extract_cached(source, cache)
    result = next(cache.glob("*/result.json"))
    lock = cache / (result.parent.name + ".lock")
    waiting = threading.Event()
    actual_flock = fcntl.flock
    def observed(fd, operation):
        if threading.current_thread() is not threading.main_thread():
            waiting.set()
        return actual_flock(fd, operation)
    with lock.open("a+") as held:
        actual_flock(held.fileno(), fcntl.LOCK_EX)
        monkeypatch.setattr(fcntl, "flock", observed)
        with ThreadPoolExecutor(max_workers=1) as pool:
            future = pool.submit(extraction_cache.extract_cached, source, cache)
            try:
                assert waiting.wait(timeout=5)
                source.write_text("New evidence")
            finally:
                actual_flock(held.fileno(), fcntl.LOCK_UN)
            with pytest.raises(extraction.ExtractionError, match="waiting"):
                future.result(timeout=5)
    assert len(fake_extract) == 1
    assert json.loads(result.read_text())["data"]["chunks"][0]["content"] == "Old evidence"


@pytest.mark.parametrize("role", ["material", "notes", "web"])
def test_changed_source_descriptor_rejects_before_replacing_existing_pack(tmp_path, monkeypatch, role):
    import mynotes.sourcepack as sourcepack
    workspace = _workspace(tmp_path)
    stable = tmp_path / "stable.md"
    stable.write_text("Stable lecture evidence")
    build_source_pack(workspace, "OPS-L01", materials=[stable])
    pack = workspace.runs_dir / "OPS-L01/source_pack.json"
    original_pack = pack.read_bytes()
    changing = tmp_path / "changing.md"
    changing.write_text("Old source")
    def edit_after_extract(path, *args, **kwargs):
        old = path.read_text()
        path.write_text("Externally changed source")
        return [{"content": old}], []
    monkeypatch.setattr(sourcepack, "extract_cached", edit_after_extract)
    options = {"materials": [changing]} if role == "material" else {"personal_notes": changing} if role == "notes" else {"web_snapshots": [changing]}
    with pytest.raises(sourcepack.SourcePackError, match="Source changed"):
        build_source_pack(workspace, "OPS-L01", **options)
    assert pack.read_bytes() == original_pack


def test_real_pdf_docx_repeat_prepare_reuses_extraction_preserving_sources_and_visuals(tmp_path, monkeypatch):
    """Real local PDF rendering and DOCX extraction, with no ASR or AI."""
    from docx import Document
    from test_ingestion_sourcepack import _create_test_pdf
    from mynotes.utils import file_sha256
    workspace = _workspace(tmp_path)
    pdf, notes = tmp_path / "slides.pdf", tmp_path / "notes.docx"
    _create_test_pdf(pdf, text="Inventory trade-off", draw_paths=4)
    document = Document()
    document.add_heading("Lecture 1", level=2)
    document.add_paragraph("Inventory balances service and cost.")
    document.add_paragraph("NOTE AI: Keep this imported note as evidence.")
    document.save(notes)
    source_hashes = {p: file_sha256(p) for p in (pdf, notes)}
    actual = extraction.extract_material
    calls = []
    def counted(*args, **kwargs):
        calls.append(args[0])
        return actual(*args, **kwargs)
    monkeypatch.setattr(extraction, "extract_material", counted)
    first = build_source_pack(workspace, "OPS-L01", materials=[pdf], personal_notes=notes)
    pack = workspace.runs_dir / "OPS-L01/source_pack.json"
    snapshot = pack.read_bytes()
    asset_hashes = {v["image_path"]: file_sha256(Path(v["image_path"])) for v in first["visuals"] if v.get("image_path")}
    assert asset_hashes and len(calls) == 2
    second = build_source_pack(workspace, "OPS-L01", materials=[pdf], personal_notes=notes)
    assert len(calls) == 2
    assert pack.read_bytes() == snapshot
    assert first["evidence"] == second["evidence"] and first["visuals"] == second["visuals"]
    assert {p: file_sha256(p) for p in (pdf, notes)} == source_hashes
    assert all(file_sha256(Path(path)) == checksum for path, checksum in asset_hashes.items())
