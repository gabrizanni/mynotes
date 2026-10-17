from copy import deepcopy
import json
from pathlib import Path

import pytest

from mynotes.batch import BatchBudget, generate_batch, plan_batch
from mynotes.configuration import course_template
from mynotes.providers import ProviderQuotaError, ProviderResponse, _context
from mynotes.sourcepack import build_source_pack
from mynotes.utils import atomic_json_write, file_sha256
from mynotes.workflow import WorkflowError
from mynotes.workspace import Workspace


def workspace_and_deck(tmp_path, lessons=2, slides=30):
    workspace = Workspace.from_value(tmp_path / "private")
    workspace.initialize()
    config = workspace.configuration_store().load()
    course = course_template()
    course.update(code="OPS", name="Operations")
    config["courses"] = [course]
    config["providers"]["default"] = "codex"
    config["providers"]["configured"] = {"codex": {"model": "test-model", "effort": "high"}}
    workspace.configuration_store().save(config)
    deck = tmp_path / "deck.pdf"
    deck.write_bytes(b"synthetic shared deck")
    digest = file_sha256(deck)
    for number in range(1, lessons + 1):
        lecture = f"OPS-L{number:02d}"
        evidence = [{"id": f"S{page:03d}", "kind": "slide", "page": page, "source_file": str(deck),
                     "content": f"Distinct topic {page} with source supported content."} for page in range(1, slides + 1)]
        pack = {"schema_version": 1, "lecture_id": lecture, "course": {"code": "OPS", "name": "Operations", "lesson_number": number},
                "availability": {"materials": True, "transcript": False}, "evidence": evidence,
                "visuals": [], "directives": [], "sources": [{"path": str(deck), "sha256": digest}]}
        atomic_json_write(workspace.runs_dir / lecture / "source_pack.json", pack)
    selection = {f"sha256:{digest}:page:{page}": {"owner": "OPS-L01" if page <= 20 else "OPS-L02", "status": "taught"} for page in range(1, slides + 1)}
    return workspace, deck, selection


def notes_for(request, lecture):
    own = [item for item in request.batch_context["ownership"] if item["owner"] == lecture]
    visible = request.batch_context["lesson_evidence"][lecture]
    first = own[0]["evidence_id"]
    paragraphs = [{"kind": "paragraph", "text": f"Lecture {lecture} explains source slide {item['evidence_id']} through its distinct source supported mechanism.",
                   "evidence_ids": [item["evidence_id"]], "coverage_role": "supplement" if item["status"] == "supplement" else "full"} for item in own]
    # A shuffled supplementary explanation must be placed between its anchors.
    paragraphs.sort(key=lambda block: int(block["evidence_ids"][0][1:]))
    return {"title": lecture, "summary": "Study material", "bullets": [], "table": [], "formula": None,
            "formulas": [], "sections": [], "evidence_coverage": [{"evidence_id": identifier, "disposition": "used" if identifier in [item["evidence_id"] for item in own] else "excluded", "reason": "Owned material or explained in its other owner lesson"} for identifier in visible],
            "visual_decisions": [{"evidence_id": item["evidence_id"], "decision": "exclude", "reason": "Prose is sufficient", "inspected": False} for item in request.source_pack["visuals"] if item["evidence_id"] in visible],
            "model_knowledge": [], "review_questions": [], "blocks": [
                {"kind": "overview", "items": ["Source concepts"], "evidence_ids": [first]},
                *paragraphs,
                {"kind": "qa", "question": "How does the mechanism work?", "answer": "The source describes the mechanism.", "evidence_ids": [first]},
                {"kind": "recap", "text": f"Distinct recap for {lecture}.", "evidence_ids": [first]}]}


class Adapter:
    def __init__(self, transform=None, fail_at=None):
        self.calls = []
        self.transform = transform
        self.fail_at = fail_at

    def generate(self, request):
        self.calls.append(request)
        if self.fail_at == len(self.calls):
            raise ProviderQuotaError("synthetic quota")
        response = {"lessons": [{"lecture_id": lecture, "notes": notes_for(request, lecture)} for lecture in request.batch_context["lessons"]]}
        if self.transform:
            self.transform(response, request)
        return ProviderResponse(response, "codex", "test-model", "high", {"input_tokens": 100, "output_tokens": 50}, f"attempt-{len(self.calls)}")


VERIFIED = BatchBudget(verified_for_model="test-model", context_tokens=1_000_000, output_tokens=100_000, output_reserve_per_lesson=4000)


def test_structural_slide_exemption_requires_trusted_image_bound_review(tmp_path):
    from PIL import Image
    workspace, deck, selection = workspace_and_deck(tmp_path)
    path = workspace.runs_dir / "OPS-L01/source_pack.json"
    pack = json.loads(path.read_text())
    image = path.parent / "title.png"
    Image.new("RGB", (20, 20), "white").save(image)
    pack["visuals"] = [{"evidence_id": "S001", "source_file": str(deck), "page": 1,
                        "image_path": str(image), "reason": "Title slide reviewed in source planning"}]
    atomic_json_write(path, pack)
    key = next(iter(selection))

    def omit(response, request):
        notes = response["lessons"][0]["notes"]
        notes["blocks"] = [b for b in notes["blocks"] if not (b.get("coverage_role") == "full" and b["evidence_ids"] == ["S001"])]
        notes["evidence_coverage"][0].update(disposition="excluded", reason="Reviewed title")

    # Neither a model exclusion nor a heading/overview proves source coverage.
    with pytest.raises(WorkflowError, match="owned slide"):
        generate_batch(workspace, "OPS", [1, 2], selection=selection, provider_adapter=Adapter(omit))
    selection[key] = {"owner": "OPS-L01", "status": "material", "structural_review": {
        "kind": "title", "reason": "Only course title and decorative artwork, no teaching content",
        "image_sha256": "0" * 64}}
    with pytest.raises(WorkflowError, match="inspected source image"):
        plan_batch(workspace, "OPS", [1, 2], selection=selection)
    selection[key]["structural_review"]["image_sha256"] = file_sha256(image)
    # A changed trusted plan has its own checkpoint; it may reuse the saved
    # transport only after local review, never an automatic provider retry.
    result = generate_batch(workspace, "OPS", [1, 2], selection=selection, provider_adapter=Adapter(omit))
    candidate = json.loads(Path(result["candidates"][0]["candidate"]).read_text())
    assert candidate["notes"]["evidence_coverage"][0]["disposition"] == "excluded"
    assert candidate["slide_ownership"][key]["structural_review"]["image_sha256"] == file_sha256(image)
    Image.new("RGB", (20, 20), "black").save(image)
    with pytest.raises(WorkflowError, match="inspected source image"):
        plan_batch(workspace, "OPS", [1, 2], selection=selection)


def test_plan_only_is_read_only_and_shared_ownership_needs_selection(tmp_path):
    workspace, deck, selection = workspace_and_deck(tmp_path)
    before = sorted(str(path.relative_to(workspace.root)) for path in workspace.root.rglob("*"))
    ambiguous = generate_batch(workspace, "OPS", [2, 1], plan_only=True)
    assert ambiguous["status"] == "NEEDS_SELECTION" and len(ambiguous["unresolved"]) == 30
    assert sorted(str(path.relative_to(workspace.root)) for path in workspace.root.rglob("*")) == before
    plan = plan_batch(workspace, "OPS", [2, 1, 1], selection=selection, budgets=VERIFIED)
    assert plan["lectures"] == ["OPS-L01", "OPS-L02"] and plan["groups"] == [["OPS-L01", "OPS-L02"]]
    assert deck.read_bytes() == b"synthetic shared deck"


def test_completed_manifest_matches_returned_ready_state(tmp_path):
    from mynotes.batch import load_batch_manifest
    workspace, _, selection = workspace_and_deck(tmp_path)
    result = generate_batch(workspace, "OPS", [1, 2], selection=selection, provider_adapter=Adapter())
    saved = load_batch_manifest(workspace, result["batch_id"])
    assert saved["status"] == result["status"] == "CANDIDATE_READY"
    assert saved["completed_groups"] == 2 and saved["candidates"] == result["candidates"]


def test_unified_shared_thirty_slide_deck_has_one_full_explanation_per_slide(tmp_path):
    workspace, deck, selection = workspace_and_deck(tmp_path)
    adapter = Adapter()
    result = generate_batch(workspace, "OPS", [2, 1], selection=selection, budgets=VERIFIED, provider_adapter=adapter)
    assert result["status"] == "CANDIDATE_READY" and len(adapter.calls) == 1
    assert len(adapter.calls[0].source_pack["evidence"]) == 30
    assert len(result["shared_attempts"]) == 1
    counts = []
    for entry in result["candidates"]:
        payload = json.loads(Path(entry["candidate"]).read_text())
        counts.append(sum(block["kind"] == "paragraph" for block in payload["notes"]["blocks"]))
        assert payload["published_to_word"] is False and payload["usage"] == {}
        assert len(payload["notes"]["evidence_coverage"]) == 30
    assert counts == [20, 10]
    repeated = generate_batch(workspace, "OPS", [1, 2], selection=selection, budgets=VERIFIED, provider_adapter=adapter)
    assert len(adapter.calls) == 1 and all(entry["reused"] for entry in repeated["candidates"])


def test_unknown_limits_use_sequential_groups_with_untrusted_continuity(tmp_path):
    workspace, _, selection = workspace_and_deck(tmp_path)
    adapter = Adapter()
    result = generate_batch(workspace, "OPS", [1, 2], selection=selection, provider_adapter=adapter)
    assert len(adapter.calls) == 2 and result["groups"] == [["OPS-L01"], ["OPS-L02"]]
    context = _context(adapter.calls[1])
    trusted, untrusted = context.split("UNTRUSTED SOURCE DATA", 1)
    assert "Distinct recap for OPS-L01" not in trusted
    assert "Distinct recap for OPS-L01" in untrusted
    assert "narrative_fingerprints" not in context


def test_skipped_middle_slide_is_labelled_inline_between_topic_anchors(tmp_path):
    workspace, deck, selection = workspace_and_deck(tmp_path)
    key = f"sha256:{file_sha256(deck)}:page:10"
    selection[key] = {"owner": "OPS-L01", "status": "supplement"}
    plan = plan_batch(workspace, "OPS", [1, 2], selection=selection)
    assert plan["ownership"][key]["anchor"] == {"after": "S009", "before": "S011"}
    result = generate_batch(workspace, "OPS", [1, 2], selection=selection, provider_adapter=Adapter())
    first = json.loads(Path(result["candidates"][0]["candidate"]).read_text())["notes"]["blocks"]
    supplement = next(block for block in first if block.get("coverage_role") == "supplement")
    assert supplement["text"].startswith("Supplementary slide material: ")
    paragraphs = [block for block in first if block["kind"] == "paragraph"]
    assert paragraphs[9] == supplement and paragraphs[8]["evidence_ids"] == ["S009"] and paragraphs[10]["evidence_ids"] == ["S011"]


def test_supplement_placed_in_end_appendix_is_rejected(tmp_path):
    workspace, deck, selection = workspace_and_deck(tmp_path)
    selection[f"sha256:{file_sha256(deck)}:page:10"] = {"owner": "OPS-L01", "status": "supplement"}
    def wrong(response, request):
        blocks = response["lessons"][0]["notes"]["blocks"]
        supplement = next(block for block in blocks if block.get("coverage_role") == "supplement")
        blocks.remove(supplement)
        blocks.insert(-2, supplement)
    with pytest.raises(WorkflowError, match="inline"):
        generate_batch(workspace, "OPS", [1, 2], selection=selection, provider_adapter=Adapter(wrong))
    assert not (workspace.runs_dir / "OPS-L01/lecture_notes.candidate.json").exists()


def test_foreign_full_explanation_and_duplicate_narrative_are_rejected(tmp_path):
    workspace, _, selection = workspace_and_deck(tmp_path)
    def wrong(response, request):
        first = response["lessons"][0]["notes"]["blocks"]
        first.insert(1, {"kind": "paragraph", "text": "This wrongly covers the next lesson's material.", "evidence_ids": ["S021"], "coverage_role": "full"})
    with pytest.raises(WorkflowError, match="non-owned"):
        generate_batch(workspace, "OPS", [1, 2], selection=selection, budgets=VERIFIED, provider_adapter=Adapter(wrong))
    def duplicate(response, request):
        for row in response["lessons"]:
            row["notes"]["blocks"][1]["text"] = "The identical unsupported repeated full explanation appears in both lessons."
    # A different rejected draft requires a separate workspace: retries now
    # retain the first successful transport instead of paying for another.
    workspace, _, selection = workspace_and_deck(tmp_path / "duplicate-case")
    with pytest.raises(WorkflowError, match="identical narrative"):
        generate_batch(workspace, "OPS", [1, 2], selection=selection, budgets=VERIFIED, provider_adapter=Adapter(duplicate))


def test_quota_resume_preserves_first_group_and_usage_refs(tmp_path):
    workspace, _, selection = workspace_and_deck(tmp_path)
    adapter = Adapter(fail_at=2)
    with pytest.raises(ProviderQuotaError):
        generate_batch(workspace, "OPS", [1, 2], selection=selection, provider_adapter=adapter)
    first = (workspace.runs_dir / "OPS-L01/lecture_notes.candidate.json").read_bytes()
    adapter.fail_at = None
    result = generate_batch(workspace, "OPS", [1, 2], selection=selection, provider_adapter=adapter)
    assert len(adapter.calls) == 3 and result["candidates"][0]["reused"]
    assert (workspace.runs_dir / "OPS-L01/lecture_notes.candidate.json").read_bytes() == first
    assert len(result["shared_attempts"]) == 2


def test_corrupt_phase_and_edited_candidate_do_not_trigger_provider(tmp_path):
    workspace, _, selection = workspace_and_deck(tmp_path)
    adapter = Adapter()
    result = generate_batch(workspace, "OPS", [1, 2], selection=selection, budgets=VERIFIED, provider_adapter=adapter)
    checkpoint = next(workspace.runs_dir.rglob("writing-*.json"))
    original = checkpoint.read_bytes()
    checkpoint.write_text("{}")
    with pytest.raises(WorkflowError, match="damaged"):
        generate_batch(workspace, "OPS", [1, 2], selection=selection, budgets=VERIFIED, provider_adapter=adapter)
    assert len(adapter.calls) == 1
    checkpoint.write_bytes(original)
    Path(result["candidates"][0]["candidate"]).write_text("user edited candidate")
    with pytest.raises(WorkflowError, match="edited or removed"):
        generate_batch(workspace, "OPS", [1, 2], selection=selection, budgets=VERIFIED, provider_adapter=adapter)
    assert len(adapter.calls) == 1


def test_quality_and_run_instructions_are_fingerprinted_and_preserved(tmp_path):
    workspace, _, selection = workspace_and_deck(tmp_path)
    adapter = Adapter()
    first = generate_batch(workspace, "OPS", [1, 2], selection=selection, quality="approfondito", budgets=VERIFIED,
                           instructions="Explain the assumptions.", image_limit=0, provider_adapter=adapter)
    assert [request.phase for request in adapter.calls] == ["writing", "scoped_review"]
    payload = json.loads(Path(first["candidates"][0]["candidate"]).read_text())
    assert payload["request_options"] == {"instructions": "Explain the assumptions.", "image_limit": 0}
    second = plan_batch(workspace, "OPS", [1, 2], selection=selection, quality="approfondito", budgets=VERIFIED, instructions="Different instructions.", image_limit=0)
    assert second["fingerprint"] != first["fingerprint"]


def test_sources_changing_during_generation_preserve_prior_candidate(tmp_path):
    workspace, deck, selection = workspace_and_deck(tmp_path)
    def change(response, request):
        deck.write_bytes(b"externally changed source")
    with pytest.raises(WorkflowError, match="Sources changed"):
        generate_batch(workspace, "OPS", [1, 2], selection=selection, budgets=VERIFIED, provider_adapter=Adapter(change))
    assert not (workspace.runs_dir / "OPS-L01/lecture_notes.candidate.json").exists()


def test_image_only_slide_identity_and_inventory(tmp_path):
    pytest.importorskip("pypdfium2")
    from PIL import Image
    workspace, _, _ = workspace_and_deck(tmp_path, lessons=1)
    image = Image.new("RGB", (20, 20), "red")
    deck = tmp_path / "image-only.pdf"
    image.save(deck, "PDF")
    pack = build_source_pack(workspace, "OPS-L01", materials=[deck])
    slide = pack["evidence"][0]
    assert slide["content"] == "" and slide["textless"]
    assert slide["canonical_slide_id"] == f"sha256:{file_sha256(deck)}:page:1"
    assert pack["visuals"][0]["evidence_id"] == slide["id"]
    assert len(pack["slide_inventory"]) == 1


def publish_synthetically(workspace, lecture, candidate_name):
    """Bind a generated candidate to a synthetic Word without a renderer/app."""
    import sqlite3
    from docx import Document
    from mynotes.workflow import WorkflowState
    path = Path(candidate_name)
    payload = json.loads(path.read_text())
    output = workspace.runs_dir / "synthetic-published.docx"
    document = Document()
    document.add_paragraph(f"MYNOTES START {lecture}")
    document.add_paragraph("Synthetic source-bound publication fixture.")
    document.add_paragraph(f"MYNOTES END {lecture}")
    document.save(output)
    payload.update(published_to_word=True, published_path=str(output), published_hash=file_sha256(output))
    atomic_json_write(path, payload)
    state = WorkflowState(workspace.database_path)
    state.record(lecture, "synthetic-boundary", ProviderResponse(payload["notes"], "codex", "test-model", "high"), "standard", output, {"materials": True})
    state.close()
    with sqlite3.connect(workspace.database_path) as connection:
        connection.execute("UPDATE candidate_runs SET status='PUBLISHED',candidate_hash=? WHERE lecture_id=?", (file_sha256(path), lecture))
    return output


def test_verified_prior_published_coverage_is_not_regenerated_in_later_lesson(tmp_path):
    workspace, _, selection = workspace_and_deck(tmp_path)
    first = generate_batch(workspace, "OPS", [1, 2], selection=selection, budgets=VERIFIED, provider_adapter=Adapter())
    publish_synthetically(workspace, "OPS-L01", first["candidates"][0]["candidate"])
    plan = plan_batch(workspace, "OPS", [2])
    assert plan["status"] == "PLANNED"
    assert sum(choice["owner"] == "OPS-L01" for choice in plan["ownership"].values()) == 20
    assert "OPS-L01" in plan["prior_publications"]
    adapter = Adapter()
    result = generate_batch(workspace, "OPS", [2], provider_adapter=adapter)
    later = json.loads(Path(result["candidates"][0]["candidate"]).read_text())
    assert sum(block["kind"] == "paragraph" for block in later["notes"]["blocks"]) == 10
    assert adapter.calls[0].batch_context["continuity"][0]["lecture_id"] == "OPS-L01"


def test_changed_prior_published_word_blocks_generation_for_selection(tmp_path):
    workspace, _, selection = workspace_and_deck(tmp_path)
    first = generate_batch(workspace, "OPS", [1, 2], selection=selection, budgets=VERIFIED, provider_adapter=Adapter())
    output = publish_synthetically(workspace, "OPS-L01", first["candidates"][0]["candidate"])
    output.write_bytes(b"externally changed Word")
    adapter = Adapter()
    plan = generate_batch(workspace, "OPS", [2], provider_adapter=adapter)
    assert plan["status"] == "NEEDS_SELECTION" and adapter.calls == []
    assert any("prior published" in row["reason"] for row in plan["unresolved"])


def test_prior_source_manifest_ownership_corruption_is_not_trusted(tmp_path):
    workspace, _, selection = workspace_and_deck(tmp_path)
    first = generate_batch(workspace, "OPS", [1, 2], selection=selection, budgets=VERIFIED, provider_adapter=Adapter())
    publish_synthetically(workspace, "OPS-L01", first["candidates"][0]["candidate"])
    Path(first["manifest"]).write_text("{}")
    plan = plan_batch(workspace, "OPS", [2])
    assert plan["status"] == "NEEDS_SELECTION"


@pytest.mark.parametrize("speaker_role", ["Professor", "UNKNOWN"])
def test_repeated_slide_can_add_teacher_insight_but_cannot_repeat_full_text(tmp_path, speaker_role):
    workspace, _, selection = workspace_and_deck(tmp_path)
    pack_path = workspace.runs_dir / "OPS-L02/source_pack.json"
    pack = json.loads(pack_path.read_text())
    pack["evidence"].append({"id": "T001", "kind": "transcript", "text": "Additional source supported teacher insight.", "start_seconds": 0, "end_seconds": 5, "speaker_role": speaker_role})
    atomic_json_write(pack_path, pack)
    def insight(response, request):
        if "OPS-L02" not in request.batch_context["lessons"]:
            return
        row = next(row for row in response["lessons"] if row["lecture_id"] == "OPS-L02")
        row["notes"]["blocks"].insert(1, {"kind": "paragraph", "text": "Additional source supported teacher insight about the previous slide.", "evidence_ids": ["S001", "T001"], "coverage_role": "new_insight", "prior_lecture_id": "OPS-L01"})
    result = generate_batch(workspace, "OPS", [1, 2], selection=selection, provider_adapter=Adapter(insight))
    later = json.loads(Path(result["candidates"][1]["candidate"]).read_text())["notes"]
    assert any(block.get("coverage_role") == "new_insight" for block in later["blocks"])


def test_overlapping_concurrent_generation_is_rejected_without_second_call(tmp_path):
    from mynotes.workflow import _generation_lock
    workspace, _, selection = workspace_and_deck(tmp_path)
    adapter = Adapter()
    with _generation_lock(workspace, "OPS-L01"):
        with pytest.raises(WorkflowError, match="already running"):
            generate_batch(workspace, "OPS", [1, 2], selection=selection, provider_adapter=adapter)
    assert adapter.calls == []


def test_identical_copied_decks_have_shared_identity_and_changed_deck_is_invalidated(tmp_path):
    workspace, deck, selection = workspace_and_deck(tmp_path)
    alias = tmp_path / "copied-deck.pdf"
    alias.write_bytes(deck.read_bytes())
    path = workspace.runs_dir / "OPS-L02/source_pack.json"
    pack = json.loads(path.read_text())
    for row in pack["evidence"]:
        row["source_file"] = str(alias)
    pack["sources"] = [{"path": str(alias), "sha256": file_sha256(alias)}]
    atomic_json_write(path, pack)
    plan = plan_batch(workspace, "OPS", [1, 2], selection=selection, budgets=VERIFIED)
    assert len(plan["ownership"]) == 30
    alias.write_bytes(b"new deck version")
    changed = plan_batch(workspace, "OPS", [1, 2], selection=selection)
    assert changed["status"] == "NEEDS_PREPARATION" and changed["lessons_needing_preparation"][0]["lecture_id"] == "OPS-L02"


def test_image_only_pptx_inventory_without_office_renderer(tmp_path, monkeypatch):
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    from mynotes import extraction
    workspace, _, _ = workspace_and_deck(tmp_path, lessons=1)
    image = tmp_path / "diagram.png"
    Image.new("RGB", (20, 20), "blue").save(image)
    deck = tmp_path / "image-only.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(str(image), Inches(1), Inches(1))
    presentation.save(deck)
    monkeypatch.setattr(extraction, "_render_pptx_as_pdf", lambda *_: None)
    pack = build_source_pack(workspace, "OPS-L01", materials=[deck])
    assert len(pack["slide_inventory"]) == 1 and pack["evidence"][0]["textless"]
    assert pack["visuals"][0]["evidence_id"] == "S001"
    assert pack["visuals"][0]["render_error"]


def test_selection_and_preparation_manifests_can_resume_without_initial_lessons(tmp_path):
    from mynotes.batch import load_batch_manifest
    workspace, _, selection = workspace_and_deck(tmp_path)
    result = generate_batch(workspace, "OPS", [1, 2])
    assert result["status"] == "NEEDS_SELECTION"
    saved = load_batch_manifest(workspace, result["batch_id"])
    assert saved["lectures"] == ["OPS-L01", "OPS-L02"]
    assert saved["budgets"]["context_tokens"] is None
    assert saved["budget_basis"] == "cli_only_model_limits_unknown"
    resumed = generate_batch(workspace, "OPS", [1, 2], batch_id=result["batch_id"], selection=selection, provider_adapter=Adapter())
    assert resumed["status"] == "CANDIDATE_READY"
    (workspace.runs_dir / "OPS-L02/source_pack.json").unlink()
    missing = plan_batch(workspace, "OPS", [2])
    assert missing["status"] == "NEEDS_PREPARATION"
    assert missing["lessons_needing_preparation"][0]["lecture_id"] == "OPS-L02"
    with pytest.raises(WorkflowError, match="Invalid batch ID"):
        load_batch_manifest(workspace, "../escape")


def test_known_output_budget_rejects_oversized_lesson_before_invocation(tmp_path):
    workspace, _, selection = workspace_and_deck(tmp_path)
    budget = BatchBudget(verified_for_model="test-model", context_tokens=1_000_000, output_tokens=100, output_reserve_per_lesson=4000)
    adapter = Adapter()
    with pytest.raises(WorkflowError, match="cannot fit"):
        generate_batch(workspace, "OPS", [1, 2], selection=selection, budgets=budget, provider_adapter=adapter)
    assert adapter.calls == []


def test_known_output_budget_still_applies_when_image_cost_is_unknown(tmp_path):
    from PIL import Image
    workspace, _, selection = workspace_and_deck(tmp_path)
    path = workspace.runs_dir / "OPS-L01" / "source_pack.json"
    pack = json.loads(path.read_text())
    image = path.parent / "slide.png"
    Image.new("RGB", (4, 4), "white").save(image)
    pack["visuals"] = [{"evidence_id": "S001", "image_path": str(image), "visual_decision": "candidate"}]
    atomic_json_write(path, pack)
    budget = BatchBudget(verified_for_model="test-model", context_tokens=1_000_000,
                        output_tokens=100, output_reserve_per_lesson=4000)
    adapter = Adapter()
    with pytest.raises(WorkflowError, match="cannot fit"):
        generate_batch(workspace, "OPS", [1, 2], selection=selection, budgets=budget, provider_adapter=adapter)
    assert adapter.calls == []


def test_supplementary_table_is_labelled_in_its_inline_position(tmp_path):
    workspace, deck, selection = workspace_and_deck(tmp_path)
    selection[f"sha256:{file_sha256(deck)}:page:10"] = {"owner": "OPS-L01", "status": "supplement"}
    def table_only(response, request):
        if response["lessons"][0]["lecture_id"] != "OPS-L01":
            return
        blocks = response["lessons"][0]["notes"]["blocks"]
        index = next(i for i, block in enumerate(blocks) if block.get("coverage_role") == "supplement")
        blocks[index] = {"kind": "table", "rows": [["Concept", "Description"], ["Topic", "Source supported content"]],
                         "evidence_ids": ["S010"], "coverage_role": "supplement"}
    result = generate_batch(workspace, "OPS", [1, 2], selection=selection, provider_adapter=Adapter(table_only))
    blocks = json.loads(Path(result["candidates"][0]["candidate"]).read_text())["notes"]["blocks"]
    table_index = next(i for i, block in enumerate(blocks) if block["kind"] == "table")
    assert blocks[table_index - 1]["text"] == "Supplementary slide material."
    assert blocks[table_index - 2]["evidence_ids"] == ["S009"]
    assert blocks[table_index + 1]["evidence_ids"] == ["S011"]


def test_review_answer_cannot_reexplain_another_lessons_slide(tmp_path):
    workspace, _, selection = workspace_and_deck(tmp_path)
    def duplicate_answer(response, request):
        later = response["lessons"][1]["notes"]
        question = next(block for block in later["blocks"] if block["kind"] == "qa")
        question["evidence_ids"] = ["S001"]
        question["answer"] = "Detailed repeated explanation " * 40
    with pytest.raises(WorkflowError, match="brief reference"):
        generate_batch(workspace, "OPS", [1, 2], selection=selection, budgets=VERIFIED,
                       provider_adapter=Adapter(duplicate_answer))


def test_later_independent_source_change_reuses_completed_earlier_group(tmp_path):
    workspace, _, selection = workspace_and_deck(tmp_path)
    adapter = Adapter()
    first = generate_batch(workspace, "OPS", [1, 2], selection=selection, provider_adapter=adapter)
    assert len(adapter.calls) == 2
    earlier = Path(first["candidates"][0]["candidate"]).read_bytes()
    pack_path = workspace.runs_dir / "OPS-L02/source_pack.json"
    pack = json.loads(pack_path.read_text())
    pack["evidence"].append({"id": "H001", "kind": "personal_note", "content": "A new lesson-specific verified example", "source_file": "local-notes", "paragraph": 1})
    atomic_json_write(pack_path, pack)
    second = generate_batch(workspace, "OPS", [1, 2], selection=selection, provider_adapter=adapter)
    assert len(adapter.calls) == 3 and second["candidates"][0]["reused"]
    assert Path(second["candidates"][0]["candidate"]).read_bytes() == earlier
    assert first["fingerprint"] != second["fingerprint"]


def test_legacy_published_overlap_requires_selection_instead_of_invented_coverage(tmp_path):
    import sqlite3
    from docx import Document
    from mynotes.workflow import WorkflowState
    workspace, _, _ = workspace_and_deck(tmp_path)
    output = workspace.runs_dir / "legacy.docx"
    document = Document()
    document.add_paragraph("Existing historical notes without canonical provenance")
    document.save(output)
    state = WorkflowState(workspace.database_path)
    state.record("OPS-L01", "legacy", ProviderResponse({}, "codex", "test-model", "high"), "standard", output, {"materials": True})
    state.close()
    plan = plan_batch(workspace, "OPS", [2])
    assert plan["status"] == "NEEDS_SELECTION"
    assert any("no verified canonical ownership" in entry["reason"] for entry in plan["unresolved"])
