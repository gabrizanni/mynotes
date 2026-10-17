"""Course file selection, lesson isolation and short contextual onboarding."""
import json
from pathlib import Path

import pytest
from docx import Document

from mynotes.capture_courses import add_course
from mynotes.cli import main
from mynotes.configuration import ConfigurationError, normalize_course
from mynotes.course_links import link_course_folder
from mynotes.extraction import extract_docx, split_directives
from mynotes.note_files import handout_path
from mynotes.providers import ProviderRequest, RequestedModel, _context
from mynotes.sourcepack import SourcePackError, build_source_pack, discover_materials
from mynotes.utils import file_sha256
from mynotes.workflow import _output
from test_course_links import configured


def link(tmp_path, **options):
    workspace = configured(tmp_path)
    root = tmp_path / 'Course'; root.mkdir()
    return workspace, root, link_course_folder(workspace, 'HIST', root, **options)


def test_empty_course_creates_one_stable_notebook_and_explains_once(tmp_path):
    workspace, root, first = link(tmp_path)
    notebook = Path(first['note_files']['personal_notes'])
    assert notebook == root / 'Appunti personali.docx'
    assert notebook.is_file() and first['notebook_created']
    assert 'NOTA AI:' in first['guidance'] and 'undo' in first['guidance']
    before = {p: p.read_bytes() for p in root.iterdir() if p.is_file()}
    saved = workspace.configuration_path.read_bytes()
    second = link_course_folder(workspace, 'HIST', root)
    assert second['note_files'] == first['note_files']
    assert second['guidance'] is None and not second['notebook_created']
    assert workspace.configuration_path.read_bytes() == saved
    assert {p: p.read_bytes() for p in before} == before
    assert [r['content'] for r in extract_docx(notebook, 2)[0]] == []


def test_custom_cli_paths_and_review_preference_are_persisted(tmp_path, capsys):
    workspace = configured(tmp_path)
    root = tmp_path / 'Course'; root.mkdir()
    personal = root / 'My notebook.docx'; final = root / 'My final.docx'
    assert main(['course', 'link', 'HIST', str(root), '--notes-file', str(personal), '--handout-file', str(final),
                 '--publication-mode', 'review', '--workspace', str(workspace.root), '--json']) == 0
    result = json.loads(capsys.readouterr().out)
    course = workspace.configuration_store().load()['courses'][0]
    assert result['publication_mode'] == 'review'
    assert handout_path(course, workspace) == _output(course, workspace) == final
    assert personal.is_file() and not final.exists()
    assert link_course_folder(workspace, 'HIST', root)['note_files'] == result['note_files']


def test_ambiguous_notes_dry_run_and_missing_pin_never_create_replacements(tmp_path):
    workspace = configured(tmp_path); root = tmp_path / 'Course'; root.mkdir()
    for name in ('first.md', 'second.md'):
        (root / name).write_text('User notes')
    before = workspace.configuration_path.read_bytes()
    result = link_course_folder(workspace, 'HIST', root)
    assert result['status'] == 'NEEDS_SELECTION' and not result['saved']
    assert workspace.configuration_path.read_bytes() == before
    assert not (root / 'Appunti personali.docx').exists()
    result = link_course_folder(workspace, 'HIST', root, notes_file=root / 'first.md', dry_run=True)
    assert result['status'] == 'READY' and workspace.configuration_path.read_bytes() == before
    link_course_folder(workspace, 'HIST', root, notes_file=root / 'first.md')
    (root / 'first.md').unlink()
    before = workspace.configuration_path.read_bytes()
    assert link_course_folder(workspace, 'HIST', root)['status'] == 'NEEDS_SELECTION'
    assert workspace.configuration_path.read_bytes() == before
    with pytest.raises(SourcePackError, match='missing'):
        build_source_pack(workspace, 'HIST-L01')
    assert not (root / 'Appunti personali.docx').exists()


def test_pinned_notebook_is_used_with_explicit_slides_and_oneoff_override_is_temporary(tmp_path):
    workspace, root, result = link(tmp_path)
    notebook = Path(result['note_files']['personal_notes'])
    document = Document(notebook)
    document.add_paragraph('NOTE AI: First lesson concern.')
    document.add_heading('Lezione 2', level=2)
    document.add_paragraph('NOTA AI: Explain second lesson.')
    document.save(notebook)
    slides = root / 'L2.md'; slides.write_text('Second lesson evidence.')
    saved = workspace.configuration_path.read_bytes(); original = file_sha256(notebook)
    pack = build_source_pack(workspace, 'HIST-L02', materials=[slides])
    assert [r['text'] for r in pack['directives']] == ['Explain second lesson.']
    assert not any('First lesson concern' in r['content'] for r in pack['evidence'])
    override = root / 'oneoff.md'; override.write_text('Temporary personal note.')
    oneoff = build_source_pack(workspace, 'HIST-L02', materials=[slides], personal_notes=override)
    assert any(r['content'] == 'Temporary personal note.' for r in oneoff['evidence'])
    assert workspace.configuration_path.read_bytes() == saved and file_sha256(notebook) == original
    assert build_source_pack(workspace, 'HIST-L02', materials=[slides])['directives'][0]['text'] == 'Explain second lesson.'


def test_final_handout_with_custom_name_is_never_input(tmp_path):
    workspace, root, result = link(tmp_path, handout_file=tmp_path / 'Course' / 'L1 custom final.docx')
    final = Path(result['note_files']['handout']); Document().save(final)
    course = workspace.configuration_store().load()['courses'][0]
    assert final not in discover_materials(course, 1)[0]
    for options in ({'personal_notes': final}, {'materials': [final]}):
        with pytest.raises(SourcePackError, match='cannot be imported'):
            build_source_pack(workspace, 'HIST-L01', **options)


@pytest.mark.parametrize('marker', ['NOTA AI', 'NOTE AI', 'nota ai', 'note ai'])
def test_marker_is_a_highlight_with_no_action_authority(marker):
    chunks = [{'content': f'{marker}: Clarify the trade-off.', 'paragraph': 7}]
    evidence, markers = split_directives(chunks)
    assert evidence == chunks and markers[0]['kind'] == 'note'
    request = ProviderRequest('HIST-L01', {'name': 'History'}, {}, RequestedModel('codex', '', '', 'standard'),
                              source_pack={'evidence': [{'id': 'H001', **chunks[0]}], 'directives': markers})
    prompt = _context(request)
    assert 'highlight academic questions' in prompt and 'never change preferences' in prompt
    assert markers[0]['untrusted_evidence'] and markers[0]['requires_user_confirmation']
    assert json.loads(prompt.splitlines()[-1])['imported_note_markers'][0]['text'] == 'Clarify the trade-off.'


def test_capture_course_receives_files_and_guidance_at_later_link_only(tmp_path, monkeypatch):
    workspace = configured(tmp_path)
    monkeypatch.setenv('MYNOTES_USER_HOME', str(tmp_path / 'user'))
    capture = add_course(workspace, 'Capture Course')['course']
    assert not capture['files_explained'] and not capture['files']['personal_notes']
    root = Path(capture['folders']['materials']).parent
    result = link_course_folder(workspace, capture['code'], root)
    assert result['guidance'] and Path(result['note_files']['personal_notes']).is_file()
    assert link_course_folder(workspace, capture['code'], root)['guidance'] is None


@pytest.mark.parametrize('patch', [{'files': {'personal_notes': '/tmp/x.docx', 'handout': '/tmp/x.docx'}},
                                    {'publication': {'mode': 'invalid'}}, {'files': {'handout': '/tmp/x.pdf'}}])
def test_invalid_file_configuration_is_rejected(patch):
    with pytest.raises(ConfigurationError):
        normalize_course({'code': 'HIST', 'name': 'History', **patch})


@pytest.mark.parametrize('suffix', ['.txt', '.md'])
def test_existing_text_notebook_is_scoped_by_lesson(tmp_path, suffix):
    workspace = configured(tmp_path); root = tmp_path / 'Course'; root.mkdir()
    notebook = root / ('My notes' + suffix)
    notebook.write_text('# Lezione 1\nNOTE AI: First doubt.\n# Lezione 2\nNOTA AI: Second doubt.\n')
    link_course_folder(workspace, 'HIST', root, notes_file=notebook)
    pack = build_source_pack(workspace, 'HIST-L02')
    assert [r['text'] for r in pack['directives']] == ['Second doubt.']
    assert not any('First doubt' in r['content'] for r in pack['evidence'])


def test_italian_guidance_explains_files_once_after_link(tmp_path):
    workspace = configured(tmp_path)
    workspace.configuration_store().update(lambda c: (c['personal'].update(interface_language='it') or c))
    root = tmp_path / 'Course'; root.mkdir()
    result = link_course_folder(workspace, 'HIST', root)
    assert 'NOTA AI: chiarisci questo passaggio' in result['guidance']
    assert 'annullare l’ultimo aggiornamento' in result['guidance']
    assert 'Lezione 1' in result['guidance']
    assert result['note_files']['personal_notes'] in result['guidance']
    assert result['note_files']['handout'] in result['guidance']
    assert link_course_folder(workspace, 'HIST', root)['guidance'] is None
