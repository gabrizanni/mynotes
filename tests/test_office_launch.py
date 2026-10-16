"""Prevent native macOS renderer crashes before spawning a process."""
import subprocess
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from types import SimpleNamespace

import pytest

from mynotes import extraction, office, publication


@pytest.mark.parametrize("denied", [False, True])
def test_probe_checks_actual_macos_permissions_not_only_environment(monkeypatch, denied):
    monkeypatch.setattr(office.sys, "platform", "darwin")
    monkeypatch.delenv("CODEX_SANDBOX", raising=False)
    queries = []
    class Check:
        def __call__(self, pid, operation, filter_type, service):
            queries.append((operation, filter_type, service.value))
            return int(denied)
    monkeypatch.setattr(office.ctypes, "CDLL", lambda _: SimpleNamespace(sandbox_check=Check()))
    assert bool(office.macos_launch_restriction()) == denied
    assert queries[0] == (b"mach-lookup", 2, b"com.apple.coreservices.launchservicesd")


def test_missing_probe_uses_codex_marker_but_does_not_block_normal_terminal(monkeypatch):
    monkeypatch.setattr(office.sys, "platform", "darwin")
    def unavailable(_):
        raise OSError("unsupported")
    monkeypatch.setattr(office.ctypes, "CDLL", unavailable)
    monkeypatch.setenv("CODEX_SANDBOX", "seatbelt")
    assert office.macos_launch_restriction()
    monkeypatch.delenv("CODEX_SANDBOX")
    assert office.macos_launch_restriction() is None


def test_non_macos_does_not_probe_or_block(monkeypatch):
    monkeypatch.setattr(office.sys, "platform", "linux")
    monkeypatch.setattr(office.ctypes, "CDLL", lambda _: pytest.fail("unexpected OS probe"))
    assert office.macos_launch_restriction() is None


def test_visual_cache_identity_changes_when_launch_permissions_change(monkeypatch):
    from mynotes.extraction_cache import extractor_identity
    monkeypatch.setattr(office, "macos_launch_restriction", lambda: "LaunchServices denied")
    restricted = extractor_identity(".pptx", render_visuals=True)
    monkeypatch.setattr(office, "macos_launch_restriction", lambda: None)
    allowed = extractor_identity(".pptx", render_visuals=True)
    assert restricted != allowed
    assert allowed["renderer_launch_restriction"] is None


@pytest.mark.parametrize("kind", ["word", "slides"])
def test_restricted_conversion_never_starts_process_or_validates(kind, tmp_path, monkeypatch):
    monkeypatch.setattr(office, "macos_launch_restriction", lambda: "LaunchServices denied")
    monkeypatch.setattr(office.subprocess, "run", lambda *a, **k: pytest.fail("renderer was started"))
    monkeypatch.setattr(publication, "_soffice", lambda: "soffice")
    monkeypatch.setattr(extraction, "libreoffice_executable", lambda: "soffice")
    function = publication._render_and_validate if kind == "word" else extraction._render_pptx_as_pdf
    error = publication.PublicationError if kind == "word" else extraction.ExtractionError
    with pytest.raises(error, match="was not started.*prevent.*crash"):
        function(tmp_path / "source.docx", tmp_path / "output")
    assert not (tmp_path / "output").exists()


def test_concurrent_conversions_have_private_profiles_and_clean_up(tmp_path, monkeypatch):
    monkeypatch.setattr(office, "macos_launch_restriction", lambda: None)
    profiles = []
    def run(command, **kwargs):
        profiles.append(command[1])
        assert all(flag in command for flag in ("--headless", "--norestore", "--nofirststartwizard"))
        converted = Path(command[command.index("--outdir") + 1])
        (converted / f"{Path(command[-1]).stem}.pdf").write_bytes(b"fresh PDF")
        return subprocess.CompletedProcess(command, 0, "", "")
    monkeypatch.setattr(office.subprocess, "run", run)
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda i: office.convert_to_pdf("soffice", tmp_path / "source.pptx", tmp_path / str(i), timeout=1), range(2)))
    assert len(set(profiles)) == 2
    assert all(path.read_bytes() == b"fresh PDF" for path in results)
    from urllib.parse import unquote, urlparse
    assert all(not Path(unquote(urlparse(value.split("=", 1)[1]).path)).exists() for value in profiles)


def test_pptx_keeps_text_and_records_unverified_visuals_without_launch(tmp_path, monkeypatch):
    from pptx import Presentation
    from pptx.util import Inches
    monkeypatch.setattr(office, "macos_launch_restriction", lambda: "LaunchServices denied")
    monkeypatch.setattr(office.subprocess, "run", lambda *a, **k: pytest.fail("renderer started"))
    monkeypatch.setattr(extraction, "libreoffice_executable", lambda: "soffice")
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[5])
    slide.shapes.title.text = "Evidence remains available"
    slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(2))
    path = tmp_path / "lesson.pptx"
    slides.save(path)
    chunks, visuals = extraction.extract_pptx(path, tmp_path / "visuals")
    assert "Evidence remains available" in chunks[0]["content"]
    assert visuals[0]["image_path"] is None
    assert "was not started" in visuals[0]["render_error"]


@pytest.mark.parametrize("failure", ["missing", "empty", "abort", "timeout"])
def test_failed_conversion_does_not_accept_old_pdf_or_retry(failure, tmp_path, monkeypatch):
    monkeypatch.setattr(office, "macos_launch_restriction", lambda: None)
    output = tmp_path / "output"
    output.mkdir()
    stale = output / "source.pdf"
    stale.write_bytes(b"old PDF")
    calls = []
    def run(command, **kwargs):
        calls.append(command)
        if failure == "timeout":
            raise subprocess.TimeoutExpired(command, 1)
        converted = Path(command[command.index("--outdir") + 1])
        if failure == "empty":
            (converted / "source.pdf").touch()
        return subprocess.CompletedProcess(command, -6 if failure == "abort" else 0, "", "abort" if failure == "abort" else "")
    monkeypatch.setattr(office.subprocess, "run", run)
    with pytest.raises(office.OfficeError):
        office.convert_to_pdf("soffice", tmp_path / "source.docx", output, timeout=1)
    assert len(calls) == 1
    assert stale.read_bytes() == b"old PDF"
    assert len(list(output.iterdir())) == 1
