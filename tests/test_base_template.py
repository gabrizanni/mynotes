"""The packaged Word base contains formatting definitions, never lecture content."""
from __future__ import annotations

from importlib.resources import files
from zipfile import ZipFile

from docx import Document


def test_base_template_has_roles_and_no_source_content() -> None:
    template = files("mynotes").joinpath("resources/base-template.docx")
    with template.open("rb") as stream:
        document = Document(stream)
    assert not document.paragraphs
    assert not document.tables
    assert len(document.sections) == 1
    section = document.sections[0]
    assert abs(section.page_width.cm - 21) < 0.01
    assert abs(section.page_height.cm - 29.7) < 0.01
    for margin in (section.top_margin, section.bottom_margin, section.left_margin, section.right_margin):
        assert abs(margin.cm - 1.5) < 0.01
    assert abs(section.footer_distance.cm - 1.25) < 0.01
    for role in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Termine tecnico", "List Number", "List Bullet", "Caption", "Intense Quote"):
        assert role in document.styles

    for level, points in ((1, 16), (2, 14), (3, 12)):
        heading = document.styles[f"Heading {level}"]
        assert heading.builtin and heading.quick_style and not heading.hidden
        assert heading.font.size.pt == points
        assert heading._element.xpath("./w:rPr/w:color")[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}themeColor") is None
        assert heading._element.xpath("./w:pPr/w:outlineLvl")[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") == str(level - 1)
    term = document.styles["Termine tecnico"]
    assert term.type == 2 and term.quick_style and not term.hidden
    assert term.font.name == "Aptos" and term.font.bold
    assert str(term.font.color.rgb) == "ED7D31"

    with template.open("rb") as stream, ZipFile(stream) as archive:
        names = set(archive.namelist())
        assert names == {
            "[Content_Types].xml", "_rels/.rels", "word/document.xml",
            "word/_rels/document.xml.rels", "word/styles.xml",
            "word/numbering.xml", "word/theme/theme1.xml",
        }
        assert b"<w:p>" not in archive.read("word/document.xml")
        assert b"<w:t" not in archive.read("word/document.xml")
        assert not any("media" in name or "docProps" in name or "customXml" in name for name in names)
