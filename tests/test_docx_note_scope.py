"""Standalone notes and structured notebooks have distinct lesson scopes."""
from pathlib import Path

import pytest
from docx import Document

from mynotes.configuration import course_template
from mynotes.extraction import ExtractionError, extract_docx
from mynotes.sourcepack import build_source_pack
from mynotes.utils import file_sha256
from mynotes.workspace import Workspace


def test_standalone_docx_without_lesson_headings_keeps_all_notes_tables_markers(tmp_path):
    path = tmp_path / "standalone.docx"
    document = Document()
    document.add_paragraph("Standalone class notes")
    document.add_paragraph("Inventory balances service and cost.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "Demand", "10"
    document.add_paragraph("NOTE AI: Explain the trade-off.")
    document.save(path)
    original = file_sha256(path)
    all_items, _ = extract_docx(path)
    for number in (1, 2, 3):
        items, visuals = extract_docx(path, lesson_number=number)
        assert items == all_items and visuals == []
        assert any(item["style"] == "Table" and "Demand | 10" in item["content"] for item in items)
        assert any("NOTE AI:" in item["content"] for item in items)
    assert file_sha256(path) == original


def test_real_standalone_docx_is_available_personal_note_evidence_in_sourcepack(tmp_path):
    workspace = Workspace(tmp_path / "workspace")
    workspace.initialize()
    config = workspace.configuration_store().load()
    course = course_template()
    course.update(code="OPS", name="Operations")
    config["courses"] = [course]
    workspace.configuration_store().save(config)
    notes = tmp_path / "selected-note.docx"
    document = Document()
    document.add_paragraph("Inventory has a service-cost trade-off.")
    document.add_paragraph("NOTE AI: Explain it clearly.")
    document.save(notes)
    original = file_sha256(notes)
    pack = build_source_pack(workspace, "OPS-L02", personal_notes=notes)
    assert pack["availability"]["personal_notes"] is True
    assert len(pack["evidence"]) == 2
    assert all(row["id"].startswith("H") and row["untrusted_evidence"] is True for row in pack["evidence"])
    assert "Inventory" in pack["evidence"][0]["content"]
    assert pack["directives"][0]["requires_user_confirmation"] is True
    assert pack["directives"][0]["untrusted_evidence"] is True
    assert file_sha256(notes) == original


def test_structured_notebook_keeps_exact_lesson_slices_and_missing_lesson_empty(tmp_path):
    notes = tmp_path / "notebook.docx"
    document = Document()
    document.add_paragraph("Notebook introduction")
    document.add_heading("Lecture 1", level=2)
    document.add_paragraph("First lesson only")
    document.add_heading("L2 - following lesson", level=2)
    document.add_paragraph("Second lesson only")
    document.save(notes)
    original = file_sha256(notes)
    first, _ = extract_docx(notes, lesson_number=1)
    second, _ = extract_docx(notes, lesson_number=2)
    assert [row["content"] for row in first] == ["Lecture 1", "First lesson only"]
    assert [row["content"] for row in second] == ["L2 - following lesson", "Second lesson only"]
    assert extract_docx(notes, lesson_number=3) == ([], [])
    assert len(extract_docx(notes)[0]) == 5
    assert file_sha256(notes) == original


def test_duplicate_target_lesson_in_notebook_still_rejected(tmp_path):
    notes = tmp_path / "duplicate.docx"
    document = Document()
    document.add_heading("Lecture 1", level=2)
    document.add_paragraph("First occurrence")
    document.add_heading("Lezione 1", level=2)
    document.add_paragraph("Second occurrence")
    document.save(notes)
    original = file_sha256(notes)
    with pytest.raises(ExtractionError, match="More than one heading"):
        extract_docx(notes, lesson_number=1)
    assert file_sha256(notes) == original


def test_empty_docx_remains_empty(tmp_path):
    notes = tmp_path / "empty.docx"
    Document().save(notes)
    original = file_sha256(notes)
    assert extract_docx(notes) == ([], [])
    assert extract_docx(notes, lesson_number=2) == ([], [])
    assert file_sha256(notes) == original
