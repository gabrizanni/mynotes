"""Local automatic updates, revision, undo and publication-state recovery."""
import json
import sqlite3
from pathlib import Path

import pytest
from docx import Document

from mynotes import publication as pub
from mynotes.cli import main
from mynotes.utils import atomic_json_write, file_sha256
from test_publication import prepared, fake_render


def revise(workspace, text):
    path = workspace.runs_dir / 'OPS7-L01' / 'lecture_notes.candidate.json'
    payload = json.loads(path.read_text())
    payload['notes']['summary'] = text
    payload['published_to_word'] = False
    atomic_json_write(path, payload)
    with sqlite3.connect(workspace.database_path) as conn:
        conn.execute("UPDATE candidate_runs SET status='CANDIDATE_READY',candidate_hash=? WHERE lecture_id='OPS7-L01'", (file_sha256(path),))


def add_second(workspace):
    first = workspace.runs_dir / 'OPS7-L01'
    run = workspace.runs_dir / 'OPS7-L02'; run.mkdir()
    pack = json.loads((first / 'source_pack.json').read_text()); pack['lecture_id'] = 'OPS7-L02'
    atomic_json_write(run / 'source_pack.json', pack)
    notes = json.loads((first / 'lecture_notes.candidate.json').read_text())
    notes['notes']['lecture_id'] = 'OPS7-L02'; notes['notes']['summary'] = 'Second lesson body.'
    notes['source_pack'] = str(run / 'source_pack.json')
    notes['source_pack_hash'] = file_sha256(run / 'source_pack.json')
    notes['published_to_word'] = False
    atomic_json_write(run / 'lecture_notes.candidate.json', notes)


@pytest.fixture
def setup(tmp_path, monkeypatch):
    workspace, target = prepared(tmp_path)
    monkeypatch.setattr(pub, '_render_and_validate', fake_render)
    return workspace, target


def test_automatic_update_and_cli_undo_first_creation_preserve_safety_copy(setup, capsys):
    workspace, target = setup
    assert main(['update', 'OPS7-L01', '--workspace', str(workspace.root)]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result['status'] == 'PUBLISHED' and target.is_file()
    published = target.read_bytes()
    assert pub.update_course_notes(workspace, 'OPS7-L01')['status'] == 'SKIPPED'
    assert main(['undo', 'OPS7', '--workspace', str(workspace.root)]) == 0
    undone = json.loads(capsys.readouterr().out)
    assert not target.exists() and not undone['output_exists']
    assert Path(undone['pre_restore_backup']['path']).read_bytes() == published
    payload = json.loads((workspace.runs_dir / 'OPS7-L01' / 'lecture_notes.candidate.json').read_text())
    assert payload['published_to_word'] is False
    with sqlite3.connect(workspace.database_path) as conn:
        assert conn.execute('SELECT status FROM lesson_runs').fetchone()[0] == 'ROLLED_BACK'
    assert pub.update_course_notes(workspace, 'OPS7-L01')['status'] == 'PUBLISHED'


def test_preview_mode_does_not_publish_without_approval(setup):
    workspace, target = setup
    workspace.configuration_store().update(lambda c: (c['courses'][0]['publication'].update(mode='review') or c))
    result = pub.update_course_notes(workspace, 'OPS7-L01')
    assert result['status'] == 'AWAITING_APPROVAL' and not target.exists()
    assert pub.publish_word_candidate(workspace, 'OPS7-L01', Path(result['candidate']))['status'] == 'PUBLISHED'


def test_sequential_updates_undo_membership_and_new_update(setup):
    workspace, target = setup
    pub.update_course_notes(workspace, 'OPS7-L01'); first = target.read_bytes()
    add_second(workspace)
    pub.update_course_notes(workspace, 'OPS7-L02')
    text = '\n'.join(p.text for p in Document(target).paragraphs)
    assert 'Second lesson body.' in text and text.count('MYNOTES START OPS7-L01') == 1
    assert pub.update_course_notes(workspace, 'OPS7-L01')['status'] == 'SKIPPED'
    result = pub.undo_last_update(workspace, 'OPS7')
    assert result['lecture_id'] == 'OPS7-L02' and target.read_bytes() == first
    with sqlite3.connect(workspace.database_path) as conn:
        assert dict(conn.execute('SELECT lecture_id,status FROM lesson_runs')) == {'OPS7-L01': 'PUBLISHED', 'OPS7-L02': 'ROLLED_BACK'}
    assert pub.update_course_notes(workspace, 'OPS7-L02')['status'] == 'PUBLISHED'
    pub.undo_last_update(workspace, 'OPS7')
    pub.undo_last_update(workspace, 'OPS7')
    assert not target.exists()


def test_revision_replaces_section_and_undo_leaves_new_candidate_reusable(setup):
    workspace, target = setup
    pub.update_course_notes(workspace, 'OPS7-L01'); original = target.read_bytes()
    revise(workspace, 'Revised explanation.')
    result = pub.update_course_notes(workspace, 'OPS7-L01')
    assert result['backup'] and Path(result['backup']['path']).read_bytes() == original
    text = '\n'.join(p.text for p in Document(target).paragraphs)
    assert 'Revised explanation.' in text and text.count('MYNOTES START OPS7-L01') == 1
    pub.undo_last_update(workspace, 'OPS7')
    assert target.read_bytes() == original
    payload = json.loads((workspace.runs_dir / 'OPS7-L01' / 'lecture_notes.candidate.json').read_text())
    assert payload['notes']['summary'] == 'Revised explanation.' and not payload['published_to_word']
    assert pub.update_course_notes(workspace, 'OPS7-L01')['status'] == 'PUBLISHED'


def test_backup_failure_never_replaces_target(setup, monkeypatch):
    workspace, target = setup
    pub.update_course_notes(workspace, 'OPS7-L01'); original = target.read_bytes()
    revise(workspace, 'Revision that must not be installed.')
    def fail(*args, **kwargs):
        raise OSError('Disk full')
    monkeypatch.setattr(pub, '_archive', fail)
    with pytest.raises(OSError, match='Disk full'):
        pub.update_course_notes(workspace, 'OPS7-L01')
    assert target.read_bytes() == original


def test_external_addition_is_preserved_by_automatic_update(setup):
    workspace, target = setup
    pub.update_course_notes(workspace, 'OPS7-L01')
    document = Document(target); document.add_paragraph('User edit.'); document.save(target)
    with pytest.raises(pub.PublicationError, match='changed'):
        pub.undo_last_update(workspace, 'OPS7')
    result = pub.update_course_notes(workspace, 'OPS7-L01')
    assert result['status'] == 'PUBLISHED'
    assert 'User edit.' in '\n'.join(p.text for p in Document(target).paragraphs)


def test_missing_handout_is_not_silently_recreated(setup):
    workspace, target = setup
    pub.update_course_notes(workspace, 'OPS7-L01'); target.unlink()
    with pytest.raises(pub.PublicationError, match='missing'):
        pub.update_course_notes(workspace, 'OPS7-L01')
    assert not target.exists()


def test_corrupt_backup_stops_undo(setup):
    workspace, target = setup
    pub.update_course_notes(workspace, 'OPS7-L01')
    revise(workspace, 'New explanation.')
    result = pub.update_course_notes(workspace, 'OPS7-L01'); original = target.read_bytes()
    Path(result['backup']['path']).write_bytes(b'corrupt')
    with pytest.raises(pub.PublicationError, match='corrupted'):
        pub.undo_last_update(workspace, 'OPS7')
    assert target.read_bytes() == original


def test_explicit_restore_reconciles_state_and_reuses_safety_copy(setup):
    workspace, target = setup
    pub.update_course_notes(workspace, 'OPS7-L01'); original = target.read_bytes()
    revise(workspace, 'New revision.')
    result = pub.update_course_notes(workspace, 'OPS7-L01'); revised = target.read_bytes()
    restored = pub.restore_backup(workspace, result['backup']['backup_id'])
    assert target.read_bytes() == original
    payload = json.loads((workspace.runs_dir / 'OPS7-L01' / 'lecture_notes.candidate.json').read_text())
    assert not payload['published_to_word']
    pub.restore_backup(workspace, restored['pre_restore_backup']['backup_id'])
    assert target.read_bytes() == revised
    assert json.loads((workspace.runs_dir / 'OPS7-L01' / 'lecture_notes.candidate.json').read_text())['published_to_word']


def test_interrupted_undo_recovery_reconciles_state(setup, monkeypatch):
    workspace, target = setup
    pub.update_course_notes(workspace, 'OPS7-L01')
    real = pub._restore_publication_state
    def fail(*args, **kwargs):
        raise RuntimeError('simulated interruption')
    monkeypatch.setattr(pub, '_restore_publication_state', fail)
    with pytest.raises(RuntimeError):
        pub.undo_last_update(workspace, 'OPS7')
    assert not target.exists()
    monkeypatch.setattr(pub, '_restore_publication_state', real)
    records = pub.recover_publications(workspace)
    assert records[0]['status'] == 'UNDONE'
    with sqlite3.connect(workspace.database_path) as conn:
        assert conn.execute('SELECT status FROM lesson_runs').fetchone()[0] == 'ROLLED_BACK'
    assert pub.update_course_notes(workspace, 'OPS7-L01')['status'] == 'PUBLISHED'


def test_undo_reconciles_real_candidate_row_without_another_provider_call(tmp_path, monkeypatch):
    from mynotes.workflow import generate_candidate
    from test_providers_quality import configured_workspace, source_pack, CandidateProvider
    workspace = configured_workspace(tmp_path); source_pack(workspace)
    provider = CandidateProvider()
    monkeypatch.setattr(pub, '_render_and_validate', fake_render)
    assert generate_candidate(workspace, 'OPS7-L01', provider_adapter=provider)['status'] == 'CANDIDATE_READY'
    pub.update_course_notes(workspace, 'OPS7-L01')
    pub.undo_last_update(workspace, 'OPS7')
    path = workspace.runs_dir / 'OPS7-L01' / 'lecture_notes.candidate.json'
    with sqlite3.connect(workspace.database_path) as conn:
        assert conn.execute('SELECT status,candidate_hash FROM candidate_runs').fetchone() == ('CANDIDATE_READY', file_sha256(path))
    assert generate_candidate(workspace, 'OPS7-L01', provider_adapter=provider)['status'] == 'SKIPPED'
    assert pub.update_course_notes(workspace, 'OPS7-L01')['status'] == 'PUBLISHED'
    assert provider.calls == ['writing']


def test_restored_previous_version_can_be_undone_further(setup):
    workspace, target = setup
    pub.update_course_notes(workspace, 'OPS7-L01')
    revise(workspace, 'Second version.')
    result = pub.update_course_notes(workspace, 'OPS7-L01')
    pub.restore_backup(workspace, result['backup']['backup_id'])
    assert pub.undo_last_update(workspace, 'OPS7')['status'] == 'UNDONE'
    assert not target.exists()


@pytest.mark.parametrize('action', ['undo', 'restore'])
def test_restore_cannot_race_with_another_lessons_generation(setup, action):
    from mynotes.workflow import _generation_lock, WorkflowError
    workspace, target = setup
    pub.update_course_notes(workspace, 'OPS7-L01')
    add_second(workspace)
    result = pub.update_course_notes(workspace, 'OPS7-L02')
    before = target.read_bytes()
    with _generation_lock(workspace, 'OPS7-L01'):
        with pytest.raises(WorkflowError, match='already running'):
            if action == 'undo':
                pub.undo_last_update(workspace, 'OPS7')
            else:
                pub.restore_backup(workspace, result['backup']['backup_id'])
    assert target.read_bytes() == before
