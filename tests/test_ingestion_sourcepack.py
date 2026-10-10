from __future__ import annotations

import os
import json
import time
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches


def _create_test_pdf(path: Path, text: str = "Demand planning and a chart", draw_paths: int = 6) -> None:
    lines_pdf = [f"72 {100 + offset * 12} m 250 {130 + offset * 8} l S\n" for offset in range(draw_paths)]
    drawings = "".join(lines_pdf).encode("ascii")
    stream_content = b"BT\n/F1 14 Tf\n72 300 Td\n(" + text.encode("utf-8") + b") Tj\nET\n" + drawings
    pdf_bytes = (
        b"%PDF-1.4\n"
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n"
        b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 400 400] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>\nendobj\n"
        b"4 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n"
        b"5 0 obj\n<< /Length " + str(len(stream_content)).encode("ascii") + b" >>\nstream\n" + stream_content + b"\nendstream\nendobj\n"
        b"xref\n0 6\n0000000000 65535 f \n"
        b"0000000009 00000 n \n0000000058 00000 n \n0000000115 00000 n \n0000000234 00000 n \n0000000305 00000 n \n"
        b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n" + str(450 + len(stream_content)).encode("ascii") + b"\n%%EOF\n"
    )
    path.write_bytes(pdf_bytes)

from mynotes.asr import EngineResult, transcribe_audio
from mynotes.configuration import course_template
from mynotes.extraction import extract_material
from mynotes.ingestion import discover_recordings, parse_recording_name
from mynotes.sourcepack import build_source_pack, discover_materials, validate_source_pack
from mynotes.workspace import Workspace


class FakeEngine:
    name = "fake"

    def __init__(self, text: str = "inventory optimization reduces total cost") -> None:
        self.calls = 0
        self.text = text

    def transcribe(self, audio, output_dir, *, model, language, vocabulary, prompt, clips=None):
        self.calls += 1
        return EngineResult(
            "fake", model, "test-1",
            [{"start": 0.0, "end": 4.0, "text": self.text, "confidence": 0.96, "words": [{"word": "inventory", "start": 0.0, "end": 0.8, "confidence": 0.95}]}],
        )


def configured_workspace(tmp_path: Path, *, code: str = "OPS7", aliases: list[str] | None = None) -> tuple[Workspace, dict]:
    workspace = Workspace.from_value(tmp_path / "private")
    workspace.initialize()
    configuration = workspace.configuration_store().load()
    course = course_template()
    course.update({"code": code, "name": "Operations Analytics", "aliases": aliases or ["ops analytics"]})
    course["transcription"].update(
        {"primary_engine": "fake", "primary_model": "fake-large", "language": "en-IN", "accent": "Indian English"}
    )
    configuration["courses"] = [course]
    configuration["ingestion"]["stable_seconds"] = 90
    workspace.configuration_store().save(configuration)
    return workspace, workspace.configuration_store().load()


def test_dynamic_course_alias_and_safe_canonical_rename(tmp_path) -> None:
    workspace, configuration = configured_workspace(tmp_path)
    recordings = tmp_path / "recordings"
    recordings.mkdir()
    source = recordings / "ops analytics lecture 2.M4A"
    source.write_bytes(b"audio")
    os.utime(source, (time.time() - 120, time.time() - 120))
    configuration["paths"]["recording_dirs"] = [str(recordings)]
    workspace.configuration_store().save(configuration)

    assert parse_recording_name(source, configuration) == ("OPS7", 2)
    dry_run = discover_recordings(workspace, configuration, apply_renames=False)
    assert dry_run[0]["status"] == "RENAME_AVAILABLE"
    assert source.exists()
    applied = discover_recordings(workspace, configuration, apply_renames=True)
    assert applied[0]["status"] == "RENAMED"
    assert (recordings / "OPS7_L02.m4a").is_file()
    assert len(applied[0]["sha256"]) == 64


def test_partial_and_duplicate_recordings_are_not_accepted(tmp_path) -> None:
    workspace, configuration = configured_workspace(tmp_path)
    recordings = tmp_path / "recordings"
    recordings.mkdir()
    fresh = recordings / "OPS7_L01.m4a"
    fresh.write_bytes(b"partial")
    old = recordings / "OPS7_L02.m4a"
    duplicate = recordings / "ops analytics lecture 2.wav"
    old.write_bytes(b"one")
    duplicate.write_bytes(b"two")
    past = time.time() - 200
    os.utime(old, (past, past))
    os.utime(duplicate, (past, past))
    configuration["paths"]["recording_dirs"] = [str(recordings)]
    workspace.configuration_store().save(configuration)

    rows = discover_recordings(workspace, configuration, now=time.time())
    by_name = {Path(item["path"]).name: item for item in rows}
    assert by_name["OPS7_L01.m4a"]["status"] == "WAITING_DOWNLOAD"
    duplicate_rows = [item for item in rows if item.get("lecture_id") == "OPS7-L02"]
    assert duplicate_rows and all(item["status"] == "NEEDS_REVIEW" for item in duplicate_rows)


def test_asr_configuration_is_forwarded_and_exact_cache_is_reused(tmp_path) -> None:
    workspace, configuration = configured_workspace(tmp_path)
    course = configuration["courses"][0]
    course["transcription"]["vocabulary"] = ["newsvendor", "safety stock"]
    audio = tmp_path / "OPS7_L01.wav"
    audio.write_bytes(b"not-real-audio-but-engine-is-injected")
    engine = FakeEngine()

    first = transcribe_audio(workspace, "OPS7-L01", course, audio, engines={"fake": engine})
    second = transcribe_audio(workspace, "OPS7-L01", course, audio, engines={"fake": engine})
    assert engine.calls == 1
    assert first["language"] == "en-IN"
    assert first["accent_hint"] == "Indian English"
    assert first["passes"][0] == {"engine": "fake", "model": "fake-large", "version": "test-1"}
    assert second["cache_reused"] is True
    assert second["segments"][0]["words"][0]["word"] == "inventory"


def test_pdf_pptx_and_lesson_scoped_docx_extraction(tmp_path) -> None:
    pdf = tmp_path / "L1.pdf"
    _create_test_pdf(pdf, text="Demand planning and a chart", draw_paths=6)
    pdf_chunks, pdf_visuals = extract_material(pdf, visual_dir=tmp_path / "pdf-images")
    assert pdf_chunks[0]["page"] == 1
    assert pdf_visuals[0]["visual_decision"] == "candidate"
    assert Path(pdf_visuals[0]["image_path"]).is_file()

    pptx = tmp_path / "Lecture 1.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(2))
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(4), Inches(3), Inches(1)).table
    for row, values in zip(table.rows, (("Month", "Demand"), ("1", "97"), ("2", "90"))):
        for cell, value in zip(row.cells, values):
            cell.text = value
    presentation.save(pptx)
    ppt_chunks, ppt_visuals = extract_material(pptx, visual_dir=tmp_path / "ppt-images")
    assert "Month | Demand\n1 | 97\n2 | 90" in ppt_chunks[0]["content"]
    assert ppt_visuals[0]["signals"]["groups"] == 0

    notes = tmp_path / "notes.docx"
    word = Document()
    word.add_heading("Lecture 1", level=2)
    word.add_paragraph("First lesson note")
    word.add_heading("L2 - second", level=2)
    word.add_paragraph("Second lesson note")
    word.save(notes)
    chunks, _ = extract_material(notes, lesson_number=2)
    assert any("Second lesson" in item["content"] for item in chunks)
    assert all("First lesson" not in item["content"] for item in chunks)


def test_source_pack_labels_evidence_directives_visual_alignment_and_changes(tmp_path) -> None:
    workspace, configuration = configured_workspace(tmp_path)
    audio = tmp_path / "OPS7_L01.wav"
    audio.write_bytes(b"audio")
    slides = tmp_path / "OPS7 L1 slides.pdf"
    _create_test_pdf(slides, text="Inventory optimization and total cost", draw_paths=4)
    notes = tmp_path / "personal.docx"
    word = Document()
    word.add_heading("Lecture 1", level=2)
    word.add_paragraph("NOTE AI: Explain the trade-off clearly")
    word.add_paragraph("This Italian note is high-priority evidence.")
    word.save(notes)

    engine = FakeEngine()
    pack = build_source_pack(
        workspace, "OPS7-L01", audio=audio, materials=[slides], personal_notes=notes, engines={"fake": engine}
    )
    validate_source_pack(pack)
    assert {item["id"][0] for item in pack["evidence"]} == {"T", "S", "H"}
    assert pack["directives"][0]["kind"] == "note"
    assert pack["security"]["source_content_is_untrusted_evidence"] is True
    assert pack["visuals"][0]["evidence_id"].startswith("S")
    assert pack["visuals"][0]["alignment"]["status"] in {"matched", "review"}
    assert pack["availability"]["missing_audio"] is False

    word = Document(notes)
    word.add_paragraph("A later correction")
    word.save(notes)
    updated = build_source_pack(
        workspace, "OPS7-L01", audio=audio, materials=[slides], personal_notes=notes, engines={"fake": engine}
    )
    note_source = next(item for item in updated["sources"] if item["role"] == "personal_notes")
    assert note_source["modified_since_last_prepare"] is True


def test_material_only_source_pack_reports_missing_audio(tmp_path) -> None:
    workspace, _ = configured_workspace(tmp_path)
    material = tmp_path / "Lecture 3.md"
    material.write_text("A complete material-only lesson.", encoding="utf-8")
    pack = build_source_pack(workspace, "OPS7-L03", materials=[material])
    assert pack["availability"]["missing_audio"] is True
    assert pack["availability"]["transcript"] is False
    assert pack["evidence"][0]["id"] == "S001"
    assert any(item["lecture_id"] == "OPS7-L03" and item["status"] == "PREPARED" for item in __import__("mynotes.workflow", fromlist=["status"]).status(workspace))


def test_imported_long_transcript_keeps_speakers_and_explicit_nested_sources(tmp_path) -> None:
    workspace, configuration = configured_workspace(tmp_path)
    course = configuration["courses"][0]
    slides = tmp_path / "course" / "slide"
    slides.mkdir(parents=True)
    material = slides / "L3 inventory.md"
    material.write_text("Inventory decisions under uncertain demand.", encoding="utf-8")
    unrelated = slides / "L3_COURSE_NOTES.docx"
    Document().save(unrelated)
    notebook = tmp_path / "course" / "MY_NOTES.docx"
    document = Document()
    document.add_heading("Lecture 3", level=2)
    document.add_paragraph("Review the in-class example.")
    document.save(notebook)
    course["folders"]["materials"] = str(tmp_path / "course")
    course["lesson_sources"] = {"3": {"materials": [str(material)], "personal_notes": str(notebook)}}
    workspace.configuration_store().save(configuration)

    transcript = tmp_path / "transcript.json"
    transcript.write_text(json.dumps({"language": "en", "segments": [
        {"id": f"T{number:03d}", "start_seconds": float(number),
         "end_seconds": float(number) + 0.5, "text": f"Explanation {number}",
         "speaker": "Professor" if number % 2 else "Student 1"}
        for number in range(1, 1002)
    ]}), encoding="utf-8")

    materials, selected_notebook = discover_materials(workspace.configuration_store().load()["courses"][0], 3)
    assert materials == [material]
    assert selected_notebook == notebook
    pack = build_source_pack(workspace, "OPS7-L03", transcript_path=transcript)
    validate_source_pack(pack)
    assert len([row for row in pack["evidence"] if row["id"].startswith("T")]) == 1001
    assert next(row for row in pack["evidence"] if row["id"] == "T1001")["speaker_role"] == "Professor"
    assert {Path(row["path"]).name for row in pack["sources"]} == {material.name, notebook.name, transcript.name}
    assert pack["transcription"]["imported"] is True


def test_excel_formulas_keep_cell_and_cached_value_provenance(tmp_path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Model"
    sheet["A1"] = 12
    sheet["B1"] = "=A1*2"
    path = tmp_path / "L3 exercise.xlsx"
    book.save(path)

    chunks, visuals = extract_material(path)
    assert visuals == []
    assert chunks[0]["sheet"] == "Model"
    assert chunks[0]["row"] == 1
    assert "A1: 12" in chunks[0]["content"]
    assert "B1: formula =A1*2; cached value not available" in chunks[0]["content"]
