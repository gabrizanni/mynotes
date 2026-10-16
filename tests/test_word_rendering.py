"""Regression fixtures for native tables, escaped content and inherited margins."""
from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm

from mynotes.publication import (
    PublicationError, _append_rows, _append_text, _configure_document,
    _docx_structure, _legacy_table_text, build_word_candidate,
)
from mynotes.utils import atomic_json_write, file_sha256
from test_publication import fake_render, prepared


# Exact seven-column data shown in the user's Word screenshot. This fixture is
# a formatting regression, not an endorsement of the source's approximations.
NEWSVENDOR_ROWS = [
    ["Standard Deviation (σ)", "Order Quantity (Q*)", "Expected Sales",
     "Expected Lost Sales (E[L_S])", "Expected Leftovers (E[L_L])",
     "Expected Profit (E[π])", "Mismatch Cost"],
    ["300", "1000", "880", "120", "120", "Rs. 30,400", "Rs. 9,600"],
    ["250", "1000", "900", "100", "100", "Rs. 32,000", "Rs. 8,000"],
    ["200", "1000", "920", "80", "80", "Rs. 33,600", "Rs. 6,400"],
    ["150", "1000", "940", "60", "60", "Rs. 35,200", "Rs. 4,800"],
    ["100", "1000", "960", "40", "40", "Rs. 36,800", "Rs. 3,200"],
    ["50", "1000", "980", "20", "20", "Rs. 38,400", "Rs. 1,600"],
    ["0 (Deterministic)", "1000", "1000", "0", "0", "Rs. 40,000", "Rs. 0"],
]


def markdown_table(rows: list[list[str]]) -> str:
    return "\n".join("| " + " | ".join(row) + " |" for row in [rows[0], [":---"] * len(rows[0]), *rows[1:]])


def set_notes(workspace, blocks):
    path = workspace.runs_dir / "OPS7-L01" / "lecture_notes.candidate.json"
    payload = json.loads(path.read_text())
    payload["notes"]["blocks"] = blocks
    atomic_json_write(path, payload)


@pytest.mark.parametrize("legacy", [False, True])
def test_screenshot_table_is_native_rectangular_and_fits_printable_width(tmp_path, monkeypatch, legacy):
    workspace, _ = prepared(tmp_path)
    if legacy:
        value = ("Consider an item with mean demand 1000:\n\n" + markdown_table(NEWSVENDOR_ROWS)
                 + "\n\nAs σ approaches zero, expected profit converges to Rs. 40,000.")
        blocks = [{"kind": "paragraph", "text": value.replace("\n", r"\n")}]
    else:
        blocks = [{"kind": "table", "label": "Newsvendor comparison", "rows": NEWSVENDOR_ROWS}]
    set_notes(workspace, blocks)
    monkeypatch.setattr("mynotes.publication._render_and_validate", fake_render)
    built = build_word_candidate(workspace, "OPS7-L01")
    document = Document(built["candidate"])
    assert len(document.tables) == 1
    table = document.tables[0]
    assert [[cell.text for cell in row.cells] for row in table.rows] == NEWSVENDOR_ROWS
    assert table.autofit is False
    assert table.alignment == WD_TABLE_ALIGNMENT.CENTER
    section = document.sections[0]
    assert sum(column.width for column in table.columns) <= section.page_width - section.left_margin - section.right_margin
    assert table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    assert all(row._tr.xpath("./w:trPr/w:cantSplit") for row in table.rows)
    assert built["structure"]["lesson_tables"] == 1
    assert all(r"\n" not in paragraph.text and "| :---" not in paragraph.text for paragraph in document.paragraphs)


@pytest.mark.parametrize("rows", [
    [["A", "B"], ["1"]], [["A"], ["1", "2"]], [["A"], [42]],
])
def test_renderer_rejects_invalid_matrix_without_inventing_cells(rows):
    document = Document()
    _configure_document(document, {}, set_page_layout=True)
    with pytest.raises(PublicationError, match="rows|cells"):
        _append_rows(document, rows)
    assert not document.tables


@pytest.mark.parametrize("text", [
    "| A | B |\n| :--- | :--- |\n| 1 |",
    "| A | B |\n| 1 | 2 |",
    "| A | B |\n| :--- | :--- |\n| 1 | 2",
    r"Summary | A | B |\n| :--- | :--- |\n| 1 | 2 |",
])
def test_ambiguous_markdown_table_blocks_validation(tmp_path, monkeypatch, text):
    workspace, target = prepared(tmp_path)
    set_notes(workspace, [{"kind": "paragraph", "text": text}])
    monkeypatch.setattr("mynotes.publication._render_and_validate", fake_render)
    with pytest.raises(PublicationError, match="table|Markdown"):
        build_word_candidate(workspace, "OPS7-L01")
    assert not target.exists()
    assert not list((workspace.runs_dir / "OPS7-L01" / "word").glob("*.validation.json"))


@pytest.mark.parametrize("value", [
    r"The commands are \nu and \neq.",
    r"Open C:\new\notes.docx or .\new\notes.docx.",
    r"Keep `relative\notes.docx` unchanged.",
    r"Math $x \neq y$ and $\nu = 1$ is preserved.",
])
def test_normalization_preserves_math_commands_and_literal_paths(value):
    assert _legacy_table_text(value) == value


def test_native_table_cells_allow_emphasis_math_and_paths():
    document = Document()
    _configure_document(document, {}, set_page_layout=True)
    _append_rows(document, [["Value", "Location"], ["**Value** $x = 1$", r"C:\new\notes.docx"]])
    assert document.tables[0].cell(1, 1).text == r"C:\new\notes.docx"
    assert document.tables[0].cell(1, 0)._tc.xpath(".//m:oMath")


def test_new_candidate_changes_all_section_margins_and_keeps_baseline(tmp_path, monkeypatch):
    workspace, target = prepared(tmp_path)
    target.parent.mkdir(parents=True)
    baseline = Document()
    baseline.add_paragraph("Historical notes with a literal malformed table | :--- | :--- |")
    baseline.add_section(WD_SECTION_START.NEW_PAGE)
    baseline.add_paragraph("Second historical section")
    for section in baseline.sections:
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(2.5)
    baseline.save(target)
    before = file_sha256(target)
    set_notes(workspace, [{"kind": "table", "rows": NEWSVENDOR_ROWS}])
    monkeypatch.setattr("mynotes.publication._render_and_validate", fake_render)
    built = build_word_candidate(workspace, "OPS7-L01")
    assert file_sha256(target) == before
    document = Document(built["candidate"])
    assert len(document.sections) == 2
    assert document.paragraphs[0].text == baseline.paragraphs[0].text
    assert all(abs(margin.cm - 1.5) < 0.01 for section in document.sections for margin in
               (section.top_margin, section.bottom_margin, section.left_margin, section.right_margin))
    with zipfile.ZipFile(built["candidate"]) as archive:
        from lxml import etree
        xml = etree.fromstring(archive.read("word/document.xml"))
        namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        for margins in xml.xpath(".//w:sectPr/w:pgMar", namespaces=namespaces):
            assert all(abs(int(margins.get("{" + namespaces["w"] + "}" + side)) - 850) <= 1
                       for side in ("top", "bottom", "left", "right"))


def test_structure_checks_cells_but_ignores_historical_lesson_defects(tmp_path):
    document = Document()
    _configure_document(document, {}, set_page_layout=True)
    document.add_paragraph(r"Historical \n | :--- | :--- |")
    document.add_paragraph("MYNOTES START OPS7-L01")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = r"Broken \n table"
    document.add_paragraph("MYNOTES END OPS7-L01")
    path = tmp_path / "cells.docx"
    document.save(path)
    with pytest.raises(PublicationError, match="Markdown|newlines"):
        _docx_structure(path, "OPS7-L01", 0, False)
    table.cell(0, 0).text = "Verified value"
    document.save(path)
    assert _docx_structure(path, "OPS7-L01", 0, False)["lesson_tables"] == 1


def test_table_escaped_pipe_stays_in_single_cell():
    document = Document()
    _configure_document(document, {}, set_page_layout=True)
    _append_text(document, "| A | B |\n| --- | --- |\n| x \\| y | 2 |")
    assert document.tables[0].cell(1, 0).text == "x | y"


def test_sentence_colon_before_escaped_table_is_not_a_windows_drive():
    document = Document()
    _configure_document(document, {}, set_page_layout=True)
    _append_text(document, r"Comparison:\n\n| A | B |\n| --- | --- |\n| 1 | 2 |")
    assert len(document.tables) == 1
    assert document.paragraphs[0].text == "Comparison:"


def test_existing_cumulative_without_builtin_list_styles_gets_native_definitions(tmp_path,monkeypatch):
    workspace,target=prepared(tmp_path)
    target.parent.mkdir(parents=True,exist_ok=True)
    document=Document()
    for name in ('List Number','List Bullet','Caption','Intense Quote'):
        if name in document.styles:document.styles[name].delete()
    document.add_paragraph('Existing academic text remains intact.')
    document.save(target);before=target.read_bytes()
    monkeypatch.setattr('mynotes.publication._render_and_validate',fake_render)
    built=build_word_candidate(workspace,'OPS7-L01')
    assert target.read_bytes()==before
    candidate=Document(built['candidate'])
    assert 'Existing academic text remains intact.' in [p.text for p in candidate.paragraphs]
    for name in ('List Number','List Bullet','Caption','Intense Quote'):
        assert name in candidate.styles
    for name in ('List Number','List Bullet'):
        num_id=candidate.styles[name].element.pPr.numPr.numId.val
        assert any(n.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId')==str(num_id) for n in candidate.part.numbering_part.element)


@pytest.mark.parametrize('wrapper', ['**', '*'])
def test_nested_emphasis_preserves_technical_terms_without_visible_markers(tmp_path, monkeypatch, wrapper):
    workspace, _ = prepared(tmp_path)
    set_notes(workspace, [{'kind': 'paragraph', 'text':
        'Automated ' + wrapper + '[[term:Pseudo-Labels]]' + wrapper + ' support learning.'}])
    monkeypatch.setattr('mynotes.publication._render_and_validate', fake_render)
    built = build_word_candidate(workspace, 'OPS7-L01')
    document = Document(built['candidate'])
    paragraph = next(p for p in document.paragraphs if 'Automated' in p.text)
    assert paragraph.text == 'Automated Pseudo-Labels support learning.'
    term = next(r for r in paragraph.runs if r.text == 'Pseudo-Labels')
    assert term.style.name == 'Termine tecnico'


def test_regression_and_regularization_formulas_remain_native_and_editable():
    from mynotes.publication import append_equation
    document = Document()
    _configure_document(document, {}, set_page_layout=True)
    append_equation(document, r'\hat{y} = b + w_1 x_1')
    append_equation(document, r'z = \ln\left(\frac{p}{1 - p}\right)')
    append_equation(document, r'\mathcal{L}_{L2} = \frac{1}{N} \sum_{i=1}^N (y_i - \hat{y}_i)^2 + \lambda \sum_{j=1}^P \beta_j^2')
    xml = document._element
    assert len(xml.xpath('.//m:acc')) == 2
    from docx.oxml.ns import qn
    assert all(acc.find(qn('m:accPr')).find(qn('m:chr')) is not None for acc in xml.xpath('.//m:acc'))
    assert xml.xpath('.//m:scr[@m:val="script"]')
    assert len(xml.xpath('.//m:f')) == 2
    assert 'ln' in ''.join(xml.xpath('.//m:t/text()'))
    assert '(' in ''.join(xml.xpath('.//m:t/text()'))
    append_equation(document, r'\binom{N}{2} + \ell_1')
    assert xml.xpath('.//m:fPr/m:type[@m:val="noBar"]')
    assert xml.xpath('.//m:dPr/m:begChr[@m:val="("]')
    assert 'ℓ' in ''.join(xml.xpath('.//m:t/text()'))
    append_equation(document, r'\sum_{i \in \text{node}} y_i')
    assert '∈' in ''.join(xml.xpath('.//m:t/text()'))
    append_equation(document, r'\sum_{k \ne j} p_k')
    assert '≠' in ''.join(xml.xpath('.//m:t/text()'))
    append_equation(document, r'p_i \propto e^{z_i / T}')
    assert '∝' in ''.join(xml.xpath('.//m:t/text()'))


def test_bold_wildcard_parameter_names_are_not_mistaken_for_markdown_residue(tmp_path,monkeypatch):
    workspace,_=prepared(tmp_path)
    set_notes(workspace,[{'kind':'bullets','items':["Increase **'min_*' parameters or decrease 'max_*' parameters**."]}])
    monkeypatch.setattr('mynotes.publication._render_and_validate',fake_render)
    built=build_word_candidate(workspace,'OPS7-L01')
    paragraph=next(p for p in Document(built['candidate']).paragraphs if 'min_*' in p.text)
    assert paragraph.text=="Increase 'min_*' parameters or decrease 'max_*' parameters."
    assert any(r.bold and 'min_*' in r.text for r in paragraph.runs)


def test_percentage_math_delimiters_do_not_capture_intervening_prose():
    from mynotes.publication import _inline_math_tokens
    value=r'Female survival is $74\%$, whereas male survival is under $19\%$. Missing `Age` values total $177$ ($19.87\%$).'
    assert _inline_math_tokens(value)==[r'$74\%$',r'$19\%$',r'$177$',r'$19.87\%$']
    document=Document();_configure_document(document,{},set_page_layout=True)
    _append_text(document,value)
    math_text=''.join(document._element.xpath('.//m:t/text()'))
    assert 'whereas' not in math_text and 'Missing' not in math_text
    assert 'whereas male survival' in document.paragraphs[0].text


def test_unambiguous_bare_bounded_sum_renders_as_native_math_but_code_stays_literal():
    document=Document();_configure_document(document,{},set_page_layout=True)
    expression='sum_{j=1}^{|V|} e^{z_j}'
    _append_text(document,'Normalize by the vocabulary ('+expression+'). Keep `('+expression+')` literal and field_name unchanged.')
    paragraph=document.paragraphs[0]
    assert len(paragraph._p.xpath('.//m:oMath'))==1
    assert '∑' in ''.join(paragraph._p.xpath('.//m:t/text()'))
    assert paragraph.text.count(expression)==1
    assert 'Keep ('+expression+') literal and field_name unchanged.' in paragraph.text


def test_serialized_quote_escapes_are_removed_in_prose_but_retained_in_code():
    document=Document();_configure_document(document,{},set_page_layout=True)
    _append_text(document,r'Select \"Apple\"; keep `\"Apple\"` as code.')
    assert document.paragraphs[0].text=='Select "Apple"; keep '+r'\"Apple\"'+' as code.'


def test_native_visual_leaves_clearance_after_a_table(tmp_path):
    from PIL import Image
    from mynotes.publication import _append_visual,_append_rows
    image=tmp_path/'source.png';Image.new('RGB',(100,50),'white').save(image)
    document=Document();_configure_document(document,{},set_page_layout=True)
    _append_rows(document,[['Token','Probability'],['Apple','0.25']])
    _append_visual(document,{'image_path':str(image)},{'reason':'Source diagram'})
    picture=next(p for p in document.paragraphs if p._p.xpath('.//w:drawing'))
    assert picture.paragraph_format.space_before.pt==6
    assert picture.paragraph_format.keep_with_next is True


def test_imported_numbering_is_schema_ordered_and_new_overviews_are_independent():
    from mynotes.publication import _new_numbered_list
    from docx.oxml.ns import qn
    document=Document()
    for name in ('List Number','List Bullet'):document.styles[name].delete()
    _configure_document(document,{},set_page_layout=True)
    first=_new_numbered_list(document);second=_new_numbered_list(document)
    numbering=document.part.numbering_part.element;seen_num=False
    for item in numbering:
        if item.tag==qn('w:num'):seen_num=True
        if item.tag==qn('w:abstractNum'):assert not seen_num
    ids=[]
    for num_id in (first,second):
        number=next(n for n in numbering if n.tag==qn('w:num') and n.get(qn('w:numId'))==str(num_id))
        abstract_id=number.find(qn('w:abstractNumId')).get(qn('w:val'));ids.append(abstract_id)
        abstract=next(n for n in numbering if n.tag==qn('w:abstractNum') and n.get(qn('w:abstractNumId'))==abstract_id)
        assert not list(abstract.iter(qn('w:pStyle')))
        assert number.find(qn('w:lvlOverride')).find(qn('w:startOverride')).get(qn('w:val'))=='1'
    assert len(set(ids))==2
