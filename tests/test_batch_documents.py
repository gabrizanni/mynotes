from copy import deepcopy
import json
from pathlib import Path
import zipfile
import pytest
from docx import Document
from mynotes import batch_documents as bd,publication as pub
from mynotes.utils import atomic_json_write,file_sha256
from test_publication import prepared,fake_render


def test_standalone_has_no_historical_lesson_and_update_keeps_baseline(tmp_path,monkeypatch):
    ws,target=prepared(tmp_path)
    target.parent.mkdir(parents=True);old=Document();old.add_paragraph('Historical content preserved only in update mode');old.save(target)
    before=target.read_bytes();monkeypatch.setattr(pub,'_render_and_validate',fake_render)
    standalone=pub.build_word_candidate(ws,'OPS7-L01',mode='standalone')
    update=pub.build_word_candidate(ws,'OPS7-L01')
    assert not any('Historical' in p.text for p in Document(standalone['candidate']).paragraphs)
    assert any('Historical' in p.text for p in Document(update['candidate']).paragraphs)
    assert standalone['baseline_hash'] is None and standalone['mode']=='standalone'
    assert target.read_bytes()==before
    with pytest.raises(pub.PublicationError,match='Standalone'):
        pub.publish_word_candidate(ws,'OPS7-L01',Path(standalone['candidate']))


def batch_fixture(tmp_path,monkeypatch):
    ws,target=prepared(tmp_path)
    source=ws.runs_dir/'OPS7-L01/lecture_notes.candidate.json'
    first=json.loads(source.read_text());first['batch_id']='OPS7-0123456789ab';atomic_json_write(source,first)
    second=deepcopy(first);second['notes']['lecture_id']='OPS7-L02';second['notes']['title']='Second distinct topic'
    run=ws.runs_dir/'OPS7-L02';run.mkdir();pack=json.loads(Path(first['source_pack']).read_text());pack['lecture_id']='OPS7-L02';atomic_json_write(run/'source_pack.json',pack)
    second.update(source_pack=str(run/'source_pack.json'),source_pack_hash=file_sha256(run/'source_pack.json'));atomic_json_write(run/'lecture_notes.candidate.json',second)
    manifest={'batch_id':first['batch_id'],'course_code':'OPS7','lectures':['OPS7-L01','OPS7-L02']}
    monkeypatch.setattr('mynotes.batch.load_batch_manifest',lambda *_:manifest)
    monkeypatch.setattr('mynotes.batch_recovery.batch_status',lambda *_:{'lessons':[{'content_validated':True}]*2})
    monkeypatch.setattr(pub,'_render_and_validate',fake_render);monkeypatch.setattr(bd,'_render_and_validate',fake_render)
    return ws,manifest['batch_id']


def test_fresh_cumulative_has_each_lesson_once_and_cached_delivery_is_reused(tmp_path,monkeypatch):
    ws,batch=batch_fixture(tmp_path,monkeypatch)
    result=bd.build_batch_documents(ws,batch,mode='both')
    assert len(result['files'])==3 and not result['published'] and not result['provider_invoked']
    cum=next(r for r in result['files'] if r['kind']=='cumulative')
    with zipfile.ZipFile(cum['path']) as z:
        xml=z.read('word/document.xml')
        for lesson in ['OPS7-L01','OPS7-L02']:
            assert xml.count(('MYNOTES START '+lesson).encode())==1
            assert xml.count(('MYNOTES END '+lesson).encode())==1
        assert xml.index(b'MYNOTES START OPS7-L01')<xml.index(b'MYNOTES START OPS7-L02')
    before={r['path']:file_sha256(Path(r['path'])) for r in result['files']}
    assert bd.build_batch_documents(ws,batch)['reused']
    assert all(file_sha256(Path(p))==h for p,h in before.items())


def test_incomplete_batch_cannot_create_a_misleading_complete_delivery(tmp_path,monkeypatch):
    ws,batch=batch_fixture(tmp_path,monkeypatch)
    monkeypatch.setattr('mynotes.batch_recovery.batch_status',lambda *_:{'lessons':[{'content_validated':False}]})
    with pytest.raises(pub.PublicationError,match='All batch lessons'):bd.build_batch_documents(ws,batch)
    assert not (ws.runs_dir/'_batches'/batch/'documents').exists()


def test_changed_course_language_does_not_reuse_stale_documents(tmp_path,monkeypatch):
    ws,batch=batch_fixture(tmp_path,monkeypatch)
    first=bd.build_batch_documents(ws,batch)
    config=ws.configuration_store().load();config['courses'][0]['languages']['notes']='it'
    ws.configuration_store().save(config)
    second=bd.build_batch_documents(ws,batch)
    assert not second['reused'] and first['index']!=second['index']
