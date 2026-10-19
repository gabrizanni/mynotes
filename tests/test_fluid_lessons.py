"""Course continuity, source-bound revisions and preservation of a personal Word."""
from copy import deepcopy
import json
from pathlib import Path

import pytest
from docx import Document

from mynotes.batch import generate_batch, plan_batch, CoverageError, _continuity, _hash
from mynotes.batch_recovery import apply_correction
from mynotes import publication as pub
from mynotes.providers import ProviderResponse
from mynotes.utils import atomic_json_write, file_sha256
from test_batch_generation import workspace_and_deck, Adapter, VERIFIED
from test_publication import prepared, fake_render
from test_note_update_undo import add_second, revise


def add_turn(workspace, lecture, text='The source-supported logistics mechanism is introduced here.'):
    path=workspace.runs_dir/lecture/'source_pack.json';pack=json.loads(path.read_text())
    pack['evidence'].append({'id':'T001','kind':'transcript','text':text,'start_seconds':0,'end_seconds':30})
    pack['availability']['transcript']=True;atomic_json_write(path,pack)


def test_future_slide_can_support_table_formula_and_long_recap(tmp_path):
    ws,_,selection=workspace_and_deck(tmp_path);add_turn(ws,'OPS-L01')
    def transform(response,request):
        row=response['lessons'][0]['notes']
        row['blocks'].insert(2,{'kind':'table','rows':[['Factor','Effect'],['Lead time','More inventory']],
            'evidence_ids':['S021','T001'],'coverage_role':'new_insight','prior_lecture_id':'OPS-L02'})
        row['blocks'].insert(3,{'kind':'formula','label':'Inventory','expression':'I = D L',
            'evidence_ids':['S021','T001'],'coverage_role':'new_insight','prior_lecture_id':'OPS-L02'})
        row['blocks'][-1].update(text='A connected recap of the lecture. '*40,evidence_ids=['S001','S021','T001'])
        next(e for e in row['evidence_coverage'] if e['evidence_id']=='S021')['disposition']='used'
    result=generate_batch(ws,'OPS',[1,2],selection=selection,budgets=VERIFIED,provider_adapter=Adapter(transform))
    assert result['status']=='CANDIDATE_READY'


def test_additional_teaching_use_requires_its_substantive_contribution(tmp_path):
    ws,_,selection=workspace_and_deck(tmp_path);add_turn(ws,'OPS-L01')
    key=next(key for key in selection if key.endswith(':21'))
    selection[key]['uses']={'OPS-L01':{'role':'introduction','reason':'Introduced in the first recording','evidence_ids':['T001']}}
    with pytest.raises(CoverageError):
        generate_batch(ws,'OPS',[1,2],selection=selection,budgets=VERIFIED,provider_adapter=Adapter())
    selection[key]['uses']['OPS-L01']['evidence_ids']=['T999']
    with pytest.raises(Exception,match='unavailable transcript'):
        plan_batch(ws,'OPS',[1,2],selection=selection)


def test_context_lesson_is_imported_with_origins_and_verified_source_pack(tmp_path):
    ws,_,selection=workspace_and_deck(tmp_path);add_turn(ws,'OPS-L02','Context from a different recording.')
    original=(ws.runs_dir/'OPS-L01/source_pack.json').read_bytes()
    # Limit the selection to the actual output lesson, with foreign sources as
    # contextual material rather than pretending its recording happened here.
    one={k:{'owner':'OPS-L01','status':'material'} for k in selection}
    adapter=Adapter()
    result=generate_batch(ws,'OPS',[1],context_lessons=['OPS-L02'],selection=one,provider_adapter=adapter)
    request=adapter.calls[0]
    turn=next(e for e in request.source_pack['evidence'] if e['id'].startswith('T'))
    assert turn['origin_lecture_id']=='OPS-L02' and turn['origin_evidence_id']=='T001'
    candidate=json.loads(Path(result['candidates'][0]['candidate']).read_text())
    effective=json.loads(Path(candidate['source_pack']).read_text())
    assert effective['evidence'][-1]['origin_lecture_id']=='OPS-L02'
    assert (ws.runs_dir/'OPS-L01/source_pack.json').read_bytes()==original


def test_continuity_preserves_complete_passages_and_sources():
    text='A complete source-grounded explanation. '*100
    notes={'blocks':[{'kind':'paragraph','text':text,'evidence_ids':['S015','T045'],'coverage_role':'full'}]}
    context=_continuity('OPS-L01',notes)
    assert context['blocks'][0]['text']==text and context['source_links']==['S015','T045']


def test_reorganization_reuses_original_success_without_claiming_equal_intent(tmp_path):
    ws,_,selection=workspace_and_deck(tmp_path);add_turn(ws,'OPS-L01')
    class Recorded(Adapter):
        def generate(self,request):
            response=super().generate(request)
            path=request.run_dir/'provider-attempts/original.json'
            atomic_json_write(path,{'returncode':0,'error':None,'requested':vars(request.requested),'stdout':'original immutable response'})
            return ProviderResponse(response.notes,response.provider,response.model,response.effort,response.usage,str(path))
    writer=Recorded()
    first=generate_batch(ws,'OPS',[1,2],selection=selection,provider_adapter=writer,pause_after='OPS-L01')
    old_manifest=Path(first['manifest']);old_bytes=old_manifest.read_bytes()
    key=next(iter(selection));selection=deepcopy(selection)
    selection[key].update(reason='Recording context clarifies the introduction',uses={'OPS-L01':{'role':'introduction','reason':'Introduced here','evidence_ids':['T001']}})
    class NoProvider:
        def generate(self,request):pytest.fail('Valid saved writing must be revalidated without a provider call')
    result=generate_batch(ws,'OPS',[1,2],selection=selection,batch_id=first['batch_id'],provider_adapter=NoProvider(),pause_after='OPS-L01')
    candidate=json.loads(Path(result['candidates'][0]['candidate']).read_text())
    assert candidate['source_revalidation']['kind']=='organization_revalidation'
    assert candidate['organization_revision']['changes'] and old_manifest.read_bytes()==old_bytes


def test_hash_anchored_replacement_and_move_preserve_unrelated_blocks():
    a={'kind':'paragraph','text':'Source-grounded passage','evidence_ids':['S001'],'coverage_role':'full'}
    b={'kind':'paragraph','text':'Destination topic','evidence_ids':['S002'],'coverage_role':'full'}
    original={'lessons':[{'lecture_id':'OPS-L01','notes':{'blocks':[a], 'evidence_coverage':[], 'visual_decisions':[]}},
                         {'lecture_id':'OPS-L02','notes':{'blocks':[b], 'evidence_coverage':[], 'visual_decisions':[]}}]}
    task={'defects':[{'lecture_id':'OPS-L01','evidence_id':'S001'}],
          'editable_blocks':[{'lecture_id':'OPS-L01','block_sha256':_hash(a)}],
          'lesson_evidence':{'OPS-L01':['S001'],'OPS-L02':['S001','S002']}}
    moved=apply_correction(original,{'moves':[{'lecture_id':'OPS-L01','block_sha256':_hash(a),'to_lecture_id':'OPS-L02','after_block_sha256':_hash(b)}]},task)
    assert moved['lessons'][1]['notes']['blocks']==[b,a] and original['lessons'][0]['notes']['blocks']==[a]
    edited=apply_correction(original,{'replacements':[{'lecture_id':'OPS-L01','block_sha256':_hash(a),'block':{**a,'text':'Corrected source-grounded passage'}}]},task)
    assert edited['lessons'][1]==original['lessons'][1]
    with pytest.raises(Exception,match='hash-anchored'):
        apply_correction(original,{'deletions':[{'lecture_id':'OPS-L02','block_sha256':_hash(b)}]},task)


def word_text(path):return '\n'.join(p.text for p in Document(path).paragraphs)


def test_multi_update_one_backup_personal_additions_survive_twice_and_undo(tmp_path,monkeypatch):
    ws,target=prepared(tmp_path);monkeypatch.setattr(pub,'_render_and_validate',fake_render)
    pub.update_course_notes(ws,'OPS7-L01');add_second(ws)
    pub.update_course_notes(ws,'OPS7-L02')
    document=Document(target)
    paragraph=next(p for p in document.paragraphs if p.text=='Cost and service interact.')
    paragraph.add_run(' My own example about the campus warehouse.')
    document.save(target);before=target.read_bytes()
    revise(ws,'Revised first lesson.')
    result=pub.update_course_notes(ws,['OPS7-L01','OPS7-L02'])
    assert result['status']=='PUBLISHED' and result['lectures']==['OPS7-L01','OPS7-L02']
    assert 'My own example' in word_text(target) and 'Revised first lesson.' in word_text(target)
    assert Path(result['backup']['path']).read_bytes()==before
    revise(ws,'Further corrected first lesson.')
    pub.update_course_notes(ws,['OPS7-L01','OPS7-L02'])
    assert word_text(target).count('My own example')==1
    pub.undo_last_update(ws,'OPS7');assert 'Revised first lesson.' in word_text(target)


def test_same_paragraph_independent_edits_merge_and_real_conflict_keeps_other_lesson(tmp_path,monkeypatch):
    ws,target=prepared(tmp_path);monkeypatch.setattr(pub,'_render_and_validate',fake_render)
    pub.update_course_notes(ws,'OPS7-L01');add_second(ws);pub.update_course_notes(ws,'OPS7-L02')
    document=Document(target)
    p=next(p for p in document.paragraphs if p.text=='Cost and service interact.')
    p.add_run(' My personal example.');document.save(target)
    path=ws.runs_dir/'OPS7-L01/lecture_notes.candidate.json';payload=json.loads(path.read_text())
    payload['notes']['sections'][0]['body']='Cost and service interact through inventory.';atomic_json_write(path,payload)
    assert pub.update_course_notes(ws,'OPS7-L01')['status']=='PUBLISHED'
    assert 'inventory.' in word_text(target) and 'My personal example.' in word_text(target)
    document=Document(target);p=next(p for p in document.paragraphs if 'inventory.' in p.text)
    p.text='My contradictory personal explanation.';document.save(target)
    payload['notes']['sections'][0]['body']='A different incompatible rewrite.';atomic_json_write(path,payload)
    second=ws.runs_dir/'OPS7-L02/lecture_notes.candidate.json';data=json.loads(second.read_text());data['notes']['summary']='Independent correction.';atomic_json_write(second,data)
    result=pub.update_course_notes(ws,['OPS7-L01','OPS7-L02'])
    assert result['status']=='PARTIALLY_PUBLISHED' and result['lectures']==['OPS7-L02']
    assert 'My contradictory personal explanation.' in word_text(target) and 'Independent correction.' in word_text(target)


def test_concurrent_word_change_and_commit_failure_never_lose_original(tmp_path,monkeypatch):
    from mynotes.course_update import build_course_revision,publish_course_revision
    ws,target=prepared(tmp_path);monkeypatch.setattr(pub,'_render_and_validate',fake_render)
    pub.update_course_notes(ws,'OPS7-L01');revise(ws,'A corrected summary.')
    built=build_course_revision(ws,['OPS7-L01'])
    document=Document(target);document.add_paragraph('Concurrent addition');document.save(target);before=target.read_bytes()
    with pytest.raises(pub.PublicationError,match='changed after validation'):
        publish_course_revision(ws,Path(built['candidate']))
    assert target.read_bytes()==before


@pytest.mark.parametrize('exception',[OSError, SystemExit])
def test_metadata_failure_or_interruption_restores_word_and_all_lessons(tmp_path,monkeypatch,exception):
    from mynotes import course_update as cu
    ws,target=prepared(tmp_path);monkeypatch.setattr(pub,'_render_and_validate',fake_render)
    pub.update_course_notes(ws,'OPS7-L01');add_second(ws);pub.update_course_notes(ws,'OPS7-L02')
    before=target.read_bytes();index=cu.baseline_index(ws,'OPS7')
    revise(ws,'New content retained after the failed commit.')
    original=cu.atomic_json_write
    def fail(path,data):
        if Path(path)==cu._index_path(ws,'OPS7'):raise exception('Commit interrupted')
        return original(path,data)
    monkeypatch.setattr(cu,'atomic_json_write',fail)
    with pytest.raises(exception):pub.update_course_notes(ws,['OPS7-L01','OPS7-L02'])
    monkeypatch.setattr(cu,'atomic_json_write',original)
    if exception is SystemExit:
        result=pub.recover_publications(ws)
        assert any(r['status']=='ROLLED_BACK' for r in result)
    assert target.read_bytes()==before and cu.baseline_index(ws,'OPS7')==index
    assert pub.update_course_notes(ws,['OPS7-L01','OPS7-L02'])['status']=='PUBLISHED'
    assert 'New content retained' in word_text(target)


def test_new_earlier_lesson_is_inserted_in_numeric_order(tmp_path,monkeypatch):
    ws,target=prepared(tmp_path);add_second(ws);monkeypatch.setattr(pub,'_render_and_validate',fake_render)
    pub.update_course_notes(ws,'OPS7-L02');pub.update_course_notes(ws,'OPS7-L01')
    text=word_text(target)
    assert text.index('MYNOTES START OPS7-L01') < text.index('MYNOTES START OPS7-L02')
