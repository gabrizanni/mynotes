from zipfile import ZipFile
import pytest
from docx import Document
from lxml import etree
from mynotes.publication import PublicationError, append_equation


def test_provider_aligned_equation_preserves_three_editable_rows(tmp_path):
    document = Document()
    append_equation(document, r'\begin{aligned}R&=10\times20=200\\TC&=100+(5\times20)=200\\\pi&=R-TC=0\end{aligned}')
    path=tmp_path/'aligned.docx';document.save(path)
    with ZipFile(path) as archive:
        xml=etree.fromstring(archive.read('word/document.xml'))
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    rows=xml.xpath('//m:eqArr/m:e',namespaces=ns)
    assert len(rows)==3
    text=[''.join(row.xpath('.//m:t/text()',namespaces=ns)) for row in rows]
    assert text==['R=10×20=200','TC=100+(5×20)=200','π=R-TC=0']
    assert '&' not in ''.join(text)


def test_aligned_equation_rows_keep_native_fractions_and_text(tmp_path):
    document=Document()
    append_equation(document,r'\begin{aligned}c&=p-v\\Q&=\frac{F}{c}\quad\text{units}\end{aligned}')
    path=tmp_path/'fraction.docx';document.save(path)
    with ZipFile(path) as archive:
        xml=archive.read('word/document.xml').decode()
    assert '<m:eqArr>' in xml and '<m:f>' in xml and '<m:nor' in xml


def test_empty_aligned_row_is_rejected():
    with pytest.raises(PublicationError,match='empty'):
        append_equation(Document(),r'\begin{aligned}x=1\\\end{aligned}')
