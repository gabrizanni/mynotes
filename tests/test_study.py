from copy import deepcopy
from pathlib import Path
import zipfile

from docx import Document
import pytest

from mynotes.configuration import course_template
from mynotes.providers import ProviderResponse, _prompt
from mynotes.study import (StudyError, build_study_candidate, finalize_study_candidate,
                           generate_study_candidate, parse_lessons, plan_study,
                           validate_study_candidate)
from mynotes.utils import atomic_json_write, file_sha256
from mynotes.workspace import Workspace


def fixture(tmp_path):
    workspace = Workspace.from_value(tmp_path / "private")
    workspace.initialize()
    configuration = workspace.configuration_store().load()
    course = course_template()
    course.update({"code": "AI", "name": "Artificial Intelligence"})
    course["folders"]["materials"] = str(tmp_path / "materials")
    course["folders"]["handout"] = str(tmp_path / "handout")
    configuration["courses"] = [course]
    configuration["providers"]["default"] = "codex"
    configuration["providers"]["configured"]["codex"] = {"model": "gpt-5.6-terra", "effort": "high"}
    workspace.configuration_store().save(configuration)
    source = tmp_path / "L6.txt"
    source.write_text("Split data before scaling; fit transformations only on the training fold. Evaluate held-out recall and precision.")
    manifest = {"schema_version": 1, "course": "AI", "lessons": [6, 7],
                "sources": [{"id": "SRC1", "path": str(source), "sha256": file_sha256(source), "kind": "notebook", "lesson_ids": [6, 7], "review_status": "complete"}],
                "evidence": [{"id": "NB_C001", "source_id": "SRC1", "content": source.read_text()}],
                "topics": [{"id": "metrics", "title": "Recall", "evidence_ids": ["NB_C001"], "quantitative": True}],
                "links": [{"id": "link1", "source_id": "SRC1", "url": "https://example.org/metrics", "title": "Metrics"}], "limitations": []}
    candidate = {"schema_version": 1, "course": "AI", "title": "Study guide", "purpose": "quiz", "language": "en", "recap_language": "it",
                 "sections": [{"id": "main", "title": "Metrics and their limits", "role": "content", "evidence_ids": ["NB_C001"], "topic_ids": ["metrics"],
                               "blocks": [{"kind": "paragraph", "text": "[[term:Recall]] measures **the fraction of actual positives found**. False negatives lower recall. A high recall alone does not guarantee precise positive predictions.", "evidence_ids": ["NB_C001"]},
                                          {"kind": "formula", "expression": r"Recall = \frac{TP}{TP+FN}", "label": "Recall", "evidence_ids": ["NB_C001"]},
                                          {"kind": "paragraph", "text": "For 8 true positives and 2 false negatives, recall equals 8/(8+2), or 0.8. This example illustrates the ratio and does not describe a measured class result."},
                                          {"kind": "table", "rows": [["Measure", "Question"], ["Recall", "How many actual positives are found?"]], "evidence_ids": ["NB_C001"]}]},
                              {"id": "recap", "title": "Riepilogo in italiano", "role": "recap", "evidence_ids": ["NB_C001"], "topic_ids": [],
                               "blocks": [{"kind": "paragraph", "text": "Il recall valuta quanti positivi effettivi sono stati identificati."}]},
                              {"id": "practice", "title": "Generated practice exercises", "role": "practice", "evidence_ids": ["NB_C001"], "topic_ids": [],
                               "blocks": [{"kind": "qa", "question": "What happens when false negatives increase?", "answer": "Recall falls if true positives stay constant."}]},
                              {"id": "resources", "title": "Linked resources explained", "role": "resources", "evidence_ids": ["NB_C001"], "topic_ids": [], "link_ids": ["link1"],
                               "blocks": [{"kind": "paragraph", "text": "https://example.org/metrics explains how recall complements precision and links evaluation to the training workflow."}]}],
                 "topic_coverage": [{"topic_id": "metrics", "section_id": "main", "facets": ["definition", "intuition", "mechanism", "limits", "formula", "worked_example"]}],
                 "visual_decisions": [], "visuals": [], "editorial_review": {"complete": True}}
    candidate_path, manifest_path = tmp_path / "candidate.json", tmp_path / "manifest.json"
    atomic_json_write(candidate_path, candidate)
    atomic_json_write(manifest_path, manifest)
    return workspace, candidate, manifest, candidate_path, manifest_path


def build(tmp_path, *, render=False):
    workspace, candidate, manifest, candidate_path, manifest_path = fixture(tmp_path)
    result = build_study_candidate(workspace, "AI", candidate_path, manifest_path, tmp_path / "guide.docx", render=render)
    return workspace, candidate, manifest, candidate_path, manifest_path, result


def fake_render(path, directory):
    directory.mkdir(parents=True, exist_ok=True)
    pdf = directory / "guide.pdf"
    pdf.write_bytes(b"isolated stub")
    return {"pdf": str(pdf), "pages": [{"page": 1, "characters": 500}], "renderer": "test double"}


def test_standalone_uses_canonical_styles_native_math_and_tables_without_cumulative(tmp_path):
    workspace, _, _, _, _, result = build(tmp_path)
    assert result["status"] == "STUDY_STRUCTURE_VALIDATED"
    assert not result["visual_reviewed"] and not result["word_verified"]
    document = Document(result["candidate"])
    assert document.styles["MyNotes Body"].font.name == "Aptos"
    assert document.styles["MyNotes Body"].font.size.pt == 11
    assert document.styles["MyNotes Body"].paragraph_format.line_spacing == 1.15
    assert len(document.tables) == 1
    assert any(r.style.name == "Termine tecnico" for p in document.paragraphs for r in p.runs)
    assert "<m:f>" in document.part.element.xml
    assert all(abs(s.left_margin.cm - 1.5) < .01 for s in document.sections)
    text = " ".join(p.text for p in document.paragraphs)
    assert "**" not in text and "[[term:" not in text
    assert text.index("Riepilogo") < text.index("Linked resources explained")
    assert text.index("Generated practice") < text.index("Linked resources explained")
    assert not list(workspace.runs_dir.glob("*/lecture_notes.candidate.json"))


def test_never_reads_existing_cumulative(tmp_path):
    workspace, candidate, manifest, candidate_path, manifest_path = fixture(tmp_path)
    handout = tmp_path / "handout" / "AI_COURSE_NOTES.docx"
    handout.parent.mkdir()
    handout.write_bytes(b"deliberately not a valid Word: must never be read")
    digest = file_sha256(handout)
    result = build_study_candidate(workspace, "AI", candidate_path, manifest_path, tmp_path / "guide.docx", render=False)
    assert result["status"] == "STUDY_STRUCTURE_VALIDATED" and file_sha256(handout) == digest


@pytest.mark.parametrize("mutation,match", [
    (lambda c,m: c["sections"].reverse(), "resources"),
    (lambda c,m: c["sections"][0].update(topic_ids=[]), "topics"),
    (lambda c,m: c["sections"][0].update(evidence_ids=["missing"]), "Unknown section evidence"),
    (lambda c,m: c["topic_coverage"][0].update(facets=["definition"]), "facets"),
    (lambda c,m: c["sections"][0]["blocks"][3].update(rows=[["a","b"],["c"]]), "rectangular"),
    (lambda c,m: m["sources"][0].update(review_status="excerpt_only"), "Full source review"),
    (lambda c,m: c["sections"][-1].update(link_ids=[]), "every requested link"),
    (lambda c,m: c["sections"][-1]["blocks"][0].update(text="Link omitted"), "exact resource URLs"),
    (lambda c,m: m["topics"].append({"id":"omitted", "evidence_ids":["NB_C001"]}), "omitted topics"),
])
def test_quality_gates_reject_omissions_and_unverified_inputs(tmp_path, mutation, match):
    _, candidate, manifest, _, _ = fixture(tmp_path)
    mutation(candidate, manifest)
    with pytest.raises(StudyError, match=match):
        validate_study_candidate(candidate, manifest)


def test_docx_default_and_original_preservation(tmp_path):
    workspace, candidate, manifest, candidate_path, manifest_path = fixture(tmp_path)
    for output in (tmp_path / "notes.md", tmp_path / "notes.pdf"):
        with pytest.raises(StudyError, match=".docx"):
            build_study_candidate(workspace, "AI", candidate_path, manifest_path, output, render=False)
    output = tmp_path / "existing.docx"
    output.write_bytes(b"original")
    with pytest.raises(StudyError, match="already exists"):
        build_study_candidate(workspace, "AI", candidate_path, manifest_path, output, render=False)
    assert output.read_bytes() == b"original"


def test_ready_requires_real_checks_not_build_success(tmp_path, monkeypatch):
    monkeypatch.setattr("mynotes.study._render_and_validate", fake_render)
    *_, result = build(tmp_path, render=True)
    assert result["status"] == "STUDY_RENDER_VALIDATED"
    result = finalize_study_candidate(Path(result["candidate"]), render=False, visual_reviewed=True)
    assert result["status"] == "STUDY_RENDER_VALIDATED"
    result = finalize_study_candidate(Path(result["candidate"]), render=False, word_verified=True)
    assert result["status"] == "STUDY_READY"


def test_draft_is_not_complete_without_editorial_review(tmp_path, monkeypatch):
    workspace, candidate, manifest, cp, mp = fixture(tmp_path)
    candidate["editorial_review"] = {"complete": False}
    atomic_json_write(cp,candidate)
    monkeypatch.setattr("mynotes.study._render_and_validate", fake_render)
    result = build_study_candidate(workspace,"AI",cp,mp,tmp_path / "guide.docx")
    result = finalize_study_candidate(Path(result["candidate"]),render=False,visual_reviewed=True,word_verified=True)
    assert result["status"] == "STUDY_RENDER_VALIDATED"


def test_final_checks_fail_if_source_changes(tmp_path):
    _, _, manifest, _, _, result = build(tmp_path)
    Path(manifest["sources"][0]["path"]).write_text("changed")
    with pytest.raises(StudyError,match="changed"):
        finalize_study_candidate(Path(result["candidate"]),render=False)


def test_cannot_assert_visual_before_render(tmp_path):
    *_, result = build(tmp_path)
    with pytest.raises(StudyError,match="precede rendering"):
        finalize_study_candidate(Path(result["candidate"]),render=False,visual_reviewed=True)


def test_no_link_request_does_not_require_invented_resources(tmp_path):
    _, candidate, manifest, _, _ = fixture(tmp_path)
    manifest['links'] = []
    candidate['sections'].pop()
    assert validate_study_candidate(candidate, manifest)['topics'] == 1


def test_plan_binds_transcript_text_to_verified_recording_source(tmp_path):
    workspace, _, _, _, _ = fixture(tmp_path)
    audio = tmp_path/'selected.wav'
    audio.write_bytes(b'synthetic recording identity')
    atomic_json_write(workspace.runs_dir/'AI-L06/source_pack.json', {
        'sources':[{'path':str(audio),'sha256':file_sha256(audio),'role':'recording'}],
        'evidence':[{'id':'T001','kind':'transcript','text':'Full lesson explanation.', 'start_seconds':0.0,'end_seconds':1.0}]})
    result = plan_study(workspace,'AI','6')
    assert result['evidence'][0]['content'] == 'Full lesson explanation.'
    assert result['evidence'][0]['source_id'] == result['sources'][0]['id']
    assert result['sources'][0]['kind'] == 'recording'


def test_unavailable_link_requires_truthful_disclosure(tmp_path):
    _, candidate, manifest, _, _ = fixture(tmp_path)
    manifest['links'][0].update(review_status='unavailable', reason='No browser access in this environment')
    with pytest.raises(StudyError,match='disclose'):
        validate_study_candidate(candidate,manifest)
    candidate['sections'][-1]['blocks'][0]['text'] += ' No browser access in this environment'
    validate_study_candidate(candidate,manifest)


def test_editorial_review_can_be_recorded_without_rewriting_valid_word(tmp_path, monkeypatch):
    workspace, candidate, manifest, cp, mp = fixture(tmp_path)
    candidate['editorial_review'] = {'complete':False}
    atomic_json_write(cp,candidate)
    monkeypatch.setattr('mynotes.study._render_and_validate',fake_render)
    result = build_study_candidate(workspace,'AI',cp,mp,tmp_path/'guide.docx')
    digest = file_sha256(Path(result['candidate']))
    result = finalize_study_candidate(Path(result['candidate']),render=False,editorial_reviewed=True,visual_reviewed=True,word_verified=True)
    assert result['status'] == 'STUDY_READY'
    assert file_sha256(Path(result['candidate'])) == digest


def test_provider_uses_selected_effort_and_never_claims_its_draft_reviewed(tmp_path):
    workspace,candidate,manifest,_,mp = fixture(tmp_path)
    requests=[]
    class Fake:
        def generate(self, request):
            requests.append(request)
            return ProviderResponse(candidate,"codex","gpt-5.6-terra","high",{"input_tokens": 100},"synthetic-attempt")
    result=generate_study_candidate(workspace,"AI",mp,tmp_path/"first.docx",provider_adapter=Fake(),render=False)
    assert requests[0].requested.model == "gpt-5.6-terra" and requests[0].requested.effort == "high"
    prompt = _prompt(requests[0])
    assert "final resources section" in prompt and "NB_C001" in prompt
    assert "End with exactly one concise recap" not in prompt
    assert result["editorial_review_required"] and not result["coverage"]["editorial_review_recorded"]
    result=generate_study_candidate(workspace,"AI",mp,tmp_path/"second.docx",provider_adapter=Fake(),render=False)
    assert result["generation_reused"] and len(requests)==1


def test_plan_deduplicates_exact_files_without_claiming_review(tmp_path):
    workspace,_,_,_,_ = fixture(tmp_path)
    materialdir=tmp_path/"materials"; materialdir.mkdir()
    (materialdir/"L6.txt").write_text("Shared content")
    (materialdir/"L7.txt").write_text("Shared content")
    result=plan_study(workspace,"AI","6-7")
    assert len(result["sources"])==1
    assert result["sources"][0]["lesson_ids"]==[6,7]
    assert result["sources"][0]["review_status"]=="pending"
    assert len(result["sources"][0]["aliases"])==2 and result["limitations"]


@pytest.mark.parametrize("value,expected",[("6-10",[6,7,8,9,10]),("8,6,7,6",[6,7,8])])
def test_lesson_selection(value,expected):
    assert parse_lessons(value)==expected
