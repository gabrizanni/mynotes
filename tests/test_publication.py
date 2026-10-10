from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from mynotes.configuration import course_template
from mynotes.publication import (
    PublicationError,
    append_equation,
    build_word_candidate,
    publication_history,
    publish_word_candidate,
    preview_style_document,
    restore_backup,
)
from mynotes.utils import atomic_json_write, file_sha256
from mynotes.workspace import Workspace
from mynotes.workflow import status


def prepared(tmp_path: Path, *, visual: bool = False, formula: str | None = None) -> tuple[Workspace, Path]:
    workspace = Workspace.from_value(tmp_path / "private"); workspace.initialize()
    configuration = workspace.configuration_store().load(); course = course_template()
    course.update({"code": "OPS7", "name": "Operations"}); course["folders"]["handout"] = str(tmp_path / "handouts")
    configuration["courses"] = [course]; workspace.configuration_store().save(configuration)
    run = workspace.runs_dir / "OPS7-L01"; run.mkdir(parents=True)
    image = run / "chart.png"
    if visual:
        Image.new("RGB", (400, 200), "white").save(image)
    pack = {
        "schema_version": 1, "lecture_id": "OPS7-L01", "course": {"code": "OPS7", "name": "Operations", "lesson_number": 1},
        "availability": {}, "evidence": [{"id": "S001", "kind": "slide", "content": "Trade-off", "source_file": "deck.pdf", "page": 2}],
        "visuals": ([{"evidence_id": "S001", "image_path": str(image), "source_file": "deck.pdf", "page": 2}] if visual else []),
        "directives": [], "sources": [], "glossary": [],
    }
    source = run / "source_pack.json"; atomic_json_write(source, pack)
    notes = {
        "lecture_id": "OPS7-L01", "title": "Trade-off", "summary": "Summary", "bullets": [], "table": [], "formula": formula,
        "sections": [{"heading": "Framework", "level": 3, "body": "Cost and service interact.", "bullets": ["Cost", "Service"], "evidence_ids": ["S001"]}],
        "evidence_coverage": [{"evidence_id": "S001", "disposition": "used", "reason": "core"}],
        "visual_decisions": ([{"evidence_id": "S001", "decision": "include", "reason": "Essential chart"}] if visual else []), "model_knowledge": [],
    }
    atomic_json_write(run / "lecture_notes.candidate.json", {"notes": notes, "source_pack": str(source), "source_pack_hash": file_sha256(source), "published_to_word": False})
    return workspace, tmp_path / "handouts" / "OPS7_COURSE_NOTES.docx"


def fake_render(path: Path, output: Path) -> dict:
    output.mkdir(parents=True, exist_ok=True); pdf = output / f"{path.stem}.pdf"; pdf.write_bytes(b"test")
    return {"pdf": str(pdf), "pages": [{"page": 1, "characters": 1, "images": 0}], "renderer": "test"}


def test_equation_is_structured_omml_fraction_root_and_scripts(tmp_path) -> None:
    path = tmp_path / "equation.docx"; document = Document()
    append_equation(document, r"y = \frac{\sqrt{x_1}}{\sum_{i=1}^{n} z_i}"); document.save(path)
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    assert b"<m:f>" in xml and b"<m:rad>" in xml and b"<m:sSubSup>" in xml
    assert b"<m:sSub>" in xml or b"<m:sSubSup>" in xml


def test_equation_supports_distribution_relation(tmp_path) -> None:
    path = tmp_path / "distribution.docx"; document = Document()
    append_equation(document, r"D \sim N(\mu, \sigma)"); document.save(path)
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    assert "∼".encode() in xml


def test_real_agy_mathbb_and_inline_math_stay_native(tmp_path, monkeypatch) -> None:
    # The real Agy sample returned double-escaped commands and inline math.
    formula = r"\\mathbb{E}[\\Delta \\Pi(Q)] = (r - w) \\cdot \\mathbb{P}(D > Q) + (s - w) \\cdot \\mathbb{P}(D \\le Q)"
    workspace, _ = prepared(tmp_path, formula=formula)
    source = workspace.runs_dir / "OPS7-L01" / "lecture_notes.candidate.json"
    payload = json.loads(source.read_text())
    payload["notes"]["sections"][0]["body"] = r"Order when $\\mathbb{P}(D \\le Q^*) = \\frac{C_u}{C_u + C_o}$ holds."
    atomic_json_write(source, payload)
    monkeypatch.setattr("mynotes.publication._render_and_validate", fake_render)
    result = build_word_candidate(workspace, "OPS7-L01")
    document = Document(result["candidate"])
    xml = document.part.element.xml
    assert xml.count('m:val="double-struck"') == 4
    assert "<m:f>" in xml and "<m:oMath>" in xml
    assert "mathbb" not in xml
    assert "𝔼" in xml and "ℙ" in xml
    assert "$" not in " ".join(paragraph.text for paragraph in document.paragraphs)
    assert document.part.element.xpath('//m:r[m:rPr/m:scr[@m:val="double-struck"]]/m:t')


def test_unsupported_formula_blocks_candidate(tmp_path, monkeypatch) -> None:
    workspace, _ = prepared(tmp_path, formula=r"\begin{matrix}a\end{matrix}")
    monkeypatch.setattr("mynotes.publication._render_and_validate", fake_render)
    with pytest.raises(PublicationError, match="Unsupported formula"):
        build_word_candidate(workspace, "OPS7-L01")


def test_build_embeds_visual_and_publish_backup_restore(tmp_path, monkeypatch) -> None:
    workspace, target = prepared(tmp_path, visual=True, formula=r"c = \frac{x}{y}")
    target.parent.mkdir(parents=True); base = Document()
    base.styles["Normal"].font.name = "Times New Roman"
    base.styles["Normal"].font.size = Pt(13)
    base.add_paragraph("Manual content outside MyNotes blocks.").add_run(" Preserved emphasis.").bold = True
    base.save(target)
    original = file_sha256(target)
    monkeypatch.setattr("mynotes.publication._render_and_validate", fake_render)
    built = build_word_candidate(workspace, "OPS7-L01")
    candidate = Path(built["candidate"])
    assert Document(candidate).inline_shapes
    candidate_document = Document(candidate)
    assert candidate_document.styles["Normal"].font.name == "Times New Roman"
    assert candidate_document.styles["Normal"].font.size == Pt(13)
    assert candidate_document.paragraphs[0].runs[-1].bold
    assert next(paragraph for paragraph in candidate_document.paragraphs if "MYNOTES START" in paragraph.text).paragraph_format.page_break_before
    result = publish_word_candidate(workspace, "OPS7-L01", candidate)
    assert result["status"] == "PUBLISHED" and file_sha256(target) != original
    assert status(workspace)[0]["status"] == "PUBLISHED"
    assert "Manual content outside" in "\n".join(p.text for p in Document(target).paragraphs)
    history = publication_history(workspace, "OPS7")
    assert len(history) == 1 and history[0]["sha256"] == original
    restored = restore_backup(workspace, history[0]["backup_id"])
    assert restored["status"] == "RESTORED" and file_sha256(target) == original


def test_external_edit_after_build_blocks_publish(tmp_path, monkeypatch) -> None:
    workspace, target = prepared(tmp_path)
    monkeypatch.setattr("mynotes.publication._render_and_validate", fake_render)
    built = build_word_candidate(workspace, "OPS7-L01")
    target.parent.mkdir(parents=True); Document().save(target)
    current = file_sha256(target)
    with pytest.raises(PublicationError, match="changed after candidate"):
        publish_word_candidate(workspace, "OPS7-L01", Path(built["candidate"]))
    assert file_sha256(target) == current


def test_ordered_blocks_keep_formula_table_and_visual_in_context(tmp_path, monkeypatch) -> None:
    workspace, _ = prepared(tmp_path, visual=True)
    source = workspace.runs_dir / "OPS7-L01" / "lecture_notes.candidate.json"
    payload = json.loads(source.read_text(encoding="utf-8"))
    payload["notes"]["blocks"] = [
        {"kind": "heading", "text": "Cost trade-off", "level": 3, "evidence_ids": ["S001"]},
        {"kind": "paragraph", "text": "**Holding cost** affects service.\\nA second explanation.", "evidence_ids": ["S001"]},
        {"kind": "formula", "label": "Cost", "expression": r"C = \frac{x}{y}", "evidence_ids": ["S001"]},
        {"kind": "table", "label": "Metrics table", "rows": [["Metric", "Value"], ["[[term:Cost]]", "10"]], "evidence_ids": ["S001"]},
        {"kind": "visual", "visual_id": "S001", "evidence_ids": ["S001"]},
        {"kind": "recap", "text": "Cost and service balance.", "evidence_ids": ["S001"]},
    ]
    payload["notes"]["formula"] = None
    payload["notes"]["visual_decisions"][0]["inspected"] = True
    atomic_json_write(source, payload)
    monkeypatch.setattr("mynotes.publication._render_and_validate", fake_render)
    result = build_word_candidate(workspace, "OPS7-L01")
    document = Document(result["candidate"])
    xml = document.part.element.xml
    assert xml.index("Cost trade-off") < xml.index("m:oMath") < xml.index("Metric") < xml.index("w:drawing") < xml.index("Cost and service balance.")
    assert "**" not in " ".join(paragraph.text for paragraph in document.paragraphs)
    assert "\\n" not in " ".join(paragraph.text for paragraph in document.paragraphs)
    assert next(p for p in document.paragraphs if p.text == "Cost").paragraph_format.keep_with_next
    assert next(p for p in document.paragraphs if p.text == "Metrics table").paragraph_format.keep_with_next
    assert len(document.tables) == 1
    assert all(row._tr.xpath("./w:trPr/w:cantSplit") for row in document.tables[0].rows)
    assert document.tables[0].cell(1, 0).text == "Cost"
    assert next(run for run in document.tables[0].cell(1, 0).paragraphs[0].runs if run.text == "Cost").style.name == "Termine tecnico"


def test_legacy_markdown_table_and_escaped_newline_become_native_word(tmp_path, monkeypatch) -> None:
    workspace, _ = prepared(tmp_path)
    source = workspace.runs_dir / "OPS7-L01" / "lecture_notes.candidate.json"
    payload = json.loads(source.read_text(encoding="utf-8"))
    payload["notes"]["sections"][0]["body"] = "First line\\n| Metric | Value |\\n| --- | --- |\\n| Cost | 10 |"
    atomic_json_write(source, payload)
    monkeypatch.setattr("mynotes.publication._render_and_validate", fake_render)
    result = build_word_candidate(workspace, "OPS7-L01")
    document = Document(result["candidate"])
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 0).text == "Cost"


def test_equation_text_is_literal_and_visual_caption_uses_filename(tmp_path, monkeypatch) -> None:
    workspace, _ = prepared(tmp_path, visual=True)
    pack_path = workspace.runs_dir / "OPS7-L01" / "source_pack.json"
    pack = json.loads(pack_path.read_text(encoding="utf-8"))
    pack["visuals"][0]["source_file"] = str(tmp_path / "private" / "deck.pdf")
    atomic_json_write(pack_path, pack)
    source = workspace.runs_dir / "OPS7-L01" / "lecture_notes.candidate.json"
    payload = json.loads(source.read_text(encoding="utf-8"))
    payload["notes"]["blocks"] = [
        {"kind": "formula", "label": "Service", "expression": r"FR = \frac{\text{Demand satisfied from stock}}{\text{Total demand}}", "evidence_ids": ["S001"]},
        {"kind": "visual", "visual_id": "S001", "evidence_ids": ["S001"]},
    ]
    payload["notes"]["visual_decisions"][0]["inspected"] = True
    payload["source_pack_hash"] = file_sha256(pack_path)
    atomic_json_write(source, payload)
    monkeypatch.setattr("mynotes.publication._render_and_validate", fake_render)
    result = build_word_candidate(workspace, "OPS7-L01")
    document = Document(result["candidate"])
    xml = document.part.element.xml
    assert "Demand satisfied from stock" in xml
    assert 'm:nor m:val="1"' in xml
    assert str(tmp_path) not in " ".join(paragraph.text for paragraph in document.paragraphs)


def test_study_profile_uses_neutral_template_and_character_style(tmp_path, monkeypatch) -> None:
    workspace, _ = prepared(tmp_path)
    candidate_path = workspace.runs_dir / "OPS7-L01" / "lecture_notes.candidate.json"
    candidate = json.loads(candidate_path.read_text(encoding="utf-8"))
    candidate["notes"]["blocks"] = [
        {"kind": "overview", "items": ["Demand", "Service"]},
        {"kind": "heading", "text": "Inventory", "level": 3},
        {"kind": "heading", "text": "Supporting detail", "level": 4},
        {"kind": "paragraph", "text": "[[term:Safety stock]] protects **service**."},
        {"kind": "paragraph", "text": "Safety stock is useful; safety stocking is a different phrase."},
        {"kind": "paragraph", "text": r"Order $Q^*=\mu+z\cdot\sigma$ units."},
        {"kind": "paragraph", "text": r"*(For discrete demand, $Q^*$ is the smallest integer satisfying $F(Q^*) \ge \frac{C_u}{C_u+C_o}$.)*"},
        {"kind": "bullets", "items": ["Cost", "Availability"]},
        {"kind": "qa", "question": "Why hold stock?", "answer": "To buffer demand."},
        {"kind": "recap", "text": "Match stock to demand."},
    ]
    atomic_json_write(candidate_path, candidate)
    monkeypatch.setattr("mynotes.publication._render_and_validate", fake_render)
    result = build_word_candidate(workspace, "OPS7-L01")
    document = Document(result["candidate"])
    section = document.sections[0]
    assert abs(section.page_width.cm - 21) < 0.01
    assert abs(section.page_height.cm - 29.7) < 0.01
    assert abs(section.left_margin.cm - 1.5) < 0.01
    assert abs(section.top_margin.cm - 1.5) < 0.01
    assert document.styles["MyNotes Body"].font.name == "Aptos"
    assert document.styles["MyNotes Body"].font.size == Pt(11)
    assert document.styles["MyNotes Body"].paragraph_format.space_after == Pt(3)
    assert abs(document.styles["MyNotes Body"].paragraph_format.line_spacing - 1.15) < 0.01
    for level, points in ((2, 16), (3, 14), (4, 12)):
        style = document.styles[f"Heading {level - 1}"]
        assert style.font.name == "Aptos Display"
        assert style.font.size == Pt(points)
        assert style.font.color.rgb == RGBColor(0x0F, 0x47, 0x61)
        assert style.paragraph_format.space_before == Pt(8)
        assert style.paragraph_format.space_after == Pt(4)
    caption_style = document.styles["MyNotes Caption"]
    assert caption_style.font.bold is True
    assert caption_style.font.italic is True
    assert caption_style.font.color.rgb == RGBColor(0x59, 0x59, 0x59)
    question_style = document.styles["MyNotes Question"]
    assert question_style.base_style.name == "Intense Quote"
    assert question_style.font.bold is True and question_style.font.italic is True
    assert question_style.font.color.rgb == RGBColor(0x4F, 0x81, 0xBD)
    footer_xml = section.footer._element.xml
    assert "PAGE" in footer_xml
    assert 'w:fldCharType="begin"' in footer_xml
    assert 'w:fldCharType="end"' in footer_xml
    assert "updateFields" in document.settings._element.xml
    assert section.footer.paragraphs[0].alignment == 1
    assert not any(paragraph.style.name == "Title" for paragraph in document.paragraphs)
    assert any(paragraph.text == "Lecture 1: Trade-off" for paragraph in document.paragraphs)
    assert [paragraph.text for paragraph in document.paragraphs if paragraph.style.name == "MyNotes List Number"] == ["Demand", "Service"]
    assert all(run.bold for paragraph in document.paragraphs if paragraph.style.name == "MyNotes List Number" for run in paragraph.runs)
    assert [paragraph.text for paragraph in document.paragraphs if paragraph.style.name == "MyNotes List Bullet"] == ["Cost", "Availability"]
    accent = next(run for paragraph in document.paragraphs for run in paragraph.runs if run.text == "Safety stock")
    assert accent.style.name == "Termine tecnico"
    assert accent.bold is None and accent.font.color.rgb is None
    assert next(p for p in document.paragraphs if p.text == "Supporting detail").style.name == "Heading 3"
    assert accent._r.xpath("./w:rPr/w:rStyle")
    assert not accent._r.xpath("./w:rPr/w:b | ./w:rPr/w:color")
    assert document.styles["Termine tecnico"].quick_style
    assert not document.styles["Termine tecnico"].hidden
    assert document.styles["Termine tecnico"].priority == 20
    assert all(not run._r.xpath("./w:rPr/w:b | ./w:rPr/w:i | ./w:rPr/w:color")
               for paragraph in document.paragraphs for run in paragraph.runs
               if run.style.name == "Termine tecnico")
    assert all(abs(margin.cm - 1.5) < 0.01 for margin in
               (section.left_margin, section.right_margin, section.top_margin, section.bottom_margin))
    assert next(p for p in document.paragraphs if p.text == "Lecture 1: Trade-off").style.name == "Heading 1"
    assert next(p for p in document.paragraphs if p.text == "Inventory").style.name == "Heading 2"
    assert sum(run.style.name == "Termine tecnico" for paragraph in document.paragraphs for run in paragraph.runs) == 2
    assert "m:oMath" in document.part.element.xml
    assert "$Q^*" not in document.part.element.xml
    assert "$F(Q^*)" not in document.part.element.xml
    assert "smallest integer" in document.part.element.xml
    assert document.styles["Termine tecnico"].font.bold
    assert document.styles["Termine tecnico"].font.color.rgb == RGBColor(0xED, 0x7D, 0x31)
    regular_bold = next(run for paragraph in document.paragraphs for run in paragraph.runs if run.text == "service")
    assert regular_bold.bold and regular_bold.font.color.rgb != RGBColor(0xED, 0x7D, 0x31)
    assert "Why hold stock?" in "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_style_preview_is_self_contained_and_preferences_change_output(tmp_path) -> None:
    sparse = tmp_path / "sparse.docx"
    rich = tmp_path / "rich.docx"
    preview_style_document({"writing_style": "concise", "list_density": "prose_first", "image_policy": "none", "question_count": 1, "recap": False}, sparse)
    preview_style_document({"writing_style": "detailed", "list_density": "list_first", "image_policy": "explanatory", "question_count": 3, "recap": True}, rich)
    first, second = Document(sparse), Document(rich)
    assert not first.inline_shapes and second.inline_shapes
    assert sum(paragraph.style.name == "MyNotes List Bullet" for paragraph in first.paragraphs) == 1
    assert sum(paragraph.style.name == "MyNotes List Bullet" for paragraph in second.paragraphs) == 4
    assert sum(paragraph.style.name == "MyNotes Question" for paragraph in first.paragraphs) == 1
    assert sum(paragraph.style.name == "MyNotes Question" for paragraph in second.paragraphs) == 2
    assert "Riepilogo in italiano" not in [paragraph.text for paragraph in first.paragraphs]
    assert "Riepilogo in italiano" in [paragraph.text for paragraph in second.paragraphs]
    with zipfile.ZipFile(rich) as archive:
        assert b"PAGE" in archive.read("word/footer1.xml")
        xml = archive.read("word/document.xml")
        assert b"MYNOTES START" not in xml
        assert b"oMath" in xml and b"tblHeader" in xml
        assert b"numPr" in xml
        numbering = archive.read("word/numbering.xml")
        assert "•".encode() in numbering


def test_image_policy_none_rejects_old_candidate_with_included_visual(tmp_path, monkeypatch) -> None:
    workspace, _ = prepared(tmp_path, visual=True)
    config = workspace.configuration_store().load()
    config["courses"][0]["style"]["image_policy"] = "none"
    workspace.configuration_store().save(config)
    monkeypatch.setattr("mynotes.publication._render_and_validate", fake_render)
    with pytest.raises(PublicationError, match="Image policy is none"):
        build_word_candidate(workspace, "OPS7-L01")
